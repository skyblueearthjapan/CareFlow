"""打刻履歴の Excel 出力 (4 シート).

正典設計書: ``docs/plans/visit-history-design-2026-09-30.md`` §3-2。
シート構成・列・書式は、つなぎの道具 ``docs/tools/visit-history/build_report.py`` の
``build_xlsx`` をそのまま移植したもの (お客様が受け取っている月次レポートと同じ形)。
Phase 2 (``actual-time-adjust-design-2026-09-30.md`` §6-3) で、到着・退出は実績時刻
(合わせた時刻) になり、QR を読み取った時刻を ``読取時刻（到着）`` ``読取時刻（退出）``
の 2 列に残す。備考は「時刻調整」(理由のある過去の調整だけ「時刻調整（理由）」。
理由は 2026-10-01 の PO 決定で画面から尋ねなくなった)。

  * ``QR読み取りあり`` — 到着の記録がある訪問
  * ``全予定``         — すべての行 ＋「QR 読み取り」(あり / なし) 列
  * ``看護師別``       — 看護師ごとの件数
  * ``読み方``         — 注意書きと備考の語彙の説明

時刻は JST・分単位 (秒は切り捨て)。DB にも現在時刻にも触らない純関数。
"""

from __future__ import annotations

from collections.abc import Sequence
from datetime import date, datetime, time
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.services.checkin.history import (
    JST,
    REPORT_TITLE,
    HistoryRow,
    caution_note,
    period_label,
    reading_notes,
    staff_counts,
    xlsx_remarks,
)

_WEEKDAYS = "月火水木金土日"
_HEAD_FILL = PatternFill("solid", fgColor="F3F6F8")
_THIN = Side(style="thin", color="D7DDE2")
_COLS = [
    "日付",
    "曜日",
    "拠点",
    "利用者",
    "予定開始",
    "予定終了",
    "予定の担当",
    "訪問した看護師",
    "到着",
    "退出",
    "滞在(分)",
    "読取時刻（到着）",
    "読取時刻（退出）",
    "備考",
]
_WIDTHS = [11, 5, 7, 18, 9, 9, 14, 14, 8, 8, 9, 15, 15, 38]
# 0 始まりの列番号: 時刻の列 (予定開始・予定終了・到着・退出・読取時刻 2 列) と、
# 太字にする実績の列。
_TIME_COLS = (4, 5, 8, 9, 11, 12)
_ACTUAL_COLS = (8, 9)


def _planned_time(value: str | None) -> time | None:
    """'HH:MM' → time (Excel の時刻セルにする)。"""
    if not value:
        return None
    return time(int(value[:2]), int(value[3:5]))


def _actual_time(value: datetime | None) -> time | None:
    """JST の打刻時刻 → 分単位の time (秒は切り捨て。画面の HH:MM と揃える)。"""
    if value is None:
        return None
    return time(value.hour, value.minute)


def _append(ws: Worksheet, values: Sequence[object]) -> None:
    """1 行足す。**文字列は必ず文字列のセルとして書く** (数式として解釈させない)。

    openpyxl は ``=`` で始まる文字列を数式のセルにする。氏名などの入力値がそのまま
    セルに入るので、``=HYPERLINK(...)`` のような値を数式にしない。
    """
    ws.append(list(values))
    for cell in ws[ws.max_row]:
        if cell.data_type == "f":
            cell.data_type = "s"


def _detail_sheet(ws: Worksheet, rows: Sequence[HistoryRow], *, with_state: bool) -> None:
    _append(ws, _COLS + (["QR 読み取り"] if with_state else []))
    for r in rows:
        line: list[object] = [
            r.visit_date,
            _WEEKDAYS[r.visit_date.weekday()],
            r.office_name or "",
            r.patient_name,
            _planned_time(r.start_time),
            _planned_time(r.end_time),
            r.planned_staff_name,
            r.actual_staff_name,
            _actual_time(r.arrival_jst),
            _actual_time(r.departure_jst),
            r.stay_minutes,
            _actual_time(r.arrival_read_at.astimezone(JST) if r.arrival_read_at else None),
            _actual_time(r.departure_read_at.astimezone(JST) if r.departure_read_at else None),
            xlsx_remarks(r),
        ]
        if with_state:
            line.append("あり" if r.arrival_at is not None else "なし")
        _append(ws, line)
    for i, width in enumerate(_WIDTHS + ([11] if with_state else []), 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = _HEAD_FILL
        cell.alignment = Alignment(horizontal="center")
        cell.border = Border(bottom=Side(style="thin", color="9AA7B0"))
    for row in ws.iter_rows(min_row=2):
        row[0].number_format = "yyyy/mm/dd"
        for i in _TIME_COLS:
            row[i].number_format = "hh:mm"
            row[i].alignment = Alignment(horizontal="center")
        row[1].alignment = Alignment(horizontal="center")
        for i in _ACTUAL_COLS:
            row[i].font = Font(bold=True)
        for cell in row:
            cell.border = Border(bottom=_THIN)
    ws.freeze_panes = "E2"
    ws.auto_filter.ref = ws.dimensions
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "1:1"


def build_history_xlsx(rows: Sequence[HistoryRow], *, date_from: date, date_to: date) -> bytes:
    """打刻履歴の行 (絞り込み・並び替え済み・``future`` を除いたもの) → .xlsx のバイト列。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "QR読み取りあり"
    _detail_sheet(ws, [r for r in rows if r.arrival_at is not None], with_state=False)

    _detail_sheet(wb.create_sheet("全予定"), rows, with_state=True)

    ws3 = wb.create_sheet("看護師別")
    _append(ws3, ["看護師", "予定の件数", "到着の読み取り", "退出の読み取り", "退出なし"])
    for c in staff_counts(rows):
        _append(ws3, [c.name, c.planned, c.arrival, c.departure, c.no_departure])
    for cell in ws3[1]:
        cell.font = Font(bold=True)
        cell.fill = _HEAD_FILL
    for i, width in enumerate([16, 12, 14, 14, 10], 1):
        ws3.column_dimensions[get_column_letter(i)].width = width

    ws4 = wb.create_sheet("読み方")
    ws4.column_dimensions["A"].width = 110
    _append(ws4, [f"{REPORT_TITLE} {period_label(date_from, date_to)}"])
    ws4["A1"].font = Font(bold=True, size=13)
    for note in [caution_note(rows), *reading_notes(rows)]:
        _append(ws4, [note])
    for row in ws4.iter_rows(min_row=2):
        row[0].alignment = Alignment(wrap_text=True, vertical="top")

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
