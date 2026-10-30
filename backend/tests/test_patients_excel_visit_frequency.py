"""患者マスタ Excel の訪問頻度・起点の週・第何週・NG 曜日・枠の「訪問する週」.

正典: docs/plans/biweekly-monthly-design-2026-10-30.md §3 / §5「Excel」。

- 書き出し → そのまま取り込み で差分が出ない (往復で noop)。
- 古い表記 (隔週 / 月次)・古い列の並び (時間タイプの列あり) のファイルも取り込める。
- 月曜でない起点の週は、その週の月曜にそろえて知らせる。
- 読めない値は行エラー。
- 完全置換・カルテでも同じ項目を運ぶ。
"""

from __future__ import annotations

from datetime import date, time
from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import select

from app.core.security import create_access_token, hash_password
from app.models import Patient, PatientFixedVisit, User
from app.services.patient_excel.importer import _merge_weekly_pattern, parse_weekly_extras
from app.services.patient_excel.karte import (
    KARTE_ANCHOR_CELL,
    KARTE_NG_CELL,
    KARTE_VISIT_WEEKS_CELL,
    SHEET_KARTE,
)
from app.services.patient_excel.schema import (
    LEGACY_PFV_COLUMNS,
    LEGACY_PFV_EDIT_COLUMNS,
    PATIENT_COL_INDEX,
    PATIENT_COLUMNS,
    PFV_EDIT_COL_INDEX,
    PFV_EDIT_COLUMNS,
    SHEET_PATIENTS,
    SHEET_PFV,
    SHEET_PFV_EDIT,
    VISIT_FREQUENCY_JA_VALUES,
    parse_pfv_visit_week_cell,
)

pytestmark = pytest.mark.usefixtures("fake_visit_geocode")

_WP = {
    "frequency_per_week": 1,
    "visit_frequency": "biweekly",
    "biweekly_anchor": "2026-11-02",
    "preferred_weekdays": ["Tue", "Fri"],
    "ng_weekdays": ["Mon", "Wed"],
    "service_minutes": 30,
    "time_type": "時間帯",
}


async def _make_user(db, email: str) -> User:
    user = User(email=email, password_hash=hash_password("pw"), role="admin")
    db.add(user)
    await db.commit()
    await db.refresh(user)
    return user


def _bearer(user: User) -> dict[str, str]:
    token = create_access_token(subject=user.id, role=user.role, staff_id=user.staff_id)
    return {"Authorization": f"Bearer {token}"}


async def _make_patient(db, **overrides) -> Patient:
    values = {
        "code": "P-FREQ",
        "name": "頻度 花子",
        "sex": "female",
        "status": "active",
        "address": "千葉市稲毛区freq",
    }
    values.update(overrides)
    p = Patient(**values)
    db.add(p)
    await db.commit()
    await db.refresh(p)
    return p


async def _make_pfv(
    db, *, patient_id, weekday: int, visit_week: str = "every", slot_index: int = 0
) -> PatientFixedVisit:
    pfv = PatientFixedVisit(
        patient_id=patient_id,
        mode="normal",
        weekday=weekday,
        slot_index=slot_index,
        start_time=time(9, 0),
        duration_min=30,
        visit_week=visit_week,
    )
    db.add(pfv)
    await db.commit()
    await db.refresh(pfv)
    return pfv


async def _post(client, user, path: str, content: bytes, *, dry_run: bool):
    files = {
        "file": (
            "t.xlsx",
            content,
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
    }
    return await client.post(
        f"/api/v1/patients/import-export/{path}?dry_run={'true' if dry_run else 'false'}",
        headers=_bearer(user),
        files=files,
    )


def _book(
    *,
    patient_rows: list[dict],
    edit_rows: list[dict] | None = None,
    edit_columns: list[dict] | None = None,
) -> bytes:
    """患者マスタ + 編集用シートの workbook (列キーの dict で値を書く)."""
    wb = Workbook()
    ws_p = wb.active
    ws_p.title = SHEET_PATIENTS
    for i, col in enumerate(PATIENT_COLUMNS, start=1):
        ws_p.cell(row=1, column=i, value=str(col["header"]))
    for r, row in enumerate(patient_rows, start=2):
        for key, idx in PATIENT_COL_INDEX.items():
            if row.get(key) is not None:
                ws_p.cell(row=r, column=idx + 1, value=row[key])
    cols = edit_columns if edit_columns is not None else PFV_EDIT_COLUMNS
    ws_f = wb.create_sheet(SHEET_PFV_EDIT)
    for i, col in enumerate(cols, start=1):
        ws_f.cell(row=1, column=i, value=str(col["header"]))
    for r, row in enumerate(edit_rows or [], start=2):
        for i, col in enumerate(cols, start=1):
            v = row.get(str(col["key"]))
            if v is not None:
                ws_f.cell(row=r, column=i, value=v)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# 純関数
# ---------------------------------------------------------------------------


def test_frequency_labels_follow_the_screen() -> None:
    assert VISIT_FREQUENCY_JA_VALUES == ("毎週", "隔週（2 週おき）", "月の第何週")


def test_parse_pfv_visit_week_cell() -> None:
    assert parse_pfv_visit_week_cell(None) is None
    assert parse_pfv_visit_week_cell(" ") is None
    assert parse_pfv_visit_week_cell("毎回") == "every"
    assert parse_pfv_visit_week_cell("A の週") == "A"
    assert parse_pfv_visit_week_cell("Ｂの週") == "B"
    assert parse_pfv_visit_week_cell("every") == "every"
    with pytest.raises(ValueError):
        parse_pfv_visit_week_cell("Cの週")


def test_parse_weekly_extras_snaps_anchor_and_validates() -> None:
    updates, errors, infos = parse_weekly_extras(
        {
            "visit_frequency": "隔週",  # 古い表記
            "biweekly_anchor": "2026/11/4",  # 水曜 → 11/2 (月) にそろえる
            "visit_weeks": "１、３",
            "ng_wd_mon": "〇",
            "ng_wd_tue": "×",
        }
    )
    assert errors == []
    assert updates["biweekly_anchor"] == "2026-11-02"
    assert updates["visit_weeks"] == "1,3"
    assert updates["ng_weekdays"] == ["Mon"]
    assert len(infos) == 1 and "11/2" in infos[0]

    _, errors, _ = parse_weekly_extras({"visit_weeks": "1,6"})
    assert errors and "1〜5" in errors[0]
    _, errors, _ = parse_weekly_extras({"biweekly_anchor": "来週"})
    assert errors and "起点の週" in errors[0]
    _, errors, _ = parse_weekly_extras({"visit_frequency": "毎月"})
    assert errors and "訪問頻度" in errors[0]

    # 空 = 維持、<CLEAR> = 消す。
    updates, errors, _ = parse_weekly_extras({"biweekly_anchor": "<CLEAR>", "visit_weeks": None})
    assert errors == [] and updates == {"biweekly_anchor": None}


def test_merge_keeps_other_keys_and_handles_clear_and_empty_ng() -> None:
    existing = {"entries": [{"weekday": "Tue"}], "biweekly_anchor": "2026-11-02"}
    merged = _merge_weekly_pattern(existing, {"biweekly_anchor": None, "ng_weekdays": []})
    # entries は残す・起点の週は消す・「NG 曜日なし」はキーを足さない。
    assert merged == {"entries": [{"weekday": "Tue"}]}
    merged = _merge_weekly_pattern({"ng_weekdays": ["Mon"]}, {"ng_weekdays": []})
    assert merged == {"ng_weekdays": []}
    # 希望の無い方に「NG 曜日なし」だけ → 変えない。
    assert _merge_weekly_pattern(None, {"ng_weekdays": []}) is None


# ---------------------------------------------------------------------------
# 書き出し → 取り込み
# ---------------------------------------------------------------------------


@pytest.mark.asyncio
async def test_export_writes_new_columns_and_roundtrip_is_noop(client, db) -> None:
    admin = await _make_user(db, "freq-rt@example.com")
    p = await _make_patient(db, weekly_pattern=dict(_WP))
    await _make_pfv(db, patient_id=p.id, weekday=1, visit_week="A")
    await _make_pfv(db, patient_id=p.id, weekday=4, visit_week="B")

    res = await client.get("/api/v1/patients/import-export/export", headers=_bearer(admin))
    assert res.status_code == 200
    wb = load_workbook(BytesIO(res.content))
    ws = wb[SHEET_PATIENTS]
    col = {k: i + 1 for k, i in PATIENT_COL_INDEX.items()}
    assert ws.cell(row=2, column=col["visit_frequency"]).value == "隔週（2 週おき）"
    assert ws.cell(row=2, column=col["biweekly_anchor"]).value == "2026-11-02"
    assert ws.cell(row=2, column=col["ng_wd_mon"]).value == "〇"
    assert ws.cell(row=2, column=col["ng_wd_tue"]).value == "×"
    assert ws.cell(row=2, column=col["ng_wd_wed"]).value == "〇"
    ws_f = wb[SHEET_PFV_EDIT]
    headers = [c.value for c in ws_f[1]]
    assert "時間タイプ" not in headers
    assert ws_f.cell(row=2, column=PFV_EDIT_COL_INDEX["tue_week"] + 1).value == "Aの週"
    assert ws_f.cell(row=2, column=PFV_EDIT_COL_INDEX["fri_week"] + 1).value == "Bの週"

    res = await _post(client, admin, "import", res.content, dry_run=True)
    assert res.status_code == 200, res.text
    s = res.json()["summary"]
    assert s["patients_error"] == 0 and s["pfv_error"] == 0, res.text
    assert s["patients_update"] == 0 and s["pfv_update"] == 0, res.text


@pytest.mark.asyncio
async def test_import_sets_frequency_anchor_ng_and_visit_week(client, db) -> None:
    admin = await _make_user(db, "freq-in@example.com")
    p = await _make_patient(
        db, weekly_pattern={"frequency_per_week": 2, "entries": [{"weekday": "Tue"}]}
    )
    tue = await _make_pfv(db, patient_id=p.id, weekday=1)
    tue2 = await _make_pfv(db, patient_id=p.id, weekday=1, slot_index=1)
    content = _book(
        patient_rows=[
            {
                "patient_id": str(p.id),
                "patient_code": p.code,
                "visit_frequency": "隔週",
                "biweekly_anchor": date(2026, 11, 5),  # 木曜 → 11/2
                "ng_wd_sat": "〇",
            }
        ],
        edit_rows=[
            {
                "patient_id": str(p.id),
                "patient_code": p.code,
                "service_minutes": 30,
                "tue_time": "09:00",
                "tue_week": "Bの週",
            }
        ],
    )
    res = await _post(client, admin, "import", content, dry_run=False)
    assert res.status_code == 200, res.text
    body = res.json()
    assert body["summary"]["patients_error"] == 0, body
    row = body["patient_rows"][0]
    assert row["operation"] == "update"
    assert any(
        c["field"] == "biweekly_anchor" and "そろえました" in str(c["new_value"])
        for c in row["changes"]
    )

    await db.refresh(p)
    assert p.weekly_pattern["visit_frequency"] == "biweekly"
    assert p.weekly_pattern["biweekly_anchor"] == "2026-11-02"
    assert p.weekly_pattern["ng_weekdays"] == ["Sat"]
    # Excel の管理外のキーは消さない。
    assert p.weekly_pattern["entries"] == [{"weekday": "Tue"}]
    rows = {
        r.slot_index: r
        for r in (
            await db.scalars(select(PatientFixedVisit).where(PatientFixedVisit.patient_id == p.id))
        ).all()
    }
    await db.refresh(rows[0])
    await db.refresh(rows[1])
    assert rows[0].id == tue.id and rows[0].visit_week == "B"
    # 同じ曜日の 2 人目の枠もそろえる。
    assert rows[1].id == tue2.id and rows[1].visit_week == "B"


@pytest.mark.asyncio
async def test_import_rejects_bad_frequency_values(client, db) -> None:
    admin = await _make_user(db, "freq-bad@example.com")
    p = await _make_patient(db, weekly_pattern={"frequency_per_week": 1})
    await _make_pfv(db, patient_id=p.id, weekday=0)
    content = _book(
        patient_rows=[
            {"patient_id": str(p.id), "patient_code": p.code, "visit_weeks": "0,3"},
        ],
        edit_rows=[
            {
                "patient_id": str(p.id),
                "patient_code": p.code,
                "mon_time": "09:00",
                "mon_week": "Cの週",
            }
        ],
    )
    res = await _post(client, admin, "import", content, dry_run=True)
    assert res.status_code == 200, res.text
    body = res.json()
    assert body["patient_rows"][0]["operation"] == "error"
    assert "第何週" in body["patient_rows"][0]["error_message"]
    assert body["summary"]["pfv_error"] == 1
    assert "訪問する週" in body["pfv_rows"][0]["error_message"]


@pytest.mark.asyncio
async def test_old_format_file_still_imports(client, db) -> None:
    """2026-10-30 より前の並び (編集用シートに時間タイプの列) と古い表記."""
    admin = await _make_user(db, "freq-old@example.com")
    p = await _make_patient(db, weekly_pattern={"frequency_per_week": 1, "ng_weekdays": ["Sun"]})
    await _make_pfv(db, patient_id=p.id, weekday=0, visit_week="A")
    content = _book(
        patient_rows=[
            {"patient_id": str(p.id), "patient_code": p.code, "visit_frequency": "月次"},
        ],
        edit_rows=[
            {
                "patient_id": str(p.id),
                "patient_code": p.code,
                "service_minutes": 30,
                "time_type": "固定",
                "mon_time": "09:00",
                "wed_time": "10:00",
            }
        ],
        edit_columns=LEGACY_PFV_EDIT_COLUMNS,
    )
    res = await _post(client, admin, "import", content, dry_run=False)
    assert res.status_code == 200, res.text
    body = res.json()
    assert body["summary"]["patients_error"] == 0 and body["summary"]["pfv_error"] == 0, body
    assert body["summary"]["pfv_new"] == 1  # 水曜
    await db.refresh(p)
    assert p.weekly_pattern["visit_frequency"] == "monthly"
    # NG 曜日の列の無いファイルは NG 曜日を変えない。
    assert p.weekly_pattern["ng_weekdays"] == ["Sun"]
    pfvs = {
        r.weekday: r
        for r in (
            await db.scalars(select(PatientFixedVisit).where(PatientFixedVisit.patient_id == p.id))
        ).all()
    }
    await db.refresh(pfvs[0])
    assert pfvs[0].visit_week == "A"  # 列が無い = 変えない
    assert pfvs[2].visit_week == "every"


@pytest.mark.asyncio
async def test_old_per_visit_sheet_with_time_type_column_imports(client, db) -> None:
    """旧 1 枠 1 行のシート (時間タイプの列あり) は列の位置がずれずに読める."""
    admin = await _make_user(db, "freq-old-pv@example.com")
    p = await _make_patient(db, weekly_pattern={"frequency_per_week": 1})
    wb = Workbook()
    ws_p = wb.active
    ws_p.title = SHEET_PATIENTS
    for i, col in enumerate(PATIENT_COLUMNS, start=1):
        ws_p.cell(row=1, column=i, value=str(col["header"]))
    ws_p.cell(row=2, column=PATIENT_COL_INDEX["patient_id"] + 1, value=str(p.id))
    ws_p.cell(row=2, column=PATIENT_COL_INDEX["patient_code"] + 1, value=p.code)
    ws_f = wb.create_sheet(SHEET_PFV)
    values = {
        "patient_id": str(p.id),
        "weekday": "木",
        "slot_index": 0,
        "mode": "通常",
        "time_type": "どれでもよい",  # 使わない列 (候補外でもエラーにしない)
        "start_time": "13:30",
        "duration_min": 45,
    }
    for i, col in enumerate(LEGACY_PFV_COLUMNS, start=1):
        ws_f.cell(row=1, column=i, value=str(col["header"]))
        if values.get(str(col["key"])) is not None:
            ws_f.cell(row=2, column=i, value=values[str(col["key"])])
    buf = BytesIO()
    wb.save(buf)

    res = await _post(client, admin, "import", buf.getvalue(), dry_run=False)
    assert res.status_code == 200, res.text
    body = res.json()
    assert body["summary"]["pfv_error"] == 0 and body["summary"]["pfv_new"] == 1, body
    pfv = (
        await db.scalars(select(PatientFixedVisit).where(PatientFixedVisit.patient_id == p.id))
    ).one()
    assert (pfv.weekday, pfv.start_time, pfv.duration_min) == (3, time(13, 30), 45)
    assert pfv.visit_week == "every"


@pytest.mark.asyncio
async def test_replace_all_keeps_visit_week_and_frequency(client, db) -> None:
    admin = await _make_user(db, "freq-ra@example.com")
    p = await _make_patient(db, weekly_pattern=dict(_WP))
    await _make_pfv(db, patient_id=p.id, weekday=1, visit_week="A")
    await _make_pfv(db, patient_id=p.id, weekday=4, visit_week="B")
    res = await client.get("/api/v1/patients/import-export/export", headers=_bearer(admin))
    assert res.status_code == 200

    res = await _post(client, admin, "replace-all", res.content, dry_run=False)
    assert res.status_code == 200, res.text
    pfvs = {
        r.weekday: r.visit_week
        for r in (
            await db.scalars(select(PatientFixedVisit).where(PatientFixedVisit.patient_id == p.id))
        ).all()
    }
    assert pfvs == {1: "A", 4: "B"}
    await db.refresh(p)
    assert p.weekly_pattern["biweekly_anchor"] == "2026-11-02"
    assert p.weekly_pattern["ng_weekdays"] == ["Mon", "Wed"]


# ---------------------------------------------------------------------------
# カルテ
# ---------------------------------------------------------------------------


@pytest.mark.asyncio
async def test_karte_shows_frequency_and_roundtrip_is_noop(client, db) -> None:
    admin = await _make_user(db, "freq-karte@example.com")
    wp = {**_WP, "visit_frequency": "monthly", "visit_weeks": "1,3"}
    p = await _make_patient(db, weekly_pattern=wp)
    await _make_pfv(db, patient_id=p.id, weekday=1, visit_week="B")

    res = await client.get(
        f"/api/v1/patients/import-export/karte-export?patient_id={p.id}",
        headers=_bearer(admin),
    )
    assert res.status_code == 200, res.text
    ws = load_workbook(BytesIO(res.content))[SHEET_KARTE]
    assert ws["D9"].value == "月の第何週"
    assert ws[KARTE_ANCHOR_CELL].value == "2026-11-02"
    assert ws[KARTE_VISIT_WEEKS_CELL].value == "1,3"
    assert ws[KARTE_NG_CELL].value == "月,水"

    res2 = await _post(client, admin, "karte-import", res.content, dry_run=True)
    assert res2.status_code == 200, res2.text
    s = res2.json()["summary"]
    assert s["patients_update"] == 0 and s["patients_error"] == 0, res2.text
    assert s["pfv_update"] == 0 and s["pfv_error"] == 0, res2.text

    # カルテで NG 曜日を「なし」にすると消える。
    wb = load_workbook(BytesIO(res.content))
    wb[SHEET_KARTE][KARTE_NG_CELL] = "なし"
    buf = BytesIO()
    wb.save(buf)
    res3 = await _post(client, admin, "karte-import", buf.getvalue(), dry_run=False)
    assert res3.status_code == 200, res3.text
    await db.refresh(p)
    assert p.weekly_pattern["ng_weekdays"] == []
    assert p.weekly_pattern["visit_weeks"] == "1,3"
