"""打刻履歴 API (``/api/v1/visit-history``) のテスト — Phase 1.

正典設計書: ``docs/plans/visit-history-design-2026-09-30.md`` §5 (BE)。

* 行の範囲: 取消済み＋打刻ありは載る / 取消済み＋打刻なしは載らない /
  【検証】等の名前による特別扱いはしない。
* state 5 種・備考の語彙と順序・予定外は予定欄 null・代行判定。
* staff ロールの絞り (他人の分が見えない・``staff_id`` 指定でも自分に固定・
  主担当が空でコース担当だけの訪問は見える)。
* ``groups`` (見出し行の件数) がページングに左右されない。
* Excel / A4 の出力が ``audit_logs`` に残る。
* 期間 92 日超は 422・summary がページングに左右されない。
* xlsx が開けてシート 4 枚・report が ``no-store`` (JS 無しでも読める HTML)。

日付は「今日 (JST)」からの相対で組む (昨日 = 過去・明日 = 未来)。「今日」は実時計では
なく固定の時刻 ``_NOW`` から取り、API 側の時計 (``visit_history.py`` の
``datetime.now``) も同じ時刻に固定する (``_fixed_clock``)。実時計に依存すると、
JST 0 時をまたいで走ったときにテスト側の「今日」とサーバ側の「今日」がずれる。
TX 後始末: 各テストの終端で ``await db.rollback()``。
"""

from __future__ import annotations

from datetime import UTC, date, datetime, time, timedelta
from io import BytesIO
from zoneinfo import ZoneInfo

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app.api.v1 import visit_history as visit_history_api
from app.core.security import create_access_token, hash_password
from app.models import (
    AuditLog,
    Course,
    Office,
    Patient,
    Staff,
    User,
    Visit,
    VisitCheckin,
    VisitStaffAssignment,
)
from app.models.accompaniment import Accompaniment
from app.services.checkin.actuals import load_actuals

JST = ZoneInfo("Asia/Tokyo")
URL = "/api/v1/visit-history"

#: テストの「いま」(水曜の昼)。
_NOW = datetime(2026, 9, 30, 12, 0, tzinfo=JST)


@pytest.fixture(autouse=True)
def _fixed_clock(monkeypatch) -> None:
    """``visit_history.py`` の時計を ``_NOW`` に固定する (``_today()`` と同じ日になる)。"""

    class _FrozenDatetime(datetime):
        @classmethod
        def now(cls, tz=None):  # noqa: ANN001, ANN206 - stdlib シグネチャに合わせる
            return _NOW.astimezone(tz) if tz is not None else _NOW.replace(tzinfo=None)

    monkeypatch.setattr(visit_history_api, "datetime", _FrozenDatetime)


def _today() -> date:
    return _NOW.date()


def _yesterday() -> date:
    return _today() - timedelta(days=1)


def _tomorrow() -> date:
    return _today() + timedelta(days=1)


def _bearer(user: User) -> dict[str, str]:
    token = create_access_token(subject=user.id, role=user.role, staff_id=user.staff_id)
    return {"Authorization": f"Bearer {token}"}


def _range(day_from: date | None = None, day_to: date | None = None) -> dict[str, str]:
    return {
        "from": (day_from or _yesterday()).isoformat(),
        "to": (day_to or _tomorrow()).isoformat(),
    }


async def _admin(db, email: str = "vh-admin@example.com") -> User:
    user = User(email=email, password_hash=hash_password("x"), role="admin")
    db.add(user)
    await db.commit()
    await db.refresh(user)
    return user


async def _staff(db, name: str) -> Staff:
    staff = Staff(name=name)
    db.add(staff)
    await db.commit()
    await db.refresh(staff)
    return staff


async def _staff_user(db, email: str, staff: Staff | None) -> User:
    user = User(
        email=email,
        password_hash=hash_password("x"),
        role="staff",
        staff_id=staff.id if staff is not None else None,
    )
    db.add(user)
    await db.commit()
    await db.refresh(user)
    return user


async def _patient(db, code: str, name: str, office: Office | None = None) -> Patient:
    patient = Patient(
        code=code, name=name, primary_office_id=office.id if office is not None else None
    )
    db.add(patient)
    await db.commit()
    await db.refresh(patient)
    return patient


async def _visit(
    db,
    patient: Patient,
    staff: Staff | None,
    day: date,
    *,
    start: time = time(9, 0),
    end: time = time(10, 0),
    status: str = "planned",
    is_unplanned: bool = False,
    deleted: bool = False,
) -> Visit:
    visit = Visit(
        patient_id=patient.id,
        primary_staff_id=staff.id if staff is not None else None,
        visit_date=day,
        start_time=start,
        end_time=end,
        type="regular",
        status=status,
        is_unplanned=is_unplanned,
        deleted_at=datetime.now(UTC) if deleted else None,
    )
    db.add(visit)
    await db.commit()
    await db.refresh(visit)
    return visit


async def _checkin(
    db,
    visit: Visit,
    staff: Staff,
    kind: str,
    at: time,
    *,
    second: int = 0,
    match_status: str = "match",
    source: str = "qr",
) -> VisitCheckin:
    """``visit.visit_date`` の JST ``at`` に打刻した行を直接入れる。"""
    moment = datetime.combine(visit.visit_date, at.replace(second=second), tzinfo=JST)
    row = VisitCheckin(
        visit_id=visit.id,
        patient_id=visit.patient_id,
        staff_id=staff.id,
        kind=kind,
        scanned_at=moment.astimezone(UTC),
        match_status=match_status,
        checkin_source=source,
        threshold_snapshot={"v": 1},
    )
    db.add(row)
    await db.commit()
    await db.refresh(row)
    return row


def _by_patient(body: dict) -> dict[str, dict]:
    return {item["patient_name"]: item for item in body["items"]}


# ---------------------------------------------------------------------------
# 行の範囲
# ---------------------------------------------------------------------------


async def test_cancelled_visit_is_listed_only_when_it_has_a_checkin(client, db) -> None:
    """取消・削除済みでも打刻があれば載せる。打刻が無ければ載せない。"""
    admin = await _admin(db)
    nurse = await _staff(db, "看護 一郎")
    day = _yesterday()
    kept = await _visit(
        db, await _patient(db, "VH-1", "取消打刻あり"), nurse, day, status="cancelled"
    )
    await _checkin(db, kept, nurse, "arrival", time(9, 2))
    await _visit(db, await _patient(db, "VH-2", "取消打刻なし"), nurse, day, status="cancelled")
    gone = await _visit(db, await _patient(db, "VH-3", "削除打刻あり"), nurse, day, deleted=True)
    await _checkin(db, gone, nurse, "departure", time(9, 40))
    await _visit(db, await _patient(db, "VH-4", "削除打刻なし"), nurse, day, deleted=True)
    # no_show は実績に数えない = 取消済みの訪問を載せる理由にならない。
    no_show = await _visit(
        db, await _patient(db, "VH-5", "取消未訪問記録"), nurse, day, status="cancelled"
    )
    await _checkin(db, no_show, nurse, "no_show", time(9, 30))

    res = await client.get(URL, headers=_bearer(admin), params=_range())
    assert res.status_code == 200, res.text
    rows = _by_patient(res.json())
    assert set(rows) == {"取消打刻あり", "削除打刻あり"}
    assert rows["取消打刻あり"]["is_cancelled"] is True
    assert "取消済みの予定に記録" in rows["取消打刻あり"]["remarks"]
    assert rows["削除打刻あり"]["is_cancelled"] is True
    await db.rollback()


async def test_verification_patient_is_not_special_cased(client, db) -> None:
    """【検証】で始まる患者も普通の 1 行 (つなぎの道具と違い、名前で除外しない)。"""
    admin = await _admin(db)
    nurse = await _staff(db, "看護 一郎")
    await _visit(db, await _patient(db, "VH-V", "【検証】テスト 太郎"), nurse, _yesterday())

    res = await client.get(URL, headers=_bearer(admin), params=_range())
    assert [i["patient_name"] for i in res.json()["items"]] == ["【検証】テスト 太郎"]
    await db.rollback()


async def test_visits_outside_the_period_are_not_listed(client, db) -> None:
    admin = await _admin(db)
    nurse = await _staff(db, "看護 一郎")
    await _visit(db, await _patient(db, "VH-IN", "期間内"), nurse, _yesterday())
    await _visit(db, await _patient(db, "VH-OUT", "期間外"), nurse, _today() - timedelta(days=5))

    res = await client.get(URL, headers=_bearer(admin), params=_range())
    assert set(_by_patient(res.json())) == {"期間内"}
    await db.rollback()


# ---------------------------------------------------------------------------
# state / 実績時刻
# ---------------------------------------------------------------------------


async def test_five_states(client, db) -> None:
    admin = await _admin(db)
    nurse = await _staff(db, "看護 一郎")
    done = await _visit(db, await _patient(db, "VH-S1", "完了"), nurse, _yesterday())
    await _checkin(db, done, nurse, "arrival", time(9, 1))
    await _checkin(db, done, nurse, "departure", time(9, 45))
    in_progress = await _visit(
        db,
        await _patient(db, "VH-S2", "訪問中"),
        nurse,
        _today(),
        start=time(0, 0),
        end=time(0, 30),
    )
    await _checkin(db, in_progress, nurse, "arrival", time(0, 0))
    no_departure = await _visit(db, await _patient(db, "VH-S3", "退出なし"), nurse, _yesterday())
    await _checkin(db, no_departure, nurse, "arrival", time(9, 1))
    await _visit(db, await _patient(db, "VH-S4", "打刻なし"), nurse, _yesterday())
    await _visit(db, await _patient(db, "VH-S5", "これから"), nurse, _tomorrow())

    res = await client.get(URL, headers=_bearer(admin), params=_range())
    assert res.status_code == 200, res.text
    body = res.json()
    states = {name: item["state"] for name, item in _by_patient(body).items()}
    assert states == {
        "完了": "done",
        "訪問中": "in_progress",
        "退出なし": "no_departure",
        "打刻なし": "none",
        "これから": "future",
    }
    # summary は future を除いて数える。total は future を含む。
    assert body["total"] == 5
    assert body["summary"] == {
        "visits": 4,
        "with_arrival": 3,
        "with_departure": 1,
        "no_departure": 1,
        "none": 1,
        "adjusted": 0,
    }
    await db.rollback()


async def test_actual_times_use_latest_checkin_per_kind(client, db) -> None:
    """再スキャンは最新を採る。時刻は UTC で返し、滞在は JST の分に切り捨てて引く。"""
    admin = await _admin(db)
    nurse = await _staff(db, "看護 一郎")
    other = await _staff(db, "看護 二郎")
    visit = await _visit(db, await _patient(db, "VH-T1", "打ち直し"), nurse, _yesterday())
    await _checkin(db, visit, other, "arrival", time(8, 50), match_status="mismatch")
    await _checkin(db, visit, nurse, "arrival", time(9, 0), second=50)
    await _checkin(db, visit, nurse, "departure", time(9, 30), second=10)
    await _checkin(db, visit, nurse, "no_show", time(9, 40))

    res = await client.get(URL, headers=_bearer(admin), params=_range())
    row = res.json()["items"][0]
    expected_arrival = datetime.combine(_yesterday(), time(9, 0, 50), tzinfo=JST).astimezone(UTC)
    assert datetime.fromisoformat(row["arrival_at"]) == expected_arrival
    assert datetime.fromisoformat(row["arrival_at"]).utcoffset() == timedelta(0)
    # 09:00:50 → 09:30:10 は秒まで引くと 29 分 20 秒。分に切り捨ててから引くので 30。
    assert row["stay_minutes"] == 30
    assert row["actual_staff_name"] == "看護 一郎"
    # 位置判定は最新の到着 (担当本人・位置 OK) で見る = 打ち直す前の不一致は出ない。
    assert row["match_status"] == "match"
    # 代行は訪問モニターと同じ判定: 到着・退出の全打刻者のいずれかが担当集合の外。
    # 担当本人が後から打ち直しても、先に担当外が打った事実は消さない。
    assert row["is_substitute"] is True
    assert row["remarks"] == ["代行（予定: 看護 一郎）"]
    # 調整が無ければ読取時刻 = 実績時刻。
    assert row["arrival_read_at"] == row["arrival_at"]
    assert row["arrival_adjusted"] is False
    assert row["departure_manual"] is False
    assert row["adjustments"] == []

    actuals = await load_actuals(db, [visit.id])
    assert actuals[visit.id].checkin_staff_ids == [nurse.id, other.id]
    await db.rollback()


async def test_departure_only_visit_shows_the_departure_staff(client, db) -> None:
    """到着を読み忘れて退出だけがある訪問: 打った人を出し、future にはしない。"""
    admin = await _admin(db)
    nurse = await _staff(db, "看護 一郎")
    visit = await _visit(db, await _patient(db, "VH-T2", "退出だけ"), nurse, _tomorrow())
    await _checkin(db, visit, nurse, "departure", time(9, 40))

    res = await client.get(URL, headers=_bearer(admin), params=_range())
    row = res.json()["items"][0]
    assert row["arrival_at"] is None
    assert row["departure_at"] is not None
    assert row["actual_staff_name"] == "看護 一郎"
    assert row["stay_minutes"] is None
    assert row["state"] == "none"
    await db.rollback()


# ---------------------------------------------------------------------------
# 備考 / 予定外 / 代行
# ---------------------------------------------------------------------------


async def test_remarks_vocabulary_and_order(client, db) -> None:
    admin = await _admin(db)
    owner = await _staff(db, "予定 太郎")
    sub = await _staff(db, "代行 花子")
    visit = await _visit(
        db, await _patient(db, "VH-R1", "全部入り"), owner, _yesterday(), status="cancelled"
    )
    await _checkin(db, visit, sub, "arrival", time(9, 5), match_status="mismatch", source="manual")
    short = await _visit(db, await _patient(db, "VH-R2", "短い滞在"), owner, _yesterday())
    await _checkin(db, short, owner, "arrival", time(9, 0), match_status="no_gps")
    await _checkin(db, short, owner, "departure", time(9, 4))

    res = await client.get(URL, headers=_bearer(admin), params=_range())
    rows = _by_patient(res.json())
    assert rows["全部入り"]["remarks"] == [
        "退出なし",
        "代行（予定: 予定 太郎）",
        "QRなし",
        "場所 要確認",
        "取消済みの予定に記録",
    ]
    assert rows["全部入り"]["checkin_source"] == "manual"
    assert rows["短い滞在"]["remarks"] == ["場所 要確認", "到着と退出が近い"]
    assert rows["短い滞在"]["stay_minutes"] == 4
    await db.rollback()


async def test_unplanned_visit_has_no_planned_columns(client, db) -> None:
    """予定外は打刻時刻が予定欄に入っているだけなので、予定として見せない。"""
    admin = await _admin(db)
    nurse = await _staff(db, "看護 一郎")
    visit = await _visit(
        db,
        await _patient(db, "VH-U1", "予定外"),
        nurse,
        _yesterday(),
        start=time(14, 3),
        end=time(14, 33),
        status="completed",
        is_unplanned=True,
    )
    await _checkin(db, visit, nurse, "arrival", time(14, 3))
    await _checkin(db, visit, nurse, "departure", time(14, 40))

    res = await client.get(URL, headers=_bearer(admin), params=_range())
    row = res.json()["items"][0]
    assert row["is_unplanned"] is True
    assert row["start_time"] is None
    assert row["end_time"] is None
    assert row["planned_staff_id"] is None
    assert row["planned_staff_name"] is None
    assert row["actual_staff_name"] == "看護 一郎"
    assert row["is_substitute"] is False
    assert row["remarks"] == ["予定外の訪問"]
    await db.rollback()


async def test_substitute_means_outside_the_assigned_set(client, db) -> None:
    """担当集合 (primary / secondary / mentor / assignments / 同行) の外だけが代行。"""
    admin = await _admin(db)
    owner = await _staff(db, "予定 太郎")
    assigned = await _staff(db, "割当 次郎")
    companion = await _staff(db, "同行 三郎")
    outsider = await _staff(db, "代行 花子")
    day = _yesterday()

    by_assigned = await _visit(db, await _patient(db, "VH-A1", "割当が打刻"), owner, day)
    db.add(VisitStaffAssignment(visit_id=by_assigned.id, staff_id=assigned.id))
    by_companion = await _visit(db, await _patient(db, "VH-A2", "同行が打刻"), owner, day)
    db.add(
        Accompaniment(
            accompanying_staff_id=companion.id, target_type="visit", visit_id=by_companion.id
        )
    )
    by_secondary = await _visit(db, await _patient(db, "VH-A3", "副担当が打刻"), owner, day)
    by_secondary.secondary_staff_id = assigned.id
    by_outsider = await _visit(db, await _patient(db, "VH-A4", "担当外が打刻"), owner, day)
    await db.commit()
    await _checkin(db, by_assigned, assigned, "arrival", time(9, 0))
    await _checkin(db, by_companion, companion, "arrival", time(9, 0))
    await _checkin(db, by_secondary, assigned, "arrival", time(9, 0))
    await _checkin(db, by_outsider, outsider, "arrival", time(9, 0))

    res = await client.get(URL, headers=_bearer(admin), params=_range())
    rows = _by_patient(res.json())
    assert {name: row["is_substitute"] for name, row in rows.items()} == {
        "割当が打刻": False,
        "同行が打刻": False,
        "副担当が打刻": False,
        "担当外が打刻": True,
    }
    assert rows["担当外が打刻"]["planned_staff_name"] == "予定 太郎"
    assert rows["担当外が打刻"]["actual_staff_name"] == "代行 花子"
    assert "代行（予定: 予定 太郎）" in rows["担当外が打刻"]["remarks"]
    await db.rollback()


async def test_course_staff_is_the_planned_staff_when_primary_is_empty(client, db) -> None:
    """主担当が空の訪問は、コース担当が「予定の担当」(スマホの ``staff_name`` と同じ規則)。

    コース担当本人の打刻は代行にしない。``manual_staff_override`` の訪問と、コースの
    無い訪問にはフォールバックしない (従来どおり「未割当」)。
    """
    admin = await _admin(db)
    course_staff = await _staff(db, "コース 担当")
    outsider = await _staff(db, "代行 花子")
    office = Office(name="東拠点")
    db.add(office)
    await db.commit()
    await db.refresh(office)
    day = _yesterday()
    iso = day.isocalendar()
    course = Course(
        iso_year=iso.year,
        iso_week=iso.week,
        weekday=day.weekday(),
        code="A",
        course_status="course_fixed",
        office_id=office.id,
        assigned_staff_id=course_staff.id,
    )
    db.add(course)
    await db.commit()
    await db.refresh(course)

    by_course_staff = await _visit(db, await _patient(db, "VH-CF1", "コース担当が打刻"), None, day)
    by_course_staff.course_id = course.id
    by_outsider = await _visit(db, await _patient(db, "VH-CF2", "担当外が打刻"), None, day)
    by_outsider.course_id = course.id
    not_read = await _visit(db, await _patient(db, "VH-CF3", "読み取りなし"), None, day)
    not_read.course_id = course.id
    overridden = await _visit(db, await _patient(db, "VH-CF4", "この訪問だけ外した"), None, day)
    overridden.course_id = course.id
    overridden.manual_staff_override = True
    no_course = await _visit(db, await _patient(db, "VH-CF5", "コースなし"), None, day)
    await db.commit()
    for visit in (by_course_staff, overridden, no_course):
        await _checkin(db, visit, course_staff, "arrival", time(9, 0))
    await _checkin(db, by_outsider, outsider, "arrival", time(9, 0))
    headers = _bearer(admin)

    res = await client.get(URL, headers=headers, params=_range())
    assert res.status_code == 200, res.text
    rows = _by_patient(res.json())
    for name in ("コース担当が打刻", "担当外が打刻", "読み取りなし"):
        assert rows[name]["planned_staff_id"] == str(course_staff.id), name
        assert rows[name]["planned_staff_name"] == "コース 担当", name
    assert rows["コース担当が打刻"]["is_substitute"] is False
    assert rows["コース担当が打刻"]["remarks"] == ["退出なし"]
    assert rows["担当外が打刻"]["is_substitute"] is True
    assert rows["担当外が打刻"]["remarks"] == ["退出なし", "代行（予定: コース 担当）"]
    assert rows["読み取りなし"]["is_substitute"] is False
    # フォールバックしない訪問は従来どおり。
    for name in ("この訪問だけ外した", "コースなし"):
        assert rows[name]["planned_staff_id"] is None, name
        assert rows[name]["planned_staff_name"] is None, name
        assert rows[name]["is_substitute"] is True, name
        assert rows[name]["remarks"] == ["退出なし", "代行（予定: 未割当）"], name

    # スマホ・PC の訪問詳細 (``VisitRead.staff_name``) と同じ名前になる。
    detail = await client.get(f"/api/v1/visits/{by_course_staff.id}", headers=headers)
    assert detail.json()["staff_name"] == rows["コース担当が打刻"]["planned_staff_name"]
    assert detail.json()["primary_staff_id"] is None

    # 絞り込み (予定の担当 または 実際の打刻者) と Excel「看護師別」の予定件数にも入る。
    res = await client.get(
        URL, headers=headers, params={**_range(), "staff_id": str(course_staff.id)}
    )
    # 「担当外が打刻」「読み取りなし」は、予定の担当 (コース担当) としてだけ一致する。
    assert set(_by_patient(res.json())) == set(rows)
    res = await client.get(URL, headers=headers, params={**_range(), "staff_id": str(outsider.id)})
    assert set(_by_patient(res.json())) == {"担当外が打刻"}
    res = await client.get(f"{URL}/export", headers=headers, params=_range())
    wb = load_workbook(BytesIO(res.content))
    assert list(wb["看護師別"].iter_rows(min_row=2, values_only=True)) == [
        ("コース 担当", 3, 3, 0, 3),
        ("代行 花子", 0, 1, 0, 1),
    ]
    await db.rollback()


# ---------------------------------------------------------------------------
# 絞り込み / 並び替え
# ---------------------------------------------------------------------------


async def test_filters_and_sort(client, db) -> None:
    admin = await _admin(db)
    east = Office(name="東拠点")
    west = Office(name="西拠点")
    db.add_all([east, west])
    await db.commit()
    await db.refresh(east)
    await db.refresh(west)
    aoki = await _staff(db, "青木 看護")
    ueda = await _staff(db, "上田 看護")
    day = _yesterday()
    p_sato = await _patient(db, "VH-F1", "佐藤 一", east)
    p_kato = await _patient(db, "VH-F2", "加藤 二", west)
    p_ito = await _patient(db, "VH-F3", "伊藤 三", west)
    v_sato = await _visit(db, p_sato, aoki, day, start=time(9, 0))
    await _checkin(db, v_sato, aoki, "arrival", time(9, 0))
    v_kato = await _visit(db, p_kato, aoki, day, start=time(10, 0), end=time(11, 0))
    await _checkin(db, v_kato, ueda, "arrival", time(10, 0))  # 上田が代行
    await _visit(db, p_ito, ueda, day, start=time(8, 0), end=time(9, 0))
    headers = _bearer(admin)

    async def names(**params: str) -> list[str]:
        res = await client.get(URL, headers=headers, params={**_range(), **params})
        assert res.status_code == 200, res.text
        return [i["patient_name"] for i in res.json()["items"]]

    # 既定は日付 → 予定開始 → 患者名。
    assert await names() == ["伊藤 三", "佐藤 一", "加藤 二"]
    # 看護師名 (実際の打刻者 ?? 予定の担当) → 日付。
    assert await names(sort="staff") == ["伊藤 三", "加藤 二", "佐藤 一"]
    assert await names(sort="patient") == ["伊藤 三", "佐藤 一", "加藤 二"]

    assert await names(patient_id=str(p_kato.id)) == ["加藤 二"]
    assert await names(office_id=str(east.id)) == ["佐藤 一"]
    # staff_id は予定の担当 **または** 実際の打刻者。
    assert await names(staff_id=str(aoki.id)) == ["佐藤 一", "加藤 二"]
    assert await names(staff_id=str(ueda.id)) == ["伊藤 三", "加藤 二"]
    assert await names(state="in") == ["佐藤 一", "加藤 二"]
    assert await names(state="nodep") == ["佐藤 一", "加藤 二"]
    assert await names(state="none") == ["伊藤 三"]
    assert await names(state="special") == ["加藤 二"]
    # q は患者名・予定担当名・打刻者名の部分一致。2 文字未満は無視。
    assert await names(q="加藤") == ["加藤 二"]
    assert await names(q="上田") == ["伊藤 三", "加藤 二"]
    assert await names(q="加") == ["伊藤 三", "佐藤 一", "加藤 二"]

    res = await client.get(URL, headers=headers, params={**_range(), "state": "bogus"})
    assert res.status_code == 422
    row = (await client.get(URL, headers=headers, params=_range())).json()["items"][1]
    assert row["office_id"] == str(east.id)
    assert row["office_name"] == "東拠点"
    await db.rollback()


# ---------------------------------------------------------------------------
# 権限
# ---------------------------------------------------------------------------


async def test_staff_role_sees_only_own_rows(client, db) -> None:
    """staff は自分が担当集合に入る訪問と自分が打刻した訪問だけ。staff_id 指定も効かない。"""
    me_staff = await _staff(db, "自分 看護")
    other = await _staff(db, "他人 看護")
    me = await _staff_user(db, "vh-me@example.com", me_staff)
    day = _yesterday()
    await _visit(db, await _patient(db, "VH-P1", "自分の担当"), me_staff, day)
    helped = await _visit(db, await _patient(db, "VH-P2", "自分が代行"), other, day)
    await _checkin(db, helped, me_staff, "arrival", time(9, 0))
    assigned = await _visit(db, await _patient(db, "VH-P3", "自分も割当"), other, day)
    db.add(VisitStaffAssignment(visit_id=assigned.id, staff_id=me_staff.id))
    await db.commit()
    theirs = await _visit(db, await _patient(db, "VH-P4", "他人の担当"), other, day)
    await _checkin(db, theirs, other, "arrival", time(9, 0))

    mine = {"自分の担当", "自分が代行", "自分も割当"}
    res = await client.get(URL, headers=_bearer(me), params=_range())
    assert res.status_code == 200, res.text
    assert set(_by_patient(res.json())) == mine
    assert res.json()["summary"]["visits"] == 3

    # staff_id に他人を指定しても自分に固定する。
    res = await client.get(URL, headers=_bearer(me), params={**_range(), "staff_id": str(other.id)})
    assert set(_by_patient(res.json())) == mine

    # Excel / A4 も同じ範囲 (他人の患者名が出ない)。
    res = await client.get(f"{URL}/report", headers=_bearer(me), params=_range())
    assert "他人の担当" not in res.text
    res = await client.get(f"{URL}/export", headers=_bearer(me), params=_range())
    wb = load_workbook(BytesIO(res.content))
    assert [r[3] for r in wb["全予定"].iter_rows(min_row=2, values_only=True)] == sorted(mine)
    await db.rollback()


async def test_staff_role_sees_course_fallback_visits(client, db) -> None:
    """主担当が空でコース担当だけの訪問は、スマホの「今日の訪問」と同じく自分の分。

    ``api/v1/visits._course_fallback_condition`` と同じ規則: ``manual_staff_override``
    の訪問と、主担当が入っている訪問にはフォールバックしない。
    """
    me_staff = await _staff(db, "自分 看護")
    other = await _staff(db, "他人 看護")
    me = await _staff_user(db, "vh-course@example.com", me_staff)
    office = Office(name="東拠点")
    db.add(office)
    await db.commit()
    await db.refresh(office)
    day = _yesterday()
    iso = day.isocalendar()
    course = Course(
        iso_year=iso.year,
        iso_week=iso.week,
        weekday=day.weekday(),
        code="A",
        course_status="course_fixed",
        office_id=office.id,
        assigned_staff_id=me_staff.id,
    )
    db.add(course)
    await db.commit()
    await db.refresh(course)

    fallback = await _visit(db, await _patient(db, "VH-C1", "コース担当だけ"), None, day)
    fallback.course_id = course.id
    overridden = await _visit(db, await _patient(db, "VH-C2", "この訪問だけ外した"), None, day)
    overridden.course_id = course.id
    overridden.manual_staff_override = True
    owned = await _visit(db, await _patient(db, "VH-C3", "主担当は他人"), other, day)
    owned.course_id = course.id
    await db.commit()

    res = await client.get(URL, headers=_bearer(me), params=_range())
    assert res.status_code == 200, res.text
    assert set(_by_patient(res.json())) == {"コース担当だけ"}
    await db.rollback()


async def test_staff_role_without_staff_link_gets_nothing(client, db) -> None:
    nurse = await _staff(db, "看護 一郎")
    unlinked = await _staff_user(db, "vh-unlinked@example.com", None)
    await _visit(db, await _patient(db, "VH-N1", "誰かの訪問"), nurse, _yesterday())

    res = await client.get(URL, headers=_bearer(unlinked), params=_range())
    assert res.status_code == 200, res.text
    assert res.json() == {
        "items": [],
        "total": 0,
        "summary": {
            "visits": 0,
            "with_arrival": 0,
            "with_departure": 0,
            "no_departure": 0,
            "none": 0,
            "adjusted": 0,
        },
        "groups": [],
    }
    await db.rollback()


async def test_requires_authentication(client) -> None:
    for path in (URL, f"{URL}/export", f"{URL}/report"):
        res = await client.get(path, params=_range())
        assert res.status_code == 401, path


# ---------------------------------------------------------------------------
# 期間の検証 / ページング
# ---------------------------------------------------------------------------


async def test_period_validation(client, db) -> None:
    admin = await _admin(db)
    headers = _bearer(admin)
    start = date(2026, 7, 1)

    ok = await client.get(URL, headers=headers, params=_range(start, start + timedelta(days=91)))
    assert ok.status_code == 200, ok.text  # 両端を含めてちょうど 92 日
    for path in (URL, f"{URL}/export", f"{URL}/report"):
        too_long = await client.get(
            path, headers=headers, params=_range(start, start + timedelta(days=92))
        )
        assert too_long.status_code == 422, path
    reversed_range = await client.get(
        URL, headers=headers, params=_range(start, start - timedelta(days=1))
    )
    assert reversed_range.status_code == 422
    missing = await client.get(URL, headers=headers, params={"from": "2026-07-01"})
    assert missing.status_code == 422
    await db.rollback()


async def test_summary_is_not_affected_by_paging(client, db) -> None:
    admin = await _admin(db)
    nurse = await _staff(db, "看護 一郎")
    for i in range(3):
        visit = await _visit(
            db,
            await _patient(db, f"VH-G{i}", f"利用者{i}"),
            nurse,
            _yesterday(),
            start=time(9 + i, 0),
            end=time(10 + i, 0),
        )
        await _checkin(db, visit, nurse, "arrival", time(9 + i, 0))
    headers = _bearer(admin)

    full = (await client.get(URL, headers=headers, params=_range())).json()
    page = await client.get(URL, headers=headers, params={**_range(), "limit": 1, "offset": 1})
    assert page.status_code == 200, page.text
    assert page.headers["cache-control"] == "no-store"
    body = page.json()
    assert [i["patient_name"] for i in body["items"]] == ["利用者1"]
    assert body["total"] == 3
    assert body["summary"] == full["summary"]
    assert body["summary"]["with_arrival"] == 3

    too_many = await client.get(URL, headers=headers, params={**_range(), "limit": 201})
    assert too_many.status_code == 422
    await db.rollback()


async def test_groups_count_the_whole_result_not_the_page(client, db) -> None:
    """見出し行の件数は、ページングする前の絞り込み結果全体から数える。"""
    admin = await _admin(db)
    aoki = await _staff(db, "青木 看護")
    ueda = await _staff(db, "上田 看護")
    day = _yesterday()
    p_sato = await _patient(db, "VH-GR1", "佐藤 一")
    p_kato = await _patient(db, "VH-GR2", "加藤 二")
    hit = await _visit(db, p_sato, aoki, day, start=time(9, 0))
    await _checkin(db, hit, aoki, "arrival", time(9, 0))
    await _visit(db, p_sato, aoki, _tomorrow(), start=time(9, 0))  # future も数える
    helped = await _visit(db, p_kato, aoki, day, start=time(10, 0), end=time(11, 0))
    await _checkin(db, helped, ueda, "arrival", time(10, 0))  # 上田が代行 → 上田の行
    await _visit(db, p_kato, None, day, start=time(11, 0), end=time(12, 0))  # 担当なし
    headers = _bearer(admin)

    res = await client.get(
        URL, headers=headers, params={**_range(), "sort": "staff", "limit": 1, "offset": 1}
    )
    assert res.status_code == 200, res.text
    body = res.json()
    assert len(body["items"]) == 1
    # 並びは画面と同じ (看護師名の順・名前なしは末尾)。
    assert body["groups"] == [
        {"label": "上田 看護", "count": 1, "with_arrival": 1},
        {"label": "青木 看護", "count": 2, "with_arrival": 1},
        {"label": "（担当なし）", "count": 1, "with_arrival": 0},
    ]

    res = await client.get(URL, headers=headers, params={**_range(), "sort": "patient"})
    assert res.json()["groups"] == [
        {"label": "佐藤 一", "count": 2, "with_arrival": 1},
        {"label": "加藤 二", "count": 2, "with_arrival": 1},
    ]
    # 絞り込みの後の件数。
    res = await client.get(
        URL, headers=headers, params={**_range(), "sort": "patient", "state": "in"}
    )
    assert res.json()["groups"] == [
        {"label": "佐藤 一", "count": 1, "with_arrival": 1},
        {"label": "加藤 二", "count": 1, "with_arrival": 1},
    ]
    res = await client.get(URL, headers=headers, params={**_range(), "sort": "date"})
    assert res.json()["groups"] == []
    await db.rollback()


# ---------------------------------------------------------------------------
# Excel / A4
# ---------------------------------------------------------------------------


async def _seed_export(db) -> None:
    nurse = await _staff(db, "看護 一郎")
    sub = await _staff(db, "代行 花子")
    hit = await _visit(db, await _patient(db, "VH-X1", "到着あり"), nurse, _yesterday())
    await _checkin(db, hit, sub, "arrival", time(9, 2), second=30)
    await _checkin(db, hit, sub, "departure", time(9, 50))
    await _visit(
        db,
        await _patient(db, "VH-X2", "打刻なし"),
        nurse,
        _yesterday(),
        start=time(11, 0),
        end=time(12, 0),
    )
    await _visit(db, await _patient(db, "VH-X3", "これから"), nurse, _tomorrow())


async def test_export_returns_xlsx_with_four_sheets(client, db) -> None:
    admin = await _admin(db)
    await _seed_export(db)
    params = _range()

    res = await client.get(f"{URL}/export", headers=_bearer(admin), params=params)
    assert res.status_code == 200, res.text
    assert res.headers["cache-control"] == "no-store"
    assert res.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert res.headers["content-disposition"] == (
        f'attachment; filename="visit-history_{params["from"]}_{params["to"]}.xlsx"'
    )

    wb = load_workbook(BytesIO(res.content))
    assert wb.sheetnames == ["QR読み取りあり", "全予定", "看護師別", "読み方"]

    hit_rows = list(wb["QR読み取りあり"].iter_rows(values_only=True))
    assert hit_rows[0] == (
        "日付",
        "曜日",
        "拠点",
        "利用者",
        "予定開始",
        "予定終了",
        "予定の担当",
        "訪問した看護師",
        "到着",
        "退出",
        "滞在(分)",
        "読取時刻（到着）",
        "読取時刻（退出）",
        "備考",
    )
    assert len(hit_rows) == 2
    row = hit_rows[1]
    assert row[0].date() == _yesterday()
    assert row[3] == "到着あり"
    assert row[4:6] == (time(9, 0), time(10, 0))
    assert row[6:8] == ("看護 一郎", "代行 花子")
    # 時刻は JST・分単位 (09:02:30 → 09:02)。
    assert row[8:11] == (time(9, 2), time(9, 50), 48)
    # 読取時刻の 2 列 (調整が無ければ到着・退出と同じ時刻)。
    assert row[11:13] == (time(9, 2), time(9, 50))
    assert row[13] == "代行（予定: 看護 一郎）"

    # future (これから) は載せない。
    all_rows = list(wb["全予定"].iter_rows(values_only=True))
    assert all_rows[0][-1] == "QR 読み取り"
    assert [(r[3], r[-1]) for r in all_rows[1:]] == [("到着あり", "あり"), ("打刻なし", "なし")]

    assert list(wb["看護師別"].iter_rows(values_only=True)) == [
        ("看護師", "予定の件数", "到着の読み取り", "退出の読み取り", "退出なし"),
        ("代行 花子", 0, 1, 1, 0),
        ("看護 一郎", 2, 0, 0, 0),
    ]
    assert wb["読み方"]["A1"].value == (
        f"訪問時刻の記録（QR 読み取り） {_yesterday():%Y/%m/%d}〜{_tomorrow():%Y/%m/%d}"
    )
    assert "2 件の訪問のうち 1 件" in wb["読み方"]["A2"].value
    await db.rollback()


async def test_report_returns_no_store_a4_html(client, db) -> None:
    admin = await _admin(db)
    await _seed_export(db)
    headers = _bearer(admin)

    res = await client.get(f"{URL}/report", headers=headers, params=_range())
    assert res.status_code == 200, res.text
    assert res.headers["cache-control"] == "no-store"
    assert res.headers["content-type"].startswith("text/html")
    html_doc = res.text
    assert html_doc.startswith("<!doctype html>")
    assert "size:A4 portrait" in html_doc
    # JavaScript が動かなくても読める: 明細の表はサーバが返す HTML に入っていて、
    # CSS の印刷フローで改ページできる。用紙へ組み直すスクリプトはその上に載る。
    assert "thead{display:table-header-group}" in html_doc
    assert "tr,.lead,.kpi,.empty{break-inside:avoid}" in html_doc
    flow, script = html_doc.split("<script>")
    assert "<td>到着あり</td>" in flow
    assert 'id="sheets"' in flow
    # スクリプトは定数 (氏名などのデータを埋め込まない)。
    assert "到着あり" not in script
    assert "看護" not in script
    assert "この表の読み方" in html_doc
    assert "看護師別の件数" in html_doc
    assert "2 件の訪問のうち 1 件" in html_doc
    # 既定は看護師別・到着のある訪問だけ・改ページなし。
    assert "<h2>代行 花子" in html_doc
    assert "<td>到着あり</td>" in html_doc
    assert "打刻なし" not in html_doc
    assert "これから" not in html_doc
    assert "<section data-pb" not in html_doc
    # 調整のある行が無ければ、フッターは「読み取った時刻です」の注記。
    footer = flow.split("<footer>")[1].split("</footer>")[0]
    assert footer.startswith("<span>到着・退出は QR を読み取った時刻です（家に入ってから読むため")
    assert "調整" not in footer

    res = await client.get(
        f"{URL}/report",
        headers=headers,
        params={**_range(), "include_none": "true", "page_break": "true"},
    )
    assert "<td>打刻なし</td>" in res.text
    assert "これから" not in res.text  # future は include_none でも載せない
    assert res.text.count('<section data-pb="1">') == 2

    res = await client.get(
        f"{URL}/report", headers=headers, params={**_range(), "group": "patient"}
    )
    assert "<h2>到着あり" in res.text
    res = await client.get(f"{URL}/report", headers=headers, params={**_range(), "group": "date"})
    assert "<h2>明細" in res.text
    res = await client.get(f"{URL}/report", headers=headers, params={**_range(), "group": "x"})
    assert res.status_code == 422
    await db.rollback()


async def test_report_labels_a_whole_month_like_the_monthly_report(client, db) -> None:
    admin = await _admin(db)
    headers = _bearer(admin)
    res = await client.get(
        f"{URL}/report", headers=headers, params=_range(date(2026, 9, 1), date(2026, 9, 30))
    )
    assert "<h1>訪問時刻の記録（QR 読み取り）　2026 年 9 月</h1>" in res.text
    res = await client.get(
        f"{URL}/report", headers=headers, params=_range(date(2026, 9, 1), date(2026, 9, 29))
    )
    assert "<h1>訪問時刻の記録（QR 読み取り）　2026/09/01〜2026/09/29</h1>" in res.text
    await db.rollback()


async def test_report_escapes_names(client, db) -> None:
    admin = await _admin(db)
    nurse = await _staff(db, "看護 <b>一郎</b>")
    visit = await _visit(
        db, await _patient(db, "VH-E1", "<script>alert(1)</script>"), nurse, _yesterday()
    )
    await _checkin(db, visit, nurse, "arrival", time(9, 0))

    res = await client.get(f"{URL}/report", headers=_bearer(admin), params=_range())
    assert "<script>alert" not in res.text
    assert "&lt;script&gt;alert(1)&lt;/script&gt;" in res.text
    assert "看護 &lt;b&gt;一郎&lt;/b&gt;" in res.text
    await db.rollback()


async def test_export_writes_names_as_text_not_formulas(client, db) -> None:
    """``=`` で始まる氏名を数式として書かない (文字列のセルにする)。"""
    admin = await _admin(db)
    nurse = await _staff(db, '=HYPERLINK("http://example.invalid","看護")')
    visit = await _visit(db, await _patient(db, "VH-FX1", "=1+1"), nurse, _yesterday())
    await _checkin(db, visit, nurse, "arrival", time(9, 0))

    res = await client.get(f"{URL}/export", headers=_bearer(admin), params=_range())
    assert res.status_code == 200, res.text
    wb = load_workbook(BytesIO(res.content))
    for sheet in ("QR読み取りあり", "全予定"):
        row = wb[sheet][2]
        # 利用者・予定の担当・訪問した看護師。
        for cell, expected in ((row[3], "=1+1"), (row[6], nurse.name), (row[7], nurse.name)):
            assert cell.data_type == "s", (sheet, cell.coordinate)
            assert cell.value == expected
    name_cell = wb["看護師別"]["A2"]
    assert (name_cell.data_type, name_cell.value) == ("s", nurse.name)
    # 数式のセルは 1 つも無い。
    assert not [c.coordinate for ws in wb for r in ws.iter_rows() for c in r if c.data_type == "f"]
    await db.rollback()


async def test_export_and_report_are_recorded_in_audit_logs(client, db) -> None:
    """氏名と時刻の一括出力 (Excel / A4) は audit_logs に残す。画面の一覧は残さない。"""
    admin = await _admin(db)
    await _seed_export(db)
    params = _range()
    headers = _bearer(admin)

    assert (await client.get(URL, headers=headers, params=params)).status_code == 200
    assert (await client.get(f"{URL}/export", headers=headers, params=params)).status_code == 200
    assert (await client.get(f"{URL}/report", headers=headers, params=params)).status_code == 200

    logs = (
        await db.scalars(
            select(AuditLog).where(AuditLog.target_table == "visit_history").order_by(AuditLog.id)
        )
    ).all()
    assert [(log.action, log.path) for log in logs] == [
        ("export_read", "/api/v1/visit-history/export"),
        ("report_read", "/api/v1/visit-history/report"),
    ]
    for log in logs:
        assert log.actor_user_id == admin.id
        assert (log.role, log.method, log.status_code) == ("admin", "GET", 200)
        assert log.target_id == f"{params['from']}_{params['to']}"
        # future (これから) を除いた 2 行を出力した。
        assert log.after == {
            "from": params["from"],
            "to": params["to"],
            "rows": 2,
            "scope": None,
        }
    await db.rollback()
