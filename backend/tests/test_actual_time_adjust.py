"""「実績の時刻を合わせる」のテスト (打刻履歴 Phase 2).

正典設計書: ``docs/plans/actual-time-adjust-design-2026-09-30.md`` §9 (BE)。

* 読取時刻 (§3): ``device_time`` が妥当なら採用 / 未来・別日・古すぎは ``scanned_at``。
* 調整 (§6-1): 到着・退出の範囲、手入力の退出と status、打刻し直し、戻す、分単位。
* 権限: 本人・担当集合・代行で打刻した本人は可 / 他人は 404 / 8 日前は 403 / admin は可。
* 反映 (§5 / §6-3): ``VisitRead``・訪問モニター・未訪問の検知・打刻履歴・Excel・A4。
* 予定外訪問の ``start_time`` / ``end_time`` の追随。通常の訪問の予定は動かない。
* 同梱 (§6-2): ``adjusted_time`` つきの打刻 / 不正でも打刻は記録される。
* 監査ログ・migration 0088 の往復。

日付は「今日 (JST)」からの相対で組む (昨日 = staff が合わせられる過去の訪問)。
「今日」は実時計ではなく固定の時刻 ``_NOW`` から取り、API 側の時計
(``visits.py`` / ``visit_history.py`` の ``datetime.now``) も同じ時刻に固定する
(``_fixed_clock``)。実時計に依存すると、JST 0 時をまたいで走ったときにテスト側の
「今日」とサーバ側の「今日」がずれる。時刻まで決めたいテストは ``_freeze_now`` で
``visits.py`` の時計を上書きする。

テスト用セッション (``db``) の commit は、**書き込み系の API を呼ぶ前に済ませる**:
監査ミドルウェアが応答の後に (待たずに) ``audit_logs`` へ書くため、共有の SQLite
接続で直後に commit すると "SQL statements in progress" で落ちることがある
(``test_visits.py`` の既知の不安定さと同じ)。
"""

from __future__ import annotations

import importlib.util
import itertools
import sys
from datetime import UTC, date, datetime, time, timedelta
from io import BytesIO
from pathlib import Path
from zoneinfo import ZoneInfo

import pytest
import sqlalchemy as sa
from alembic.config import Config
from alembic.operations import Operations
from alembic.runtime.migration import MigrationContext
from alembic.script import ScriptDirectory
from openpyxl import load_workbook
from sqlalchemy import create_engine, func, inspect, select

from app.api.v1 import visit_history as visit_history_api
from app.api.v1 import visits as visits_api
from app.core.security import create_access_token, hash_password
from app.models import (
    AuditLog,
    Course,
    Office,
    Patient,
    Staff,
    User,
    Visit,
    VisitCheckin,
    VisitStaffAssignment,
    VisitTimeAdjustment,
)
from app.models.accompaniment import Accompaniment
from app.services.checkin import adjust as adjust_service
from app.services.checkin.actuals import load_actuals, resolve_read_time, stay_minutes
from app.services.checkin.adjust import adjust_actual_time
from app.services.checkin.monitor import build_monitor
from app.services.checkin.notify import run_check_missing

JST = ZoneInfo("Asia/Tokyo")
HISTORY_URL = "/api/v1/visit-history"

# 画面に出る文言に使わない語 (PO 決定: 遅れて記録されるのは看護師の誤りではない)。
FORBIDDEN_WORDS = ("直す", "直し", "修正", "補正")

#: テストの「いま」(水曜の夜)。「今日」の訪問の打刻 (13 時台) がすべて過去になる時刻。
_NOW = datetime(2026, 9, 30, 20, 0, tzinfo=JST)


def _frozen_datetime(fixed: datetime) -> type[datetime]:
    """``now()`` だけが ``fixed`` を返す ``datetime`` (``test_qr_open_checkin`` と同じ)。

    呼ぶたびに 1 ミリ秒だけ進める: 調整の ``created_at`` は API の ``now`` なので、完全に
    止めると続けて入れた調整が同時刻になり、「どれが最新か」が決まらなくなる。
    """
    ticks = itertools.count()

    class _FrozenDatetime(datetime):
        @classmethod
        def now(cls, tz=None):  # noqa: ANN001, ANN206 - stdlib シグネチャに合わせる
            moment = fixed + timedelta(milliseconds=next(ticks))
            return moment.astimezone(tz) if tz is not None else moment.replace(tzinfo=None)

    return _FrozenDatetime


@pytest.fixture(autouse=True)
def _fixed_clock(monkeypatch) -> None:
    """API 側の時計を ``_NOW`` に固定する (テスト側の ``_today()`` と同じ日になる)。"""
    frozen = _frozen_datetime(_NOW)
    monkeypatch.setattr(visits_api, "datetime", frozen)
    monkeypatch.setattr(visit_history_api, "datetime", frozen)


def _today() -> date:
    return _NOW.date()


def _yesterday() -> date:
    return _today() - timedelta(days=1)


def _at(day: date, hour: int, minute: int, second: int = 0) -> datetime:
    """``day`` の JST ``hour:minute:second`` を UTC aware で返す。"""
    return datetime.combine(day, time(hour, minute, second), tzinfo=JST).astimezone(UTC)


def _hm(iso: str | None) -> str | None:
    """API の ISO 文字列 → JST の ``HH:MM:SS``。"""
    if iso is None:
        return None
    value = datetime.fromisoformat(iso)
    if value.tzinfo is None:  # SQLite は生の打刻の時刻を naive (= UTC) で返す。
        value = value.replace(tzinfo=UTC)
    return value.astimezone(JST).strftime("%H:%M:%S")


def _bearer(user: User) -> dict[str, str]:
    token = create_access_token(subject=user.id, role=user.role, staff_id=user.staff_id)
    return {"Authorization": f"Bearer {token}"}


def _freeze_now(monkeypatch, fixed: datetime) -> None:
    """``visits.py`` の ``datetime.now`` を ``fixed`` に固定し直す (時刻まで決めたいテスト用)。"""
    monkeypatch.setattr(visits_api, "datetime", _frozen_datetime(fixed))


async def _staff(db, name: str) -> Staff:
    staff = Staff(name=name)
    db.add(staff)
    await db.commit()
    await db.refresh(staff)
    return staff


async def _user(db, email: str, *, role: str = "staff", staff: Staff | None = None) -> User:
    user = User(
        email=email,
        password_hash=hash_password("x"),
        role=role,
        staff_id=staff.id if staff is not None else None,
    )
    db.add(user)
    await db.commit()
    await db.refresh(user)
    return user


async def _nurse(db, name: str, email: str) -> tuple[Staff, User]:
    staff = await _staff(db, name)
    return staff, await _user(db, email, staff=staff)


async def _patient(db, code: str, name: str = "利用者", **kwargs) -> Patient:
    patient = Patient(code=code, name=name, **kwargs)
    db.add(patient)
    await db.commit()
    await db.refresh(patient)
    return patient


async def _visit(
    db,
    patient: Patient,
    staff: Staff | None,
    day: date,
    *,
    start: time = time(13, 0),
    end: time = time(13, 35),
    status: str = "in_progress",
    **kwargs,
) -> Visit:
    visit = Visit(
        patient_id=patient.id,
        primary_staff_id=staff.id if staff is not None else None,
        visit_date=day,
        start_time=start,
        end_time=end,
        type="regular",
        status=status,
        **kwargs,
    )
    db.add(visit)
    await db.commit()
    await db.refresh(visit)
    return visit


async def _checkin(
    db,
    visit: Visit,
    staff: Staff,
    kind: str,
    scanned_at: datetime,
    *,
    device_time: datetime | None = None,
    created_at: datetime | None = None,
    reason: str | None = None,
) -> VisitCheckin:
    """打刻行を直接入れる。``created_at`` は既定で ``scanned_at`` (= 受信した時刻)。"""
    row = VisitCheckin(
        visit_id=visit.id,
        patient_id=visit.patient_id,
        staff_id=staff.id,
        kind=kind,
        scanned_at=scanned_at,
        device_time=device_time,
        match_status="match",
        checkin_source="qr",
        reason=reason,
        threshold_snapshot={"v": 1},
        created_at=created_at or scanned_at,
    )
    db.add(row)
    await db.commit()
    await db.refresh(row)
    return row


async def _put(
    client,
    user: User,
    visit: Visit,
    kind: str,
    hhmm: str,
    *,
    surface: str | None = None,
    reason_code: str | None = None,
):
    headers = _bearer(user)
    if surface is not None:
        headers["X-Client-Surface"] = surface
    return await client.put(
        f"/api/v1/visits/{visit.id}/actual-time",
        headers=headers,
        json={"kind": kind, "time": hhmm, "reason_code": reason_code, "reason_text": None},
    )


async def _delete(client, user: User, visit: Visit, kind: str):
    return await client.delete(
        f"/api/v1/visits/{visit.id}/actual-time", headers=_bearer(user), params={"kind": kind}
    )


async def _adjustments(db, visit: Visit) -> list[VisitTimeAdjustment]:
    return list(
        (
            await db.scalars(
                select(VisitTimeAdjustment)
                .where(VisitTimeAdjustment.visit_id == visit.id)
                .order_by(VisitTimeAdjustment.created_at, VisitTimeAdjustment.id)
                .execution_options(populate_existing=True)
            )
        ).all()
    )


async def _reload_visit(db, visit: Visit) -> Visit:
    return await db.scalar(
        select(Visit).where(Visit.id == visit.id).execution_options(populate_existing=True)
    )


async def _arrived_visit(db, code: str, *, day: date | None = None, **visit_kwargs):
    """担当本人が 13:06:40 に到着を読んだ訪問 (予定 13:00–13:35) を 1 件作る。"""
    day = day or _yesterday()
    staff, user = await _nurse(db, f"看護 {code}", f"{code.lower()}@example.com")
    visit = await _visit(db, await _patient(db, code), staff, day, **visit_kwargs)
    await _checkin(db, visit, staff, "arrival", _at(day, 13, 6, 40))
    return staff, user, visit


# ---------------------------------------------------------------------------
# 読取時刻 (§3)
# ---------------------------------------------------------------------------


def test_read_time_prefers_valid_device_time() -> None:
    day = date(2026, 9, 29)
    scanned = _at(day, 13, 6, 40)
    # 妥当 (少し前) → device_time。
    assert resolve_read_time(_at(day, 13, 6, 25), scanned) == _at(day, 13, 6, 25)
    # 圏外の後送り (同じ日の数時間前) → device_time。
    assert resolve_read_time(_at(day, 9, 30), scanned) == _at(day, 9, 30)
    # 端末の時計が少し進んでいる (120 秒以内) → device_time。
    assert resolve_read_time(scanned + timedelta(seconds=120), scanned) == scanned + timedelta(
        seconds=120
    )
    # device_time 無し → scanned_at。
    assert resolve_read_time(None, scanned) == scanned
    # naive (SQLite が返す形) は UTC とみなす。
    assert resolve_read_time(None, scanned.replace(tzinfo=None)) == scanned


def test_read_time_falls_back_to_scanned_at() -> None:
    day = date(2026, 9, 29)
    scanned = _at(day, 13, 6, 40)
    # 未来 (120 秒を超えて進んでいる) → scanned_at。
    assert resolve_read_time(scanned + timedelta(seconds=121), scanned) == scanned
    # 別の日 (JST) → scanned_at。
    assert resolve_read_time(_at(day - timedelta(days=1), 23, 50), scanned) == scanned
    late = _at(day, 0, 1)
    assert resolve_read_time(_at(day - timedelta(days=1), 23, 59, 30), late) == late
    # 同じ日でも 18 時間より古い → scanned_at。
    night = _at(day, 23, 0)
    assert resolve_read_time(_at(day, 4, 59), night) == night
    assert resolve_read_time(_at(day, 5, 0), night) == _at(day, 5, 0)


async def test_visit_read_uses_read_time_and_keeps_latest_checkin_raw(client, db) -> None:
    """実績時刻は読取時刻 (妥当な device_time)。``latest_checkin`` は生の打刻のまま。"""
    day = _yesterday()
    staff, user = await _nurse(db, "看護 読取", "read@example.com")
    visit = await _visit(db, await _patient(db, "AT-R1"), staff, day)
    await _checkin(db, visit, staff, "arrival", _at(day, 13, 7, 5), device_time=_at(day, 13, 6, 40))

    res = await client.get(f"/api/v1/visits/{visit.id}", headers=_bearer(user))
    assert res.status_code == 200, res.text
    body = res.json()
    assert _hm(body["actual_arrival_at"]) == "13:06:40"
    assert _hm(body["actual_arrival_read_at"]) == "13:06:40"
    assert _hm(body["latest_checkin"]["scanned_at"]) == "13:07:05"
    assert body["actual_arrival_adjusted"] is False
    assert body["actual_departure_at"] is None
    assert body["actual_departure_read_at"] is None
    assert body["actual_departure_adjusted"] is False
    assert body["actual_departure_manual"] is False
    assert body["actual_adjust_allowed"] is True
    await db.rollback()


# ---------------------------------------------------------------------------
# 到着を合わせる (§6-1)
# ---------------------------------------------------------------------------


async def test_adjust_arrival_changes_actual_time_but_not_the_plan(client, db) -> None:
    _, user, visit = await _arrived_visit(db, "AT-A1")

    res = await _put(
        client, user, visit, "arrival", "12:56", surface="mobile", reason_code="intercom_wait"
    )
    assert res.status_code == 200, res.text
    body = res.json()
    assert _hm(body["actual_arrival_at"]) == "12:56:00"
    assert _hm(body["actual_arrival_read_at"]) == "13:06:40"
    assert body["actual_arrival_adjusted"] is True
    assert body["actual_adjust_allowed"] is True
    # 生の打刻はそのまま。
    assert _hm(body["latest_checkin"]["scanned_at"]) == "13:06:40"
    # 通常の訪問の予定には触れない。
    assert (body["visit_date"], body["start_time"], body["end_time"]) == (
        _yesterday().isoformat(),
        "13:00:00",
        "13:35:00",
    )
    assert body["status"] == "in_progress"
    saved = await _reload_visit(db, visit)
    assert (saved.start_time, saved.end_time) == (time(13, 0), time(13, 35))

    rows = await _adjustments(db, visit)
    assert len(rows) == 1
    row = rows[0]
    assert (row.kind, row.source, row.reason_code, row.reason_text) == (
        "arrival",
        "mobile",
        "intercom_wait",
        None,
    )
    assert row.created_by_user_id == user.id
    assert row.created_by_staff_id == user.staff_id
    checkin = await db.scalar(select(VisitCheckin).where(VisitCheckin.visit_id == visit.id))
    assert row.base_checkin_id == checkin.id
    # 一覧 (スマホの今日の訪問と同じ API) にも同じ値が出る。
    listed = await client.get(
        "/api/v1/visits",
        headers=_bearer(user),
        params={"week_start": _yesterday().isoformat(), "week_end": _yesterday().isoformat()},
    )
    item = next(v for v in listed.json() if v["id"] == str(visit.id))
    assert _hm(item["actual_arrival_at"]) == "12:56:00"
    assert item["actual_arrival_adjusted"] is True
    assert item["actual_adjust_allowed"] is True
    await db.rollback()


async def test_source_defaults_to_pc_without_the_surface_header(client, db) -> None:
    _, user, visit = await _arrived_visit(db, "AT-A2")
    assert (await _put(client, user, visit, "arrival", "13:00")).status_code == 200
    assert (
        await _put(client, user, visit, "arrival", "12:58", surface="tablet")
    ).status_code == 200
    assert [r.source for r in await _adjustments(db, visit)] == ["pc", "pc"]
    await db.rollback()


async def test_arrival_range_for_staff_and_admin(client, db) -> None:
    """読取より後は不可・さかのぼれるのは 90 分 (admin は下限なし)。"""
    _, user, visit = await _arrived_visit(db, "AT-A3")
    admin = await _user(db, "at-a3-admin@example.com", role="admin")

    # 読取 13:06:40 → 分に切り捨てた 13:06 までは可、13:07 は不可。
    late = await _put(client, user, visit, "arrival", "13:07")
    assert late.status_code == 422
    assert late.json()["detail"] == "到着は読み取った時刻（13:06）より後にはできません"
    assert (await _put(client, user, visit, "arrival", "13:06")).status_code == 200

    # 90 分前 = 11:36 までは可、11:35 は不可。
    too_early = await _put(client, user, visit, "arrival", "11:35")
    assert too_early.status_code == 422
    assert "90 分前（11:36）" in too_early.json()["detail"]
    assert (await _put(client, user, visit, "arrival", "11:36")).status_code == 200

    # admin は 90 分の下限なし (同じ日の 00:00 以降)。読取より後は admin でも不可。
    res = await _put(client, admin, visit, "arrival", "00:00")
    assert res.status_code == 200, res.text
    assert _hm(res.json()["actual_arrival_at"]) == "00:00:00"
    assert (await _put(client, admin, visit, "arrival", "13:07")).status_code == 422
    await db.rollback()


async def test_arrival_range_reads_checkin_settings(client, db) -> None:
    """さかのぼれる上限は checkin_settings.arrival_max_back_min (mig 0089) から読む。"""
    from app.models.checkin_settings import CheckinSettings

    _, user, visit = await _arrived_visit(db, "AT-A3S")
    db.add(CheckinSettings(is_singleton=True, arrival_max_back_min=30))
    await db.commit()

    # 読取 13:06 → 30 分前 = 12:36 までは可、12:35 は不可。
    too_early = await _put(client, user, visit, "arrival", "12:35")
    assert too_early.status_code == 422
    assert "30 分前（12:36）" in too_early.json()["detail"]
    assert (await _put(client, user, visit, "arrival", "12:36")).status_code == 200
    await db.rollback()


async def test_arrival_must_be_before_the_departure(client, db) -> None:
    day = _yesterday()
    staff, user, visit = await _arrived_visit(db, "AT-A4", status="completed")
    await _checkin(db, visit, staff, "departure", _at(day, 13, 20, 10))
    assert (await _put(client, user, visit, "arrival", "12:50")).status_code == 200
    assert (await _put(client, user, visit, "departure", "13:00")).status_code == 200

    # 退出の実績 13:00 以降にはできない (読取 13:06 より前でも)。
    res = await _put(client, user, visit, "arrival", "13:00")
    assert res.status_code == 422
    assert res.json()["detail"] == "到着は退出（13:00）より前の時刻にしてください"
    assert (await _put(client, user, visit, "arrival", "12:59")).status_code == 200
    await db.rollback()


async def test_time_must_be_hh_mm_and_errors_are_plain_strings(client, db) -> None:
    """分単位。4xx の ``detail`` は画面にそのまま出せる文字列 (配列や dict にしない)。"""
    _, user, visit = await _arrived_visit(db, "AT-A5")
    url = f"/api/v1/visits/{visit.id}/actual-time"
    headers = _bearer(user)

    details: list[str] = []
    for hhmm in ("12:56:30", "25:00", "1256", "", "12:60"):
        res = await _put(client, user, visit, "arrival", hhmm)
        assert res.status_code == 422, hhmm
        details.append(res.json()["detail"])
    for body in (
        {"kind": "arrival"},  # time 無し
        {"kind": "arrival", "time": 1256},  # 文字列でない
        {"time": "12:56"},  # kind 無し
        {"kind": "no_show", "time": "12:56"},
        {"kind": "arrival", "time": "12:56", "reason_code": "forgot"},
        {"kind": "arrival", "time": "12:56", "reason_text": "あ" * 201},
    ):
        res = await client.put(url, headers=headers, json=body)
        assert res.status_code == 422, body
        details.append(res.json()["detail"])
    res = await client.delete(url, headers=headers)  # kind 無し
    assert res.status_code == 422
    details.append(res.json()["detail"])
    res = await client.delete(url, headers=headers, params={"kind": "no_show"})
    assert res.status_code == 422
    details.append(res.json()["detail"])

    for detail in details:
        assert isinstance(detail, str) and detail, detail
        assert not any(word in detail for word in FORBIDDEN_WORDS), detail
    assert await _adjustments(db, visit) == []
    await db.rollback()


async def test_arrival_without_a_read_is_409(client, db) -> None:
    staff, user = await _nurse(db, "看護 未着", "at-a6@example.com")
    visit = await _visit(db, await _patient(db, "AT-A6"), staff, _yesterday(), status="planned")

    res = await _put(client, user, visit, "arrival", "12:56")
    assert res.status_code == 409
    assert res.json()["detail"] == "到着の記録がありません"
    # 読み取りの無い退出 (手入力) も、到着の読み取りが必要。
    res = await _put(client, user, visit, "departure", "13:30")
    assert res.status_code == 409
    assert res.json()["detail"] == "到着の記録がありません"
    await db.rollback()


async def test_deleted_visit_is_409_and_cancelled_with_a_checkin_is_allowed(client, db) -> None:
    admin = await _user(db, "at-a7-admin@example.com", role="admin")
    _, _, deleted = await _arrived_visit(db, "AT-A7", deleted_at=datetime.now(UTC))
    _, _, cancelled = await _arrived_visit(db, "AT-A8", status="cancelled")

    res = await _put(client, admin, deleted, "arrival", "12:56")
    assert res.status_code == 409
    assert isinstance(res.json()["detail"], str)

    # 取消済みでも打刻があれば合わせられる (取込が打刻済みの予定を取り消した実例)。
    res = await _put(client, admin, cancelled, "arrival", "12:56")
    assert res.status_code == 200, res.text
    assert res.json()["status"] == "cancelled"
    await db.rollback()


# ---------------------------------------------------------------------------
# 退出を合わせる・手で入れる
# ---------------------------------------------------------------------------


async def test_departure_with_a_read_has_a_range(client, db) -> None:
    day = _yesterday()
    staff, user, visit = await _arrived_visit(db, "AT-D1", status="completed")
    await _checkin(db, visit, staff, "departure", _at(day, 13, 40, 20))

    # 到着の実績 (13:06) より後。
    res = await _put(client, user, visit, "departure", "13:06")
    assert res.status_code == 422
    assert res.json()["detail"] == "退出は到着（13:06）より後の時刻にしてください"
    # 読取 13:40 + 30 分 = 14:10 まで。
    res = await _put(client, user, visit, "departure", "14:11")
    assert res.status_code == 422
    assert "30 分後（14:10）" in res.json()["detail"]
    res = await _put(client, user, visit, "departure", "14:10")
    assert res.status_code == 200, res.text
    body = res.json()
    assert _hm(body["actual_departure_at"]) == "14:10:00"
    assert _hm(body["actual_departure_read_at"]) == "13:40:20"
    assert body["actual_departure_adjusted"] is True
    assert body["actual_departure_manual"] is False
    assert body["end_time"] == "13:35:00"
    # 到着を合わせると、退出の下限もそれに従う。
    assert (await _put(client, user, visit, "arrival", "12:56")).status_code == 200
    assert (await _put(client, user, visit, "departure", "12:57")).status_code == 200
    await db.rollback()


async def test_manual_departure_completes_the_visit_and_reset_reopens_it(client, db) -> None:
    """読み取りの無い退出を入れると completed、読取時刻に戻すと in_progress。"""
    _, user, visit = await _arrived_visit(db, "AT-D2")

    res = await _put(client, user, visit, "departure", "13:41", reason_code="no_read")
    assert res.status_code == 200, res.text
    body = res.json()
    assert _hm(body["actual_departure_at"]) == "13:41:00"
    assert body["actual_departure_read_at"] is None
    assert body["actual_departure_adjusted"] is True
    assert body["actual_departure_manual"] is True
    assert body["status"] == "completed"
    assert body["end_time"] == "13:35:00"  # 予定は動かない
    row = (await _adjustments(db, visit))[0]
    assert row.base_checkin_id is None

    # 過去の訪問は 23:59 まで入れられる。
    assert (await _put(client, user, visit, "departure", "23:59")).status_code == 200

    res = await _delete(client, user, visit, "departure")
    assert res.status_code == 200, res.text
    body = res.json()
    assert body["actual_departure_at"] is None
    assert body["actual_departure_adjusted"] is False
    assert body["actual_departure_manual"] is False
    assert body["status"] == "in_progress"
    await db.rollback()


async def test_todays_departure_cannot_be_later_than_now(client, db, monkeypatch) -> None:
    day = _today()
    _, user, visit = await _arrived_visit(db, "AT-D3", day=day)
    _freeze_now(monkeypatch, _at(day, 13, 50, 20))

    res = await _put(client, user, visit, "departure", "13:51")
    assert res.status_code == 422
    assert res.json()["detail"] == "退出はいまの時刻より後にはできません"
    res = await _put(client, user, visit, "departure", "13:50")
    assert res.status_code == 200, res.text
    assert res.json()["status"] == "completed"
    await db.rollback()


async def test_rescan_makes_an_older_adjustment_ineffective(client, db, monkeypatch) -> None:
    """打刻し直すと、それより前の調整は効かなくなる (読み直した打刻には改めて合わせられる)。"""
    day = _today()
    staff, user = await _nurse(db, "看護 読み直し", "at-d4@example.com")
    visit = await _visit(db, await _patient(db, "AT-D4"), staff, day, status="planned")
    checkin_url = f"/api/v1/visits/{visit.id}/checkin"

    _freeze_now(monkeypatch, _at(day, 13, 6, 40))
    assert (await client.post(checkin_url, headers=_bearer(user), json={})).status_code == 200
    _freeze_now(monkeypatch, _at(day, 13, 8))
    res = await _put(client, user, visit, "arrival", "12:56")
    assert res.status_code == 200, res.text
    assert _hm(res.json()["actual_arrival_at"]) == "12:56:00"

    # 調整の後に読み直す → 新しい読取時刻が実績になり、前の調整は効かない。
    _freeze_now(monkeypatch, _at(day, 13, 15))
    res = await client.post(checkin_url, headers=_bearer(user), json={})
    assert res.status_code == 200, res.text
    body = res.json()
    assert _hm(body["actual_arrival_at"]) == "13:15:00"
    assert _hm(body["actual_arrival_read_at"]) == "13:15:00"
    assert body["actual_arrival_adjusted"] is False
    # 記録は消えない (追記専用)。
    assert len(await _adjustments(db, visit)) == 1

    # 読み直した打刻に対して、もう一度合わせられる。
    _freeze_now(monkeypatch, _at(day, 13, 16))
    res = await _put(client, user, visit, "arrival", "13:10")
    assert res.status_code == 200, res.text
    assert _hm(res.json()["actual_arrival_at"]) == "13:10:00"
    assert res.json()["actual_arrival_adjusted"] is True
    assert len(await _adjustments(db, visit)) == 2
    await db.rollback()


async def test_reset_appends_a_null_row_and_is_idempotent(client, db) -> None:
    _, user, visit = await _arrived_visit(db, "AT-D5")

    # 調整が 1 件も無い状態で戻しても、エラーにせず今の VisitRead を返す (行を足さない)。
    res = await _delete(client, user, visit, "arrival")
    assert res.status_code == 200, res.text
    assert _hm(res.json()["actual_arrival_at"]) == "13:06:40"
    assert res.json()["actual_arrival_adjusted"] is False
    res = await _delete(client, user, visit, "departure")
    assert res.status_code == 200, res.text
    assert await _adjustments(db, visit) == []

    assert (await _put(client, user, visit, "arrival", "12:56")).status_code == 200
    res = await _delete(client, user, visit, "arrival")
    assert res.status_code == 200, res.text
    body = res.json()
    assert _hm(body["actual_arrival_at"]) == "13:06:40"
    assert body["actual_arrival_adjusted"] is False
    rows = await _adjustments(db, visit)
    assert [r.adjusted_at is None for r in rows] == [False, True]
    # もう一度戻しても行は増えない。
    assert (await _delete(client, user, visit, "arrival")).status_code == 200
    assert len(await _adjustments(db, visit)) == 2
    # 打刻の行は 1 件のまま書き換わっていない。
    checkins = (
        await db.scalars(select(VisitCheckin).where(VisitCheckin.visit_id == visit.id))
    ).all()
    assert len(checkins) == 1
    await db.rollback()


# ---------------------------------------------------------------------------
# 権限
# ---------------------------------------------------------------------------


async def test_permissions(client, db) -> None:
    day = _yesterday()
    owner, owner_user, visit = await _arrived_visit(db, "AT-P1")
    assigned, assigned_user = await _nurse(db, "看護 割当", "at-p1-assigned@example.com")
    db.add(VisitStaffAssignment(visit_id=visit.id, staff_id=assigned.id))
    await db.commit()
    substitute, substitute_user = await _nurse(db, "看護 代行", "at-p1-sub@example.com")
    _, stranger_user = await _nurse(db, "看護 他人", "at-p1-stranger@example.com")
    unlinked_user = await _user(db, "at-p1-unlinked@example.com")
    admin = await _user(db, "at-p1-admin@example.com", role="admin")
    # 代行が退出を読んだ (担当集合の外だが、自分が打刻した訪問)。
    await _checkin(db, visit, substitute, "departure", _at(day, 13, 40))

    assert (await _put(client, owner_user, visit, "arrival", "13:00")).status_code == 200
    assert (await _put(client, assigned_user, visit, "arrival", "12:59")).status_code == 200
    assert (await _put(client, substitute_user, visit, "arrival", "12:58")).status_code == 200
    assert (await _put(client, admin, visit, "arrival", "12:57")).status_code == 200
    # 関わっていない staff・staff 未紐付けの staff には存在を見せない。
    for other in (stranger_user, unlinked_user):
        res = await _put(client, other, visit, "arrival", "12:56")
        assert res.status_code == 404
        assert (await _delete(client, other, visit, "arrival")).status_code == 404
    res = await client.put(
        f"/api/v1/visits/{visit.id}/actual-time", json={"kind": "arrival", "time": "12:56"}
    )
    assert res.status_code == 401

    # 代行は QR 無しの GET では見えない (従来どおり) が、合わせることはできている。
    rows = await _adjustments(db, visit)
    assert [r.created_by_staff_id for r in rows] == [owner.id, assigned.id, substitute.id, None]
    await db.rollback()


async def test_staff_can_adjust_only_the_last_seven_days(client, db) -> None:
    admin = await _user(db, "at-p2-admin@example.com", role="admin")
    _, user7, seven = await _arrived_visit(db, "AT-P2", day=_today() - timedelta(days=7))
    _, user8, eight = await _arrived_visit(db, "AT-P3", day=_today() - timedelta(days=8))

    assert (await _put(client, user7, seven, "arrival", "12:56")).status_code == 200
    res = await _put(client, user8, eight, "arrival", "12:56")
    assert res.status_code == 403
    assert isinstance(res.json()["detail"], str)
    assert "7 日前まで" in res.json()["detail"]
    assert (await _delete(client, user8, eight, "arrival")).status_code == 403
    # 画面はこの値でボタンを出し分ける。
    shown = await client.get(f"/api/v1/visits/{eight.id}", headers=_bearer(user8))
    assert shown.json()["actual_adjust_allowed"] is False
    shown = await client.get(f"/api/v1/visits/{seven.id}", headers=_bearer(user7))
    assert shown.json()["actual_adjust_allowed"] is True
    # admin は期間に依らず合わせられる。
    res = await _put(client, admin, eight, "arrival", "12:56")
    assert res.status_code == 200, res.text
    assert res.json()["actual_adjust_allowed"] is True
    await db.rollback()


async def test_course_fallback_staff_can_adjust(client, db) -> None:
    """主担当が空でコース担当だけの訪問 (打刻と同じ可視性) も本人が合わせられる。"""
    day = _yesterday()
    office = Office(name="拠点 AT")
    db.add(office)
    await db.commit()
    await db.refresh(office)
    course_staff, course_user = await _nurse(db, "看護 コース", "at-p4@example.com")
    other = await _staff(db, "看護 打刻")
    iso = day.isocalendar()
    course = Course(
        iso_year=iso.year,
        iso_week=iso.week,
        weekday=day.weekday(),
        code="A",
        course_status="course_fixed",
        office_id=office.id,
        assigned_staff_id=course_staff.id,
    )
    db.add(course)
    await db.commit()
    await db.refresh(course)
    visit = await _visit(db, await _patient(db, "AT-P4"), None, day, course_id=course.id)
    await _checkin(db, visit, other, "arrival", _at(day, 13, 6, 40))

    res = await _put(client, course_user, visit, "arrival", "12:56")
    assert res.status_code == 200, res.text
    await db.rollback()


# ---------------------------------------------------------------------------
# 予定外訪問の追随 (§5 の例外)
# ---------------------------------------------------------------------------


async def test_unplanned_visit_follows_the_adjusted_times(client, db) -> None:
    day = _yesterday()
    staff, user = await _nurse(db, "看護 予定外", "at-u1@example.com")
    visit = await _visit(
        db,
        await _patient(db, "AT-U1"),
        staff,
        day,
        start=time(14, 3),
        end=time(15, 3),
        is_unplanned=True,
    )
    await _checkin(db, visit, staff, "arrival", _at(day, 14, 3, 20))

    # 到着を合わせたら start_time も同じ時刻に。
    res = await _put(client, user, visit, "arrival", "13:55")
    assert res.status_code == 200, res.text
    assert (res.json()["start_time"], res.json()["end_time"]) == ("13:55:00", "15:03:00")
    # 退出を手で入れたら end_time も同じ時刻に。
    res = await _put(client, user, visit, "departure", "14:40")
    assert res.status_code == 200, res.text
    assert (res.json()["start_time"], res.json()["end_time"]) == ("13:55:00", "14:40:00")
    assert res.json()["status"] == "completed"
    # 到着を読取時刻に戻したら start_time も読取時刻 (分) に戻る。
    res = await _delete(client, user, visit, "arrival")
    assert res.status_code == 200, res.text
    assert res.json()["start_time"] == "14:03:00"
    assert res.json()["visit_date"] == day.isoformat()
    await db.rollback()


# ---------------------------------------------------------------------------
# 打刻リクエストへの同梱 (§6-2)
# ---------------------------------------------------------------------------


async def test_checkin_with_adjusted_time_creates_one_adjustment(client, db, monkeypatch) -> None:
    day = _today()
    staff, user = await _nurse(db, "看護 同梱", "at-b1@example.com")
    visit = await _visit(db, await _patient(db, "AT-B1"), staff, day, status="planned")
    _freeze_now(monkeypatch, _at(day, 13, 6, 50))

    res = await client.post(
        f"/api/v1/visits/{visit.id}/checkin",
        headers=_bearer(user),
        json={
            "at": _at(day, 13, 6, 40).isoformat(),
            "adjusted_time": "12:56",
            "adjust_reason_code": "intercom_wait",
        },
    )
    assert res.status_code == 200, res.text
    body = res.json()
    assert _hm(body["actual_arrival_at"]) == "12:56:00"
    # 読取時刻は QR を読んだ瞬間 (at)。scanned_at はサーバが受け取った時刻。
    assert _hm(body["actual_arrival_read_at"]) == "13:06:40"
    assert _hm(body["latest_checkin"]["scanned_at"]) == "13:06:50"
    assert body["actual_arrival_adjusted"] is True
    assert body["actual_adjust_allowed"] is True
    assert (body["status"], body["start_time"]) == ("in_progress", "13:00:00")

    rows = await _adjustments(db, visit)
    assert len(rows) == 1
    assert (rows[0].kind, rows[0].source, rows[0].reason_code) == (
        "arrival",
        "checkin",
        "intercom_wait",
    )
    checkin = await db.scalar(select(VisitCheckin).where(VisitCheckin.visit_id == visit.id))
    assert rows[0].base_checkin_id == checkin.id
    await db.rollback()


@pytest.mark.parametrize("adjusted_time", ["13:30", "09:00", "99:99", "12:56:30", "", 1256, None])
async def test_checkin_is_recorded_even_when_adjusted_time_is_invalid(
    client, db, monkeypatch, adjusted_time
) -> None:
    """検証に通らない ``adjusted_time`` は黙って無視し、打刻そのものは必ず記録する。"""
    day = _today()
    staff, user = await _nurse(db, "看護 同梱", "at-b2@example.com")
    visit = await _visit(db, await _patient(db, "AT-B2"), staff, day, status="planned")
    _freeze_now(monkeypatch, _at(day, 13, 6, 50))

    res = await client.post(
        f"/api/v1/visits/{visit.id}/checkin",
        headers=_bearer(user),
        json={"adjusted_time": adjusted_time, "adjust_reason_code": {"bad": True}},
    )
    assert res.status_code == 200, res.text
    body = res.json()
    assert _hm(body["actual_arrival_at"]) == "13:06:50"
    assert body["actual_arrival_adjusted"] is False
    assert body["actual_adjust_allowed"] is True
    assert body["status"] == "in_progress"
    assert (
        await db.scalar(
            select(func.count()).select_from(VisitCheckin).where(VisitCheckin.visit_id == visit.id)
        )
    ) == 1
    assert await _adjustments(db, visit) == []
    await db.rollback()


async def test_checkout_and_adhoc_checkin_accept_adjusted_time(client, db, monkeypatch) -> None:
    day = _today()
    staff, user, visit = await _arrived_visit(db, "AT-B3", day=day)
    await _patient(db, "AT-B4", qr_token="at-b4-token")
    _freeze_now(monkeypatch, _at(day, 13, 45, 10))

    res = await client.post(
        f"/api/v1/visits/{visit.id}/checkout",
        headers=_bearer(user),
        # 理由コードが不正でも、時刻の調整は記録する (理由なしになる)。
        json={"adjusted_time": "13:41", "adjust_reason_code": "unknown"},
    )
    assert res.status_code == 200, res.text
    body = res.json()
    assert _hm(body["actual_departure_at"]) == "13:41:00"
    assert _hm(body["actual_departure_read_at"]) == "13:45:10"
    assert body["actual_departure_manual"] is False
    assert body["status"] == "completed"
    rows = await _adjustments(db, visit)
    assert [(r.kind, r.source, r.reason_code) for r in rows] == [("departure", "checkin", None)]

    # 予定外訪問: 同梱した時刻が通れば、生成される訪問の start_time も追随する。
    res = await client.post(
        "/api/v1/visits/adhoc-checkin",
        headers=_bearer(user),
        json={"qr_token": "at-b4-token", "adjusted_time": "13:30"},
    )
    assert res.status_code == 200, res.text
    body = res.json()
    assert body["is_unplanned"] is True
    assert body["start_time"] == "13:30:00"
    assert _hm(body["actual_arrival_at"]) == "13:30:00"
    assert _hm(body["actual_arrival_read_at"]) == "13:45:10"
    assert body["actual_adjust_allowed"] is True
    await db.rollback()


# ---------------------------------------------------------------------------
# 監査ログ
# ---------------------------------------------------------------------------


async def test_audit_log_keeps_before_and_after(client, db) -> None:
    day = _yesterday()
    _, user, visit = await _arrived_visit(db, "AT-L1")
    assert (
        await _put(client, user, visit, "arrival", "12:56", reason_code="intercom_wait")
    ).status_code == 200
    assert (await _delete(client, user, visit, "arrival")).status_code == 200

    logs = (
        await db.scalars(
            select(AuditLog)
            .where(AuditLog.target_table == "visit_time_adjustments")
            .order_by(AuditLog.id)
        )
    ).all()
    assert [log.action for log in logs] == ["visit_actual_time_adjust", "visit_actual_time_reset"]
    adjust, reset = logs
    assert adjust.actor_user_id == user.id
    assert adjust.target_id == str(visit.id)
    read_iso = _at(day, 13, 6, 40).isoformat()
    adjusted_iso = _at(day, 12, 56).isoformat()
    assert adjust.before == {
        "kind": "arrival",
        "at": read_iso,
        "read_at": read_iso,
        "adjusted": False,
        "reason_code": None,
    }
    assert adjust.after == {
        "kind": "arrival",
        "at": adjusted_iso,
        "reason_code": "intercom_wait",
        "reason_text": None,
        "source": "pc",
    }
    assert reset.before == {
        "kind": "arrival",
        "at": adjusted_iso,
        "read_at": read_iso,
        "adjusted": True,
        "reason_code": "intercom_wait",
    }
    assert reset.after == {"kind": "arrival", "at": read_iso, "source": "pc"}
    await db.rollback()


async def test_audit_log_records_who_and_from_where(client, db, monkeypatch) -> None:
    """調整の監査行にも role / method / path / IP / 端末が入る (Excel・A4 の監査行と同じ項目)。"""
    day = _today()
    staff, user = await _nurse(db, "看護 監査", "at-l2@example.com")
    visit = await _visit(db, await _patient(db, "AT-L2"), staff, day, status="planned")
    url = f"/api/v1/visits/{visit.id}/actual-time"
    headers = {
        **_bearer(user),
        "X-Forwarded-For": "203.0.113.7, 10.0.0.1",
        "User-Agent": "rakusuke-test",
    }
    _freeze_now(monkeypatch, _at(day, 13, 6, 50))

    # 打刻に同梱した調整 → 合わせる → 戻す。
    res = await client.post(
        f"/api/v1/visits/{visit.id}/checkin", headers=headers, json={"adjusted_time": "13:00"}
    )
    assert res.status_code == 200, res.text
    res = await client.put(url, headers=headers, json={"kind": "arrival", "time": "12:56"})
    assert res.status_code == 200, res.text
    res = await client.delete(url, headers=headers, params={"kind": "arrival"})
    assert res.status_code == 200, res.text

    logs = (
        await db.scalars(
            select(AuditLog)
            .where(AuditLog.target_table == "visit_time_adjustments")
            .order_by(AuditLog.id)
        )
    ).all()
    assert [(log.action, log.method, log.path) for log in logs] == [
        ("visit_actual_time_adjust", "POST", f"/api/v1/visits/{visit.id}/checkin"),
        ("visit_actual_time_adjust", "PUT", url),
        ("visit_actual_time_reset", "DELETE", url),
    ]
    for log in logs:
        assert (log.role, log.status_code) == ("staff", 200)
        assert log.ip_address == "203.0.113.7"
        assert log.user_agent == "rakusuke-test"
        assert log.actor_user_id == user.id
    await db.rollback()


# ---------------------------------------------------------------------------
# 訪問モニター / 未訪問の検知 (実績時刻基準)
# ---------------------------------------------------------------------------


def _monitor_visit(monitor, visit: Visit):
    return next(v for row in monitor.staff for v in row.visits if v.visit_id == visit.id)


async def test_monitor_uses_actual_times(client, db) -> None:
    day = _yesterday()
    staff, user = await _nurse(db, "看護 モニタ", "at-m1@example.com")
    admin = await _user(db, "at-m1-admin@example.com", role="admin")
    visit = await _visit(db, await _patient(db, "AT-M1"), staff, day, status="completed")
    await _checkin(db, visit, staff, "arrival", _at(day, 13, 20))
    await _checkin(db, visit, staff, "departure", _at(day, 13, 40))
    now = _at(day, 14, 0)

    before = _monitor_visit(await build_monitor(db, day, now=now), visit)
    assert (before.phase, before.alert_level) == ("done", "review")  # 20 分の遅れ
    assert (before.arrival_delay_min, before.stay_minutes) == (20, 20)
    assert before.arrival_at == _at(day, 13, 20)
    assert before.arrival_read_at == _at(day, 13, 20)
    assert (before.arrival_adjusted, before.departure_adjusted, before.departure_manual) == (
        False,
        False,
        False,
    )
    assert before.adjustments == []

    res = await _put(client, user, visit, "arrival", "13:00", reason_code="intercom_wait")
    assert res.status_code == 200, res.text

    after = _monitor_visit(await build_monitor(db, day, now=now), visit)
    assert (after.phase, after.alert_level) == ("done", "none")
    assert (after.arrival_delay_min, after.stay_minutes) == (0, 40)
    assert after.arrival_at == _at(day, 13, 0)
    assert after.arrival_read_at == _at(day, 13, 20)
    assert after.departure_at == _at(day, 13, 40)
    assert after.arrival_adjusted is True
    # 生の打刻は書き換わらない。
    assert after.arrival.scanned_at.replace(tzinfo=UTC) == _at(day, 13, 20)
    assert len(after.adjustments) == 1
    adjustment = after.adjustments[0]
    assert (
        adjustment.kind,
        adjustment.reason_code,
        adjustment.reason_label,
        adjustment.reason_text,
        adjustment.by_name,
    ) == ("arrival", "intercom_wait", "インターホン待ち", None, "看護 モニタ")
    assert adjustment.created_at is not None

    # API の応答にも同じ項目が載る。
    res = await client.get(
        "/api/v1/monitor", headers=_bearer(admin), params={"date": day.isoformat()}
    )
    assert res.status_code == 200, res.text
    item = res.json()["staff"][0]["visits"][0]
    assert _hm(item["arrival_at"]) == "13:00:00"
    assert _hm(item["arrival_read_at"]) == "13:20:00"
    assert _hm(item["arrival"]["scanned_at"]) == "13:20:00"
    assert item["arrival_adjusted"] is True
    assert item["departure_manual"] is False
    assert item["adjustments"][0]["reason_label"] == "インターホン待ち"
    assert item["adjustments"][0]["reason_code"] == "intercom_wait"
    await db.rollback()


async def test_monitor_manual_departure_makes_the_visit_done(client, db) -> None:
    day = _yesterday()
    staff, user = await _nurse(db, "看護 手入力", "at-m2@example.com")
    visit = await _visit(db, await _patient(db, "AT-M2"), staff, day)
    await _checkin(db, visit, staff, "arrival", _at(day, 13, 6))
    now = _at(day, 18, 0)

    before = _monitor_visit(await build_monitor(db, day, now=now), visit)
    # 退出の読み取りが無いまま 240 分を超えた = 退出の記録が無い可能性。
    assert (before.phase, before.alert_level, before.stay_minutes) == ("inprogress", "review", 294)

    assert (
        await _put(client, user, visit, "departure", "13:41", reason_code="no_read")
    ).status_code == 200
    after = _monitor_visit(await build_monitor(db, day, now=now), visit)
    assert (after.phase, after.alert_level, after.stay_minutes) == ("done", "none", 35)
    assert after.departure is None  # 生の打刻は無い
    assert after.departure_at == _at(day, 13, 41)
    assert after.departure_read_at is None
    assert (after.departure_adjusted, after.departure_manual) == (True, True)
    assert [(a.kind, a.reason_label) for a in after.adjustments] == [("departure", "読み取りなし")]
    await db.rollback()


async def test_pair_start_and_missing_check_use_actual_times(db) -> None:
    """同住所ペアの補正と未訪問の検知が、相方の実績時刻 (手で入れた退出) を起点にする。"""
    day = _yesterday()
    staff, user = await _nurse(db, "看護 ペア", "at-m3@example.com")
    await _user(db, "at-m3-admin@example.com", role="admin")
    first = await _visit(
        db,
        await _patient(db, "AT-M3", "ペア 先", lat=35.6, lng=139.7),
        staff,
        day,
        start=time(9, 0),
        end=time(10, 0),
    )
    second = await _visit(
        db,
        await _patient(db, "AT-M4", "ペア 後", lat=35.6, lng=139.7),
        staff,
        day,
        start=time(9, 0),
        end=time(10, 0),
        status="planned",
    )
    await _checkin(db, first, staff, "arrival", _at(day, 9, 0))
    now = _at(day, 9, 55)

    # 相方の退出が無い間は「到着 + 予定の 60 分」= 10:00 が起点 → まだ待ち。
    before = _monitor_visit(await build_monitor(db, day, now=now), second)
    assert (before.phase, before.pair_waiting) == ("awaiting", True)
    result = await run_check_missing(db, target_date=day, now=now)
    assert (result["scanned"], result["missing"]) == (2, 0)

    # 相方の退出を 09:30 と入れたら、起点は 09:30 → 猶予 20 分を過ぎて未訪問。
    # (この後 run_check_missing が commit するので、API を通さずサービスを直接呼ぶ。)
    await adjust_actual_time(
        db,
        visit=first,
        kind="departure",
        hhmm="09:30",
        reason_code="no_read",
        reason_text=None,
        source="pc",
        actor=user,
        now=datetime.now(UTC),
    )
    await db.commit()
    after = _monitor_visit(await build_monitor(db, day, now=now), second)
    assert (after.phase, after.alert_level) == ("missing", "missing")
    result = await run_check_missing(db, target_date=day, now=now)
    assert (result["missing"], result["created"]) == (1, 1)
    await db.rollback()


# ---------------------------------------------------------------------------
# 打刻履歴・Excel・A4
# ---------------------------------------------------------------------------


async def test_history_rows_show_adjustments(client, db) -> None:
    day = _yesterday()
    admin = await _user(db, "at-h1-admin@example.com", role="admin")
    staff, user = await _nurse(db, "看護 履歴", "at-h1@example.com")
    adjusted = await _visit(
        db, await _patient(db, "AT-H1", "調整あり"), staff, day, status="completed"
    )
    await _checkin(db, adjusted, staff, "arrival", _at(day, 13, 6, 40))
    await _checkin(db, adjusted, staff, "departure", _at(day, 13, 31, 10))
    manual = await _visit(
        db,
        await _patient(db, "AT-H2", "退出を入れた"),
        staff,
        day,
        start=time(15, 0),
        end=time(15, 30),
    )
    await _checkin(db, manual, staff, "arrival", _at(day, 15, 2))
    plain = await _visit(
        db, await _patient(db, "AT-H3", "調整なし"), staff, day, start=time(16, 0), end=time(16, 30)
    )
    await _checkin(db, plain, staff, "arrival", _at(day, 16, 1))
    assert (
        await _put(client, user, adjusted, "arrival", "12:56", reason_code="intercom_wait")
    ).status_code == 200
    assert (
        await _put(client, user, manual, "departure", "15:32", reason_code="no_read")
    ).status_code == 200
    params = {"from": day.isoformat(), "to": day.isoformat()}

    res = await client.get(HISTORY_URL, headers=_bearer(admin), params=params)
    assert res.status_code == 200, res.text
    body = res.json()
    rows = {item["patient_name"]: item for item in body["items"]}
    row = rows["調整あり"]
    assert (_hm(row["arrival_at"]), _hm(row["arrival_read_at"])) == ("12:56:00", "13:06:40")
    assert (_hm(row["departure_at"]), _hm(row["departure_read_at"])) == ("13:31:10", "13:31:10")
    assert (row["arrival_adjusted"], row["departure_adjusted"], row["departure_manual"]) == (
        True,
        False,
        False,
    )
    assert row["stay_minutes"] == 35  # 読取どおりなら 25 分
    assert row["remarks"] == ["時刻調整"]
    assert row["adjust_allowed"] is True
    assert len(row["adjustments"]) == 1
    adjustment = row["adjustments"][0]
    assert set(adjustment) == {
        "kind",
        "reason_code",
        "reason_label",
        "reason_text",
        "by_name",
        "created_at",
    }
    assert (adjustment["kind"], adjustment["reason_code"], adjustment["reason_label"]) == (
        "arrival",
        "intercom_wait",
        "インターホン待ち",
    )
    assert adjustment["by_name"] == "看護 履歴"

    row = rows["退出を入れた"]
    assert row["state"] == "done"
    assert (_hm(row["departure_at"]), row["departure_read_at"]) == ("15:32:00", None)
    assert (row["departure_adjusted"], row["departure_manual"]) == (True, True)
    assert row["stay_minutes"] == 30
    assert row["remarks"] == ["時刻調整"]

    row = rows["調整なし"]
    assert (row["arrival_adjusted"], row["adjustments"], row["remarks"]) == (
        False,
        [],
        ["退出なし"],
    )
    assert body["summary"]["adjusted"] == 2
    assert body["groups"] == []

    # 絞り込み「時刻の調整あり」。
    res = await client.get(
        HISTORY_URL, headers=_bearer(admin), params={**params, "state": "adjusted"}
    )
    assert res.status_code == 200, res.text
    assert {i["patient_name"] for i in res.json()["items"]} == {"調整あり", "退出を入れた"}
    assert res.json()["summary"]["adjusted"] == 2

    # Excel: 到着・退出は実績時刻、読取時刻は別の 2 列、備考は「時刻調整（理由）」。
    res = await client.get(f"{HISTORY_URL}/export", headers=_bearer(admin), params=params)
    assert res.status_code == 200, res.text
    sheet = list(load_workbook(BytesIO(res.content))["QR読み取りあり"].iter_rows(values_only=True))
    header = sheet[0]
    assert header[8:14] == (
        "到着",
        "退出",
        "滞在(分)",
        "読取時刻（到着）",
        "読取時刻（退出）",
        "備考",
    )
    by_name = {line[3]: line for line in sheet[1:]}
    assert by_name["調整あり"][8:14] == (
        time(12, 56),
        time(13, 31),
        35,
        time(13, 6),
        time(13, 31),
        "時刻調整（インターホン待ち）",
    )
    assert by_name["退出を入れた"][8:14] == (
        time(15, 2),
        time(15, 32),
        30,
        time(15, 2),
        None,
        "時刻調整（読み取りなし）",
    )
    assert by_name["調整なし"][11:14] == (time(16, 1), None, "退出なし")

    # A4: 読取時刻の列が無いので、備考に読み取った時刻と理由を出す。
    res = await client.get(f"{HISTORY_URL}/report", headers=_bearer(admin), params=params)
    assert res.status_code == 200, res.text
    assert "調整（読取 13:06・インターホン待ち）" in res.text
    assert "調整（退出は読み取りなし）" in res.text
    # 画面に出る部分 (埋め込みスクリプトのコメントは除く) に使わない語が無いこと。
    visible = res.text.split("<script>")[0]
    assert not any(word in visible for word in FORBIDDEN_WORDS)
    # 調整のある行を含むので、フッターは「読み取った時刻です」だけでは終わらない。
    footer = visible.split("<footer>")[1].split("</footer>")[0]
    assert "備考に「調整」とある行は" in footer
    assert "合わせた時刻" in footer
    # 調整のある行を含まない出力 (調整なしの利用者だけ) は、元の注記のまま。
    plain_only = await client.get(
        f"{HISTORY_URL}/report",
        headers=_bearer(admin),
        params={**params, "patient_id": str(plain.patient_id)},
    )
    plain_footer = plain_only.text.split("<footer>")[1].split("</footer>")[0]
    assert "備考に「調整」とある行は" not in plain_footer
    assert "到着・退出は QR を読み取った時刻です" in plain_footer
    await db.rollback()


async def test_history_outputs_without_reason(client, db) -> None:
    """理由の無い調整 (PO 決定 2026-10-01 で画面から理由を外した後の標準) の出し方。

    Excel の備考は「時刻調整」だけ、A4 は読み取った時刻だけを添える。理由の欄は空。
    """
    day = _yesterday()
    admin = await _user(db, "at-h6-admin@example.com", role="admin")
    staff, user = await _nurse(db, "看護 理由なし", "at-h6@example.com")
    adjusted = await _visit(
        db, await _patient(db, "AT-H6", "理由なし調整"), staff, day, status="completed"
    )
    await _checkin(db, adjusted, staff, "arrival", _at(day, 13, 6, 40))
    await _checkin(db, adjusted, staff, "departure", _at(day, 13, 31, 10))
    manual = await _visit(
        db,
        await _patient(db, "AT-H7", "理由なし退出"),
        staff,
        day,
        start=time(15, 0),
        end=time(15, 30),
    )
    await _checkin(db, manual, staff, "arrival", _at(day, 15, 2))
    assert (await _put(client, user, adjusted, "arrival", "12:56")).status_code == 200
    assert (await _put(client, user, manual, "departure", "15:32")).status_code == 200
    params = {"from": day.isoformat(), "to": day.isoformat()}

    res = await client.get(HISTORY_URL, headers=_bearer(admin), params=params)
    assert res.status_code == 200, res.text
    rows = {item["patient_name"]: item for item in res.json()["items"]}
    for name in ("理由なし調整", "理由なし退出"):
        assert rows[name]["remarks"] == ["時刻調整"]
        adjustment = rows[name]["adjustments"][0]
        assert (
            adjustment["reason_code"],
            adjustment["reason_label"],
            adjustment["reason_text"],
        ) == (None, None, None)

    res = await client.get(f"{HISTORY_URL}/export", headers=_bearer(admin), params=params)
    assert res.status_code == 200, res.text
    sheet = list(load_workbook(BytesIO(res.content))["QR読み取りあり"].iter_rows(values_only=True))
    by_name = {line[3]: line for line in sheet[1:]}
    assert by_name["理由なし調整"][13] == "時刻調整"
    assert by_name["理由なし退出"][13] == "時刻調整"

    res = await client.get(f"{HISTORY_URL}/report", headers=_bearer(admin), params=params)
    assert res.status_code == 200, res.text
    assert "調整（読取 13:06）" in res.text
    assert "調整（退出は読み取りなし）" in res.text
    await db.rollback()


async def test_history_adjust_allowed_follows_the_permission(client, db) -> None:
    admin = await _user(db, "at-h2-admin@example.com", role="admin")
    _, recent_user, _recent = await _arrived_visit(db, "AT-H4", day=_yesterday())
    old_day = _today() - timedelta(days=8)
    staff = await db.get(Staff, recent_user.staff_id)
    old = await _visit(db, await _patient(db, "AT-H5", "8 日前"), staff, old_day)
    await _checkin(db, old, staff, "arrival", _at(old_day, 13, 6))
    params = {"from": old_day.isoformat(), "to": _today().isoformat()}

    res = await client.get(HISTORY_URL, headers=_bearer(recent_user), params=params)
    assert {i["visit_date"]: i["adjust_allowed"] for i in res.json()["items"]} == {
        _yesterday().isoformat(): True,
        old_day.isoformat(): False,
    }
    res = await client.get(HISTORY_URL, headers=_bearer(admin), params=params)
    assert [i["adjust_allowed"] for i in res.json()["items"]] == [True, True]
    await db.rollback()


# ---------------------------------------------------------------------------
# 単一ソース (actuals.py)
# ---------------------------------------------------------------------------


async def test_load_actuals_shape(client, db) -> None:
    day = _yesterday()
    staff, user, visit = await _arrived_visit(db, "AT-S1")
    other = await _staff(db, "看護 先")
    await _checkin(db, visit, other, "arrival", _at(day, 13, 1))
    await _checkin(db, visit, staff, "no_show", _at(day, 12, 0))
    untouched = await _visit(db, await _patient(db, "AT-S2"), staff, day, status="planned")
    assert (await _put(client, user, visit, "arrival", "12:56")).status_code == 200

    actuals = await load_actuals(db, [visit.id, untouched.id])
    assert untouched.id not in actuals
    mine = actuals[visit.id]
    assert mine.arrival.at == _at(day, 12, 56)
    assert mine.arrival.read_at == _at(day, 13, 6, 40)
    assert (mine.arrival.adjusted, mine.arrival.manual) == (True, False)
    assert mine.arrival.checkin.staff_id == staff.id
    assert mine.arrival.adjustment.kind == "arrival"
    assert mine.departure is None
    assert mine.no_show is not None and mine.no_show.kind == "no_show"
    # 到着・退出の全打刻者 (新しい順)。no_show は数えない。
    assert mine.checkin_staff_ids == [staff.id, other.id]
    assert mine.latest_checkin.kind == "arrival"
    assert await load_actuals(db, []) == {}
    await db.rollback()


# ---------------------------------------------------------------------------
# レビュー指摘の反映 (M-1 / M-2 / L-1〜L-5 / L-11 / L-12)
# ---------------------------------------------------------------------------


async def test_staff_who_only_checked_in_gets_the_restricted_response(client, db) -> None:
    """M-1: 担当外で打刻しただけのスタッフには、QR capability の GET と同じ範囲だけを返す。"""
    day = _yesterday()
    owner, owner_user = await _nurse(db, "看護 担当", "at-r1-owner@example.com")
    assigned = await _staff(db, "看護 割当")
    companion = await _staff(db, "看護 同行")
    substitute, substitute_user = await _nurse(db, "看護 代行", "at-r1-sub@example.com")
    admin = await _user(db, "at-r1-admin@example.com", role="admin")
    visit = await _visit(
        db,
        await _patient(db, "AT-R1", "絞り込み 太郎"),
        owner,
        day,
        status="completed",
        note="申し送り: 鍵は植木鉢の下",
        kaipoke_id="KP-R1",
    )
    db.add(VisitStaffAssignment(visit_id=visit.id, staff_id=assigned.id))
    db.add(
        Accompaniment(accompanying_staff_id=companion.id, target_type="visit", visit_id=visit.id)
    )
    await db.commit()
    # 代行が到着を読み、担当本人が退出を読んだ (最新の打刻 = 担当が理由を書いた退出)。
    await _checkin(db, visit, substitute, "arrival", _at(day, 13, 6, 40))
    await _checkin(db, visit, owner, "departure", _at(day, 13, 40), reason="担当が書いた理由")

    def assert_restricted(body: dict) -> None:
        assert body["note"] is None
        assert body["kaipoke_id"] is None
        assert body["staff_assignments"] == []
        assert body["accompaniment"] is None
        assert body["accompaniments"] == []
        assert body["latest_checkin"]["reason"] is None
        # 訪問を遂行するのに要る情報と実績時刻は残る。
        assert body["patient_name"] == "絞り込み 太郎"
        assert body["actual_adjust_allowed"] is True

    # QR なしの GET では見えない (従来どおり)。
    res = await client.get(f"/api/v1/visits/{visit.id}", headers=_bearer(substitute_user))
    assert res.status_code == 404

    # 調整が無い状態の DELETE (何もしない・200) でも全項目は読めない。
    for kind in ("arrival", "departure"):
        res = await _delete(client, substitute_user, visit, kind)
        assert res.status_code == 200, res.text
        assert_restricted(res.json())
    res = await _put(client, substitute_user, visit, "arrival", "12:56")
    assert res.status_code == 200, res.text
    assert_restricted(res.json())
    assert _hm(res.json()["actual_arrival_at"]) == "12:56:00"
    assert res.json()["actual_arrival_adjusted"] is True
    res = await _delete(client, substitute_user, visit, "arrival")
    assert res.status_code == 200, res.text
    assert_restricted(res.json())
    assert _hm(res.json()["actual_arrival_at"]) == "13:06:40"

    # 担当本人と admin には従来どおり全量が返る。
    for full_user in (owner_user, admin):
        res = await _put(client, full_user, visit, "arrival", "13:00")
        assert res.status_code == 200, res.text
        body = res.json()
        assert body["note"] == "申し送り: 鍵は植木鉢の下"
        assert body["kaipoke_id"] == "KP-R1"
        assert [a["staff_id"] for a in body["staff_assignments"]] == [str(assigned.id)]
        assert [a["staff_name"] for a in body["accompaniments"]] == ["看護 同行"]
        assert body["latest_checkin"]["reason"] == "担当が書いた理由"
        res = await _delete(client, full_user, visit, "arrival")
        assert res.status_code == 200, res.text
        assert res.json()["note"] == "申し送り: 鍵は植木鉢の下"
    await db.rollback()


async def test_resend_of_the_same_read_keeps_the_adjustment(client, db, monkeypatch) -> None:
    """M-2: 同じ読み取り (同じ ``at``) の再送では、先に入れた調整は効いたまま。"""
    day = _today()
    staff, user = await _nurse(db, "看護 再送", "at-r2@example.com")
    visit = await _visit(db, await _patient(db, "AT-R2"), staff, day, status="planned")
    checkin_url = f"/api/v1/visits/{visit.id}/checkin"
    read_at = _at(day, 13, 6, 40).isoformat()

    _freeze_now(monkeypatch, _at(day, 13, 6, 50))
    res = await client.post(checkin_url, headers=_bearer(user), json={"at": read_at})
    assert res.status_code == 200, res.text
    _freeze_now(monkeypatch, _at(day, 13, 8))
    assert (await _put(client, user, visit, "arrival", "12:56")).status_code == 200

    # 同じ読み取りがもう一度届く (スマホの再送)。打刻の行は増えるが、実績は調整のまま。
    _freeze_now(monkeypatch, _at(day, 13, 9))
    res = await client.post(checkin_url, headers=_bearer(user), json={"at": read_at})
    assert res.status_code == 200, res.text
    body = res.json()
    assert _hm(body["actual_arrival_at"]) == "12:56:00"
    assert _hm(body["actual_arrival_read_at"]) == "13:06:40"
    assert body["actual_arrival_adjusted"] is True
    assert (
        await db.scalar(
            select(func.count()).select_from(VisitCheckin).where(VisitCheckin.visit_id == visit.id)
        )
    ) == 2
    # 再送の後でも、読取時刻に戻せる・もう一度合わせられる。
    res = await _delete(client, user, visit, "arrival")
    assert res.json()["actual_arrival_adjusted"] is False
    _freeze_now(monkeypatch, _at(day, 13, 10))
    res = await _put(client, user, visit, "arrival", "13:00")
    assert _hm(res.json()["actual_arrival_at"]) == "13:00:00"

    # 本当に読み直した (読み取った瞬間が違う) 場合は、従来どおり調整が効かなくなる。
    _freeze_now(monkeypatch, _at(day, 13, 15))
    res = await client.post(
        checkin_url, headers=_bearer(user), json={"at": _at(day, 13, 14, 50).isoformat()}
    )
    assert res.status_code == 200, res.text
    body = res.json()
    assert _hm(body["actual_arrival_at"]) == "13:14:50"
    assert body["actual_arrival_adjusted"] is False
    await db.rollback()


async def test_arrival_can_go_back_to_the_read_time_in_the_same_minute(client, db) -> None:
    """L-1: 到着と退出を同じ分に読んだ訪問でも、到着を読取時刻に戻せる。"""
    day = _yesterday()
    staff, user = await _nurse(db, "看護 同分", "at-r3@example.com")
    same = await _visit(db, await _patient(db, "AT-R3"), staff, day, status="completed")
    await _checkin(db, same, staff, "arrival", _at(day, 13, 6, 10))
    await _checkin(db, same, staff, "departure", _at(day, 13, 6, 50))
    assert (await _put(client, user, same, "arrival", "13:00")).status_code == 200

    res = await _delete(client, user, same, "arrival")
    assert res.status_code == 200, res.text
    body = res.json()
    assert (_hm(body["actual_arrival_at"]), _hm(body["actual_departure_at"])) == (
        "13:06:10",
        "13:06:50",
    )
    assert body["actual_arrival_adjusted"] is False

    # 退出を合わせてあって、戻した到着が退出より前にならない場合は従来どおり 422。
    other = await _visit(db, await _patient(db, "AT-R4"), staff, day, status="completed")
    await _checkin(db, other, staff, "arrival", _at(day, 13, 6, 10))
    await _checkin(db, other, staff, "departure", _at(day, 13, 20))
    assert (await _put(client, user, other, "arrival", "12:50")).status_code == 200
    assert (await _put(client, user, other, "departure", "13:00")).status_code == 200
    res = await _delete(client, user, other, "arrival")
    assert res.status_code == 422
    detail = res.json()["detail"]
    assert "読み取った時刻（13:06）" in detail and "退出（13:00）" in detail
    assert not any(word in detail for word in FORBIDDEN_WORDS)
    await db.rollback()


async def test_reset_of_a_manual_departure_restores_the_state_before_it(client, db) -> None:
    """L-2: 手で入れた退出を戻すと、手で入れる前の status と予定外訪問の end_time に戻る。"""
    day = _yesterday()
    staff, user = await _nurse(db, "看護 復元", "at-r5@example.com")
    # 元から completed だった訪問 (到着の読み取りだけがある)。
    done = await _visit(db, await _patient(db, "AT-R5"), staff, day, status="completed")
    await _checkin(db, done, staff, "arrival", _at(day, 13, 6))
    # 予定外訪問 (end_time は生成時の暫定値 15:03)。
    unplanned = await _visit(
        db,
        await _patient(db, "AT-R6"),
        staff,
        day,
        start=time(14, 3),
        end=time(15, 3),
        is_unplanned=True,
    )
    await _checkin(db, unplanned, staff, "arrival", _at(day, 14, 3, 20))

    res = await _put(client, user, done, "departure", "13:41", reason_code="no_read")
    assert res.status_code == 200, res.text
    assert res.json()["status"] == "completed"
    res = await _delete(client, user, done, "departure")
    assert res.status_code == 200, res.text
    assert res.json()["actual_departure_at"] is None
    assert res.json()["status"] == "completed"  # in_progress に落とさない
    assert res.json()["end_time"] == "13:35:00"  # 通常の訪問の予定は動かない

    # 予定外訪問: 手で入れる → 続けてもう一度入れる → 戻す。
    res = await _put(client, user, unplanned, "departure", "14:40", reason_code="no_read")
    assert (res.json()["status"], res.json()["end_time"]) == ("completed", "14:40:00")
    res = await _put(client, user, unplanned, "departure", "14:50", reason_code="no_read")
    assert (res.json()["status"], res.json()["end_time"]) == ("completed", "14:50:00")
    res = await _delete(client, user, unplanned, "departure")
    assert res.status_code == 200, res.text
    body = res.json()
    assert body["actual_departure_at"] is None
    assert (body["status"], body["start_time"], body["end_time"]) == (
        "in_progress",
        "14:03:00",
        "15:03:00",
    )

    rows = await _adjustments(db, unplanned)
    assert [(r.prev_visit_status, r.prev_visit_end_time) for r in rows] == [
        ("in_progress", time(15, 3)),
        ("in_progress", time(15, 3)),  # 合わせ直した行は最初の値を引き継ぐ
        (None, None),  # 戻す操作の行
    ]
    rows = await _adjustments(db, done)
    assert (rows[0].prev_visit_status, rows[0].prev_visit_end_time) == ("completed", None)
    await db.rollback()


async def test_adjust_allowed_needs_an_arrival_read(client, db) -> None:
    """L-3: 到着の読み取りが無い訪問は、合わせる対象が無いので false (admin でも)。"""
    admin = await _user(db, "at-r7-admin@example.com", role="admin")
    staff, user = await _nurse(db, "看護 未読", "at-r7@example.com")
    no_read = await _visit(
        db, await _patient(db, "AT-R7", "読み取りなし"), staff, _yesterday(), status="planned"
    )
    departure_only = await _visit(
        db, await _patient(db, "AT-R8", "退出だけ"), staff, _yesterday(), start=time(15, 0)
    )
    await _checkin(db, departure_only, staff, "departure", _at(_yesterday(), 15, 30))
    arrived = await _visit(
        db, await _patient(db, "AT-R9", "到着あり"), staff, _yesterday(), start=time(16, 0)
    )
    await _checkin(db, arrived, staff, "arrival", _at(_yesterday(), 16, 1))
    future = await _visit(
        db, await _patient(db, "AT-R10", "あした"), staff, _today() + timedelta(days=1)
    )
    expected = {"読み取りなし": False, "退出だけ": False, "到着あり": True, "あした": False}
    by_id = {no_read.id: "読み取りなし", departure_only.id: "退出だけ", arrived.id: "到着あり"}
    by_id[future.id] = "あした"

    for viewer in (user, admin):
        for visit_id, name in by_id.items():
            res = await client.get(f"/api/v1/visits/{visit_id}", headers=_bearer(viewer))
            assert res.json()["actual_adjust_allowed"] is expected[name], name
        res = await client.get(
            HISTORY_URL,
            headers=_bearer(viewer),
            params={
                "from": _yesterday().isoformat(),
                "to": (_today() + timedelta(days=1)).isoformat(),
            },
        )
        assert {i["patient_name"]: i["adjust_allowed"] for i in res.json()["items"]} == expected

    # 未来日の訪問への PUT: staff は 403 (「7 日前まで」ではなく、状況に合う文言)。
    res = await _put(client, user, future, "arrival", "12:56")
    assert res.status_code == 403
    assert res.json()["detail"] == "まだ訪問日になっていない訪問の時刻は合わせられません"
    assert (await _delete(client, user, future, "arrival")).status_code == 403
    # admin は期間の制限が無いので、到着の読み取りが無いことを伝える 409。
    res = await _put(client, admin, future, "arrival", "12:56")
    assert res.status_code == 409
    assert res.json()["detail"] == "到着の記録がありません"
    await db.rollback()


def test_stay_minutes_is_floored_to_the_minute() -> None:
    """L-4: 到着・退出をそれぞれ JST の分に切り捨ててからの差 (負は 0)。"""
    day = date(2026, 9, 29)
    assert stay_minutes(_at(day, 13, 0, 50), _at(day, 13, 30, 10)) == 30
    assert stay_minutes(_at(day, 13, 0, 10), _at(day, 13, 29, 50)) == 29
    assert stay_minutes(_at(day, 13, 30), _at(day, 13, 0)) == 0
    assert stay_minutes(None, _at(day, 13, 0)) is None
    assert stay_minutes(_at(day, 13, 0), None) is None
    # 進行中は現在時刻を同じ規則で使う。SQLite が返す naive (= UTC) も同じ結果。
    assert stay_minutes(_at(day, 13, 0, 50), None, now=_at(day, 13, 30, 10)) == 30
    assert stay_minutes(_at(day, 13, 0, 50).replace(tzinfo=None), _at(day, 13, 30, 10)) == 30


async def test_stay_minutes_is_the_same_in_the_monitor_and_the_history(client, db) -> None:
    """L-4: モニターと打刻履歴の滞在分が同じ規則 (画面の HH:MM の引き算) になる。"""
    day = _yesterday()
    admin = await _user(db, "at-r11-admin@example.com", role="admin")
    staff = await _staff(db, "看護 滞在")
    done = await _visit(db, await _patient(db, "AT-R11", "完了"), staff, day, status="completed")
    await _checkin(db, done, staff, "arrival", _at(day, 13, 0, 50))
    await _checkin(db, done, staff, "departure", _at(day, 13, 30, 10))
    # 09:00:40 に到着して退出の読み取りが無い訪問 (退出の記録が無い可能性の判定用)。
    open_visit = await _visit(
        db, await _patient(db, "AT-R12", "進行中"), staff, day, start=time(9, 0), end=time(9, 30)
    )
    await _checkin(db, open_visit, staff, "arrival", _at(day, 9, 0, 40))

    # 13:00:50 → 13:30:10 は秒まで引くと 29 分 20 秒。分に切り捨ててから引くので 30。
    monitor = await build_monitor(db, day, now=_at(day, 14, 0))
    assert _monitor_visit(monitor, done).stay_minutes == 30
    res = await client.get(
        HISTORY_URL, headers=_bearer(admin), params={"from": day.isoformat(), "to": day.isoformat()}
    )
    rows = {item["patient_name"]: item for item in res.json()["items"]}
    assert rows["完了"]["stay_minutes"] == 30

    # 退出の記録が無い可能性 (滞在 > 240 分) の境界も、分に切り捨てた差で決まる。
    at_limit = _monitor_visit(await build_monitor(db, day, now=_at(day, 13, 0, 50)), open_visit)
    assert (at_limit.stay_minutes, at_limit.alert_level) == (240, "none")
    over = _monitor_visit(await build_monitor(db, day, now=_at(day, 13, 1, 10)), open_visit)
    assert (over.stay_minutes, over.alert_level) == (241, "review")
    await db.rollback()


async def test_read_on_another_day_is_409(client, db) -> None:
    """L-5: 読み取った日と訪問日が違う訪問 (取込で日付が動いた) は合わせられない。"""
    day = _yesterday()
    read_day = day - timedelta(days=1)
    admin = await _user(db, "at-r13-admin@example.com", role="admin")
    staff, user = await _nurse(db, "看護 日付", "at-r13@example.com")
    visit = await _visit(db, await _patient(db, "AT-R13"), staff, day)
    await _checkin(db, visit, staff, "arrival", _at(read_day, 13, 6, 40))

    for actor in (user, admin):
        for kind, hhmm in (("arrival", "12:56"), ("departure", "13:41")):
            res = await _put(client, actor, visit, kind, hhmm)
            assert res.status_code == 409, (kind, res.text)
            detail = res.json()["detail"]
            assert isinstance(detail, str)
            assert detail == (
                f"読み取った日（{read_day.month}/{read_day.day}）と訪問日"
                f"（{day.month}/{day.day}）が違うため、時刻を合わせられません。"
                "管理者に連絡してください"
            )
            assert not any(word in detail for word in FORBIDDEN_WORDS)
    assert await _adjustments(db, visit) == []
    saved = await _reload_visit(db, visit)
    assert (saved.visit_date, saved.start_time, saved.end_time) == (day, time(13, 0), time(13, 35))
    await db.rollback()


async def test_adjuster_without_a_staff_name_is_shown_as_admin_to_staff(client, db) -> None:
    """L-11: スタッフ未紐付けの管理者が合わせた調整は、staff には「管理者」と出す。"""
    day = _yesterday()
    admin = await _user(db, "at-r14-admin@example.com", role="admin")
    _, user, visit = await _arrived_visit(db, "AT-R14")
    assert (await _put(client, admin, visit, "arrival", "12:56")).status_code == 200
    params = {"from": day.isoformat(), "to": day.isoformat()}

    async def by_names(viewer: User) -> tuple[str, str]:
        history = await client.get(HISTORY_URL, headers=_bearer(viewer), params=params)
        assert history.status_code == 200, history.text
        monitor = await client.get(
            "/api/v1/monitor", headers=_bearer(viewer), params={"date": day.isoformat()}
        )
        assert monitor.status_code == 200, monitor.text
        return (
            history.json()["items"][0]["adjustments"][0]["by_name"],
            monitor.json()["staff"][0]["visits"][0]["adjustments"][0]["by_name"],
        )

    assert await by_names(user) == ("管理者", "管理者")
    # admin への応答は今のまま (email)。
    assert await by_names(admin) == ("at-r14-admin@example.com", "at-r14-admin@example.com")
    # スタッフが合わせた調整は、誰が見てもスタッフ名。
    assert (await _put(client, user, visit, "arrival", "12:58")).status_code == 200
    assert await by_names(user) == ("看護 AT-R14", "看護 AT-R14")
    assert await by_names(admin) == ("看護 AT-R14", "看護 AT-R14")
    await db.rollback()


async def test_checkin_survives_an_unexpected_failure_in_the_bundled_adjustment(
    client, db, monkeypatch
) -> None:
    """L-12: 同梱の調整で検証以外の例外が起きても、打刻は記録される (調整だけ捨てる)。"""
    day = _today()
    staff, user = await _nurse(db, "看護 例外", "at-r15@example.com")
    visit = await _visit(db, await _patient(db, "AT-R15"), staff, day, status="planned")
    _freeze_now(monkeypatch, _at(day, 13, 6, 50))
    calls: list[str] = []

    async def boom(session, **kwargs):  # noqa: ANN001, ANN003, ANN202
        # 調整の行を書いた後で落ちる (SAVEPOINT ごと巻き戻ることを確かめる)。
        session.add(
            VisitTimeAdjustment(visit_id=kwargs["visit"].id, kind="arrival", source="checkin")
        )
        await session.flush()
        calls.append(kwargs["hhmm"])
        raise RuntimeError("boom")

    monkeypatch.setattr(adjust_service, "adjust_actual_time", boom)

    res = await client.post(
        f"/api/v1/visits/{visit.id}/checkin",
        headers=_bearer(user),
        json={"adjusted_time": "12:56", "adjust_reason_code": "intercom_wait"},
    )
    assert res.status_code == 200, res.text
    assert calls == ["12:56"]
    body = res.json()
    assert _hm(body["actual_arrival_at"]) == "13:06:50"
    assert body["actual_arrival_adjusted"] is False
    assert body["status"] == "in_progress"
    assert (
        await db.scalar(
            select(func.count()).select_from(VisitCheckin).where(VisitCheckin.visit_id == visit.id)
        )
    ) == 1
    assert await _adjustments(db, visit) == []
    await db.rollback()


# ---------------------------------------------------------------------------
# migration 0088
# ---------------------------------------------------------------------------

_BACKEND_ROOT = Path(__file__).resolve().parent.parent


def _load_migration_module() -> object:
    path = _BACKEND_ROOT / "alembic" / "versions" / "0088_visit_time_adjustments.py"
    spec = importlib.util.spec_from_file_location("migration_0088", path)
    assert spec is not None and spec.loader is not None
    mod = importlib.util.module_from_spec(spec)
    sys.modules["migration_0088"] = mod
    spec.loader.exec_module(mod)
    return mod


def test_migration_0088_revision_chain() -> None:
    cfg = Config(str(_BACKEND_ROOT / "alembic.ini"))
    cfg.set_main_option("script_location", str(_BACKEND_ROOT / "alembic"))
    script = ScriptDirectory.from_config(cfg)
    rev = script.get_revision("0088_visit_time_adjustments")
    assert rev is not None
    assert rev.down_revision == "0087_visit_recordings_summary_edit"
    assert len(list(script.get_heads())) == 1


def test_migration_0088_sqlite_roundtrip(tmp_path: Path) -> None:
    engine = create_engine(f"sqlite:///{tmp_path / 'migration_0088.db'}")
    migration = _load_migration_module()

    def run(step) -> None:
        with engine.begin() as conn:
            with Operations.context(MigrationContext.configure(conn)):
                step()

    run(migration.upgrade)
    insp = inspect(engine)
    assert {c["name"] for c in insp.get_columns("visit_time_adjustments")} == {
        "id",
        "visit_id",
        "kind",
        "adjusted_at",
        "base_checkin_id",
        "reason_code",
        "reason_text",
        "source",
        "created_by_user_id",
        "created_by_staff_id",
        "prev_visit_status",
        "prev_visit_end_time",
        "created_at",
        "updated_at",
    }
    nullable = {c["name"]: c["nullable"] for c in insp.get_columns("visit_time_adjustments")}
    assert nullable["adjusted_at"] is True  # NULL = 読取時刻に戻す
    assert nullable["base_checkin_id"] is True
    # 手で入れた退出の行だけが持つ (それ以外は NULL)。
    assert (nullable["prev_visit_status"], nullable["prev_visit_end_time"]) == (True, True)
    assert (nullable["visit_id"], nullable["kind"], nullable["source"]) == (False, False, False)
    assert "ix_visit_time_adjustments_visit_kind_created" in {
        ix["name"] for ix in insp.get_indexes("visit_time_adjustments")
    }

    insert = (
        "INSERT INTO visit_time_adjustments (id, visit_id, kind, source) "
        "VALUES ('{id}', 'v1', '{kind}', 'pc')"
    )
    with engine.begin() as conn:
        conn.execute(sa.text(insert.format(id="a1", kind="arrival")))
        conn.execute(sa.text(insert.format(id="a2", kind="departure")))
        assert conn.execute(
            sa.text("SELECT created_at FROM visit_time_adjustments WHERE id = 'a1'")
        ).scalar()
    with pytest.raises(sa.exc.IntegrityError):
        with engine.begin() as conn:
            conn.execute(sa.text(insert.format(id="a3", kind="no_show")))

    run(migration.downgrade)
    assert "visit_time_adjustments" not in inspect(engine).get_table_names()
    run(migration.upgrade)
    assert "visit_time_adjustments" in inspect(engine).get_table_names()
    engine.dispose()
