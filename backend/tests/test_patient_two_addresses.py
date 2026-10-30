"""患者様の 2 つの住所 (保険証の住所・実際の訪問先) の試験.

正典: ``docs/plans/patient-two-addresses-design-2026-10-22.md`` §4・§6。

1. 保存のきまり 4 通り (純関数 ``normalize_patient_addresses``)。
2. API: 作成・PATCH の部分更新 (保険証の住所だけ / 同じ→別 / 別→同じ / 訪問先だけ)。
3. 権限: 職員への応答では ``insurance_address`` が null (管理者には出る)。
4. スマホの患者追加申請: 入力の住所 = 保険証の住所・訪問先は同じ。
5. Excel: 出力の列・往復 (新しい列あり)・古い形式 (列なし) で保険証の住所を消さない・
   完全置換の古い形式・カルテ A4 の住所欄。
"""

from __future__ import annotations

from datetime import date, time
from io import BytesIO

import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import select

from app.core.security import create_access_token, hash_password
from app.models import Patient, Staff, User, Visit
from app.models.pending_request import PendingRequest
from app.services.patient_address import PatientAddresses, normalize_patient_addresses
from app.services.patient_excel.importer import excel_address_updates, has_address_columns
from app.services.patient_excel.karte import build_karte_workbook, parse_karte_workbook
from app.services.patient_excel.schema import (
    PATIENT_COL_INDEX,
    PATIENT_COLUMNS,
    SHEET_PATIENTS,
    SHEET_PFV,
)
from app.services.pending_request_applier import PendingRequestApplier

# 訪問先が変わると、座標が無ければサーバーが位置を求める → 試験では偽物 (conftest)。
pytestmark = pytest.mark.usefixtures("fake_visit_geocode")

# ---------------------------------------------------------------------------
# helpers
# ---------------------------------------------------------------------------


async def _make_user(db, email: str, role: str = "admin", staff_id=None) -> User:
    user = User(email=email, password_hash=hash_password("pw"), role=role, staff_id=staff_id)
    db.add(user)
    await db.commit()
    await db.refresh(user)
    return user


def _bearer(user: User) -> dict[str, str]:
    token = create_access_token(subject=user.id, role=user.role, staff_id=user.staff_id)
    return {"Authorization": f"Bearer {token}"}


async def _make_patient(db, **overrides) -> Patient:
    values = {
        "code": "P-ADDR",
        "name": "住所 太郎",
        "sex": "male",
        "status": "active",
        "special_week_active": [],
    }
    values.update(overrides)
    p = Patient(**values)
    db.add(p)
    await db.commit()
    await db.refresh(p)
    return p


async def _reload(db, patient_id) -> Patient:
    # 他の行 (利用者など) は期限切れにしない: その患者だけ読み直す。
    p = await db.scalar(
        select(Patient).where(Patient.id == patient_id).execution_options(populate_existing=True)
    )
    assert p is not None
    return p


def _xlsx(columns: list[dict], rows: list[dict], *, header_overrides: dict | None = None) -> bytes:
    """患者マスタ + 空の固定訪問シートの workbook (列は ``columns`` の並び = 列位置)."""
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_PATIENTS
    headers = header_overrides or {}
    for i, col in enumerate(columns, start=1):
        ws.cell(row=1, column=i, value=headers.get(str(col["key"]), str(col["header"])))
    for r, row in enumerate(rows, start=2):
        for i, col in enumerate(columns, start=1):
            v = row.get(str(col["key"]))
            if v is not None:
                ws.cell(row=r, column=i, value=v)
    wb.create_sheet(title=SHEET_PFV)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


# 2 列を足す前 (NG スタッフ列まで) の古い形式。見出しも古い「住所」。
_LEGACY_COLUMNS = PATIENT_COLUMNS[: PATIENT_COL_INDEX["insurance_address"]]
_LEGACY_HEADERS = {"address": "住所"}


async def _upload(client, admin, content: bytes, *, mode: str = "import", dry_run: bool = False):
    return await client.post(
        f"/api/v1/patients/import-export/{mode}?dry_run={'true' if dry_run else 'false'}",
        headers=_bearer(admin),
        files={
            "file": (
                "t.xlsx",
                content,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )


# ---------------------------------------------------------------------------
# 1. 保存のきまり (純関数)
# ---------------------------------------------------------------------------


def test_rule1_same_copies_insurance_to_visit() -> None:
    r = normalize_patient_addresses(
        insurance_address=" 千葉市稲毛区1-1 ", address="古い住所", visit_same_as_insurance=True
    )
    assert r == PatientAddresses("千葉市稲毛区1-1", "千葉市稲毛区1-1", True)


def test_rule2_different_but_empty_visit_falls_back_to_same() -> None:
    r = normalize_patient_addresses(
        insurance_address="千葉市稲毛区1-1", address="  ", visit_same_as_insurance=False
    )
    assert r == PatientAddresses("千葉市稲毛区1-1", "千葉市稲毛区1-1", True)


def test_rule3_different_with_visit_keeps_visit() -> None:
    r = normalize_patient_addresses(
        insurance_address="千葉市稲毛区1-1",
        address="千葉市中央区2-2",
        visit_same_as_insurance=False,
    )
    assert r == PatientAddresses("千葉市稲毛区1-1", "千葉市中央区2-2", False)


def test_rule4_empty_insurance_is_not_guessed() -> None:
    # 訪問先だけ → 保険証の住所は空のまま・「別」(訪問先を消さない)。
    r = normalize_patient_addresses(
        insurance_address="", address="千葉市中央区2-2", visit_same_as_insurance=True
    )
    assert r == PatientAddresses(None, "千葉市中央区2-2", False)
    # 両方空 → 「同じ」のまま。
    r = normalize_patient_addresses(
        insurance_address=None, address=None, visit_same_as_insurance=False
    )
    assert r == PatientAddresses(None, None, True)


# ---------------------------------------------------------------------------
# 2. API: 作成・PATCH
# ---------------------------------------------------------------------------


@pytest.mark.asyncio
async def test_create_same_uses_insurance_address_and_request_coords(client, db) -> None:
    admin = await _make_user(db, "addr-create@example.com")
    res = await client.post(
        "/api/v1/patients",
        headers=_bearer(admin),
        json={
            "name": "作成 花子",
            "insurance_address": "千葉市稲毛区1-1",
            "visit_same_as_insurance": True,
            "lat": 35.6,
            "lng": 140.1,
        },
    )
    assert res.status_code == 201, res.text
    body = res.json()
    assert body["insurance_address"] == "千葉市稲毛区1-1"
    assert body["address"] == "千葉市稲毛区1-1"
    assert body["visit_same_as_insurance"] is True
    assert (body["lat"], body["lng"]) == (35.6, 140.1)


@pytest.mark.asyncio
async def test_create_different_keeps_visit_address(client, db) -> None:
    admin = await _make_user(db, "addr-create2@example.com")
    res = await client.post(
        "/api/v1/patients",
        headers=_bearer(admin),
        json={
            "name": "作成 次郎",
            "insurance_address": "千葉市稲毛区1-1",
            "visit_same_as_insurance": False,
            "address": "千葉市中央区2-2",
        },
    )
    assert res.status_code == 201, res.text
    body = res.json()
    assert (body["insurance_address"], body["address"], body["visit_same_as_insurance"]) == (
        "千葉市稲毛区1-1",
        "千葉市中央区2-2",
        False,
    )


@pytest.mark.asyncio
async def test_create_visit_only_leaves_insurance_empty(client, db) -> None:
    """古い画面・旧クライアント: address だけ → 保険証の住所は推測で埋めない (きまり 4)."""
    admin = await _make_user(db, "addr-create3@example.com")
    res = await client.post(
        "/api/v1/patients",
        headers=_bearer(admin),
        json={"name": "作成 三郎", "address": "千葉市中央区2-2"},
    )
    assert res.status_code == 201, res.text
    body = res.json()
    assert body["insurance_address"] is None
    assert body["address"] == "千葉市中央区2-2"
    assert body["visit_same_as_insurance"] is False


@pytest.mark.asyncio
async def test_patch_insurance_only_moves_visit_when_same(client, db) -> None:
    admin = await _make_user(db, "addr-patch1@example.com")
    p = await _make_patient(
        db, insurance_address="旧住所", address="旧住所", visit_same_as_insurance=True
    )
    res = await client.patch(
        f"/api/v1/patients/{p.id}",
        headers=_bearer(admin),
        json={"insurance_address": "新住所", "lat": 35.1, "lng": 140.2},
    )
    assert res.status_code == 200, res.text
    body = res.json()
    assert (body["insurance_address"], body["address"], body["visit_same_as_insurance"]) == (
        "新住所",
        "新住所",
        True,
    )
    assert (body["lat"], body["lng"]) == (35.1, 140.2)


@pytest.mark.asyncio
async def test_patch_insurance_only_does_not_move_different_visit(client, db) -> None:
    admin = await _make_user(db, "addr-patch2@example.com")
    p = await _make_patient(
        db,
        code="P-ADDR-2",
        insurance_address="旧住所",
        address="訪問先",
        visit_same_as_insurance=False,
    )
    res = await client.patch(
        f"/api/v1/patients/{p.id}",
        headers=_bearer(admin),
        json={"insurance_address": "新住所"},
    )
    assert res.status_code == 200, res.text
    body = res.json()
    assert (body["insurance_address"], body["address"], body["visit_same_as_insurance"]) == (
        "新住所",
        "訪問先",
        False,
    )


@pytest.mark.asyncio
async def test_patch_same_to_different_and_back(client, db) -> None:
    admin = await _make_user(db, "addr-patch3@example.com")
    p = await _make_patient(
        db,
        code="P-ADDR-3",
        insurance_address="保険証の住所",
        address="保険証の住所",
        visit_same_as_insurance=True,
    )
    # 同じ → 別
    res = await client.patch(
        f"/api/v1/patients/{p.id}",
        headers=_bearer(admin),
        json={"visit_same_as_insurance": False, "address": "別の場所", "lat": 35.5, "lng": 140.5},
    )
    assert res.status_code == 200, res.text
    body = res.json()
    assert (body["insurance_address"], body["address"], body["visit_same_as_insurance"]) == (
        "保険証の住所",
        "別の場所",
        False,
    )
    # 別 → 同じ (訪問先は保険証の住所に戻る)
    res = await client.patch(
        f"/api/v1/patients/{p.id}",
        headers=_bearer(admin),
        json={"visit_same_as_insurance": True},
    )
    assert res.status_code == 200, res.text
    body = res.json()
    assert (body["address"], body["visit_same_as_insurance"]) == ("保険証の住所", True)


@pytest.mark.asyncio
async def test_patch_different_with_empty_visit_falls_back(client, db) -> None:
    admin = await _make_user(db, "addr-patch4@example.com")
    p = await _make_patient(
        db,
        code="P-ADDR-4",
        insurance_address="保険証の住所",
        address="保険証の住所",
        visit_same_as_insurance=True,
    )
    res = await client.patch(
        f"/api/v1/patients/{p.id}",
        headers=_bearer(admin),
        json={"visit_same_as_insurance": False, "address": ""},
    )
    assert res.status_code == 200, res.text
    body = res.json()
    assert (body["address"], body["visit_same_as_insurance"]) == ("保険証の住所", True)


@pytest.mark.asyncio
async def test_patch_explicit_null_same_flag_is_422(client, db) -> None:
    admin = await _make_user(db, "addr-patch5@example.com")
    p = await _make_patient(db, code="P-ADDR-5")
    res = await client.patch(
        f"/api/v1/patients/{p.id}",
        headers=_bearer(admin),
        json={"visit_same_as_insurance": None},
    )
    assert res.status_code == 422, res.text


# ---------------------------------------------------------------------------
# 3. 権限: 職員には保険証の住所を返さない
# ---------------------------------------------------------------------------


@pytest.mark.asyncio
async def test_staff_does_not_get_insurance_address(client, db) -> None:
    staff = Staff(name="職員 A")
    db.add(staff)
    await db.commit()
    await db.refresh(staff)
    staff_user = await _make_user(db, "addr-staff@example.com", "staff", staff_id=staff.id)
    admin = await _make_user(db, "addr-admin@example.com")
    p = await _make_patient(
        db,
        code="P-ADDR-STAFF",
        insurance_address="保険証の住所",
        address="訪問先",
        visit_same_as_insurance=False,
    )
    db.add(
        Visit(
            patient_id=p.id,
            primary_staff_id=staff.id,
            visit_date=date(2026, 10, 22),
            start_time=time(9, 0),
            end_time=time(10, 0),
            type="regular",
            status="planned",
            source="manual",
        )
    )
    await db.commit()

    res = await client.get("/api/v1/patients", headers=_bearer(staff_user))
    assert res.status_code == 200, res.text
    [row] = [r for r in res.json() if r["id"] == str(p.id)]
    assert row["insurance_address"] is None
    assert row["address"] == "訪問先"
    assert row["visit_same_as_insurance"] is False

    res = await client.get(f"/api/v1/patients/{p.id}", headers=_bearer(staff_user))
    assert res.status_code == 200, res.text
    assert res.json()["insurance_address"] is None

    # 管理者には出る。職員の応答で DB の値は消えていない。
    res = await client.get(f"/api/v1/patients/{p.id}", headers=_bearer(admin))
    assert res.json()["insurance_address"] == "保険証の住所"
    assert (await _reload(db, p.id)).insurance_address == "保険証の住所"


# ---------------------------------------------------------------------------
# 4. スマホの患者追加申請
# ---------------------------------------------------------------------------


@pytest.mark.asyncio
async def test_pending_patient_create_uses_input_as_insurance_address(db) -> None:
    user = await _make_user(db, "addr-pending@example.com")
    pr = PendingRequest(
        requester_user_id=user.id,
        request_type="patient_create",
        payload={
            "code": "P-ADDR-PR",
            "name": "申請 太郎",
            "status": "active",
            "address": " 千葉市稲毛区3-3 ",
            "lat": 35.6,
            "lng": 140.1,
        },
        status="pending",
    )
    db.add(pr)
    await db.commit()
    await db.refresh(pr)

    await PendingRequestApplier().apply(db, pr)
    await db.commit()

    row = await db.scalar(select(Patient).where(Patient.code == "P-ADDR-PR"))
    assert row is not None
    assert row.insurance_address == "千葉市稲毛区3-3"
    assert row.address == "千葉市稲毛区3-3"
    assert row.visit_same_as_insurance is True


# ---------------------------------------------------------------------------
# 5. Excel
# ---------------------------------------------------------------------------


def test_excel_columns_appended_at_end_and_address_header_renamed() -> None:
    keys = [str(c["key"]) for c in PATIENT_COLUMNS]
    # 2026-10-30: その後ろに起点の週と NG 曜日 7 列を足した (列の位置は動かさない)。
    i = keys.index("insurance_address")
    assert keys[i : i + 2] == ["insurance_address", "visit_same_as_insurance"]
    assert keys[i + 2] == "biweekly_anchor"
    assert keys.index("ng_staff_codes") == i - 1  # 既存の列位置は動かさない
    headers = {str(c["key"]): c["header"] for c in PATIENT_COLUMNS}
    assert headers["address"] == "訪問先の住所"
    assert headers["insurance_address"] == "保険証の住所"
    assert has_address_columns(tuple(c["header"] for c in PATIENT_COLUMNS))
    assert not has_address_columns(tuple(c["header"] for c in _LEGACY_COLUMNS))


def test_excel_address_updates_keeps_legacy_inconsistent_rows_stable() -> None:
    """きまりを当てた結果が今と同じなら書き換えない (往復で noop を保つ)."""
    p = Patient(insurance_address=None, address="訪問先", visit_same_as_insurance=True)
    assert (
        excel_address_updates(
            address="訪問先", insurance_address=None, visit_same_mark=True, current=p
        )
        == {}
    )


@pytest.mark.asyncio
async def test_excel_export_and_roundtrip(client, db) -> None:
    admin = await _make_user(db, "addr-xl-rt@example.com")
    diff = await _make_patient(
        db,
        code="P-XL-DIFF",
        insurance_address="保険証の住所",
        address="訪問先",
        visit_same_as_insurance=False,
        lat=35.5,
        lng=140.5,
    )
    same = await _make_patient(
        db,
        code="P-XL-SAME",
        insurance_address="同じ住所",
        address="同じ住所",
        visit_same_as_insurance=True,
    )
    res = await client.get("/api/v1/patients/import-export/export", headers=_bearer(admin))
    assert res.status_code == 200, res.text
    exported = res.content
    ws = load_workbook(BytesIO(exported))[SHEET_PATIENTS]
    col = {k: i + 1 for k, i in PATIENT_COL_INDEX.items()}
    assert ws.cell(row=1, column=col["address"]).value == "訪問先の住所"
    rows = {ws.cell(row=r, column=col["patient_code"]).value: r for r in range(2, ws.max_row + 1)}
    r_diff, r_same = rows["P-XL-DIFF"], rows["P-XL-SAME"]
    assert ws.cell(row=r_diff, column=col["insurance_address"]).value == "保険証の住所"
    assert ws.cell(row=r_diff, column=col["visit_same_as_insurance"]).value is None
    assert ws.cell(row=r_same, column=col["visit_same_as_insurance"]).value == "○"

    # そのまま取り込む → 変化なし。
    res = await _upload(client, admin, exported, dry_run=True)
    assert res.status_code == 200, res.text
    summary = res.json()["summary"]
    assert summary["patients_update"] == 0 and summary["patients_error"] == 0, res.text

    # 「別」の方に ○ を付ける → 訪問先は保険証の住所になる (訪問先の欄は空 = 書き換えない)。
    # ○ のまま訪問先の欄に別の住所が残っていると行エラー (下の試験)。
    wb = load_workbook(BytesIO(exported))
    ws = wb[SHEET_PATIENTS]
    ws.cell(row=r_diff, column=col["visit_same_as_insurance"], value="○")
    ws.cell(row=r_diff, column=col["address"]).value = None
    # 「同じ」の方の保険証の住所を変える → 訪問先もそろい、位置は求め直す。
    ws.cell(row=r_same, column=col["insurance_address"], value="引っ越し先")
    ws.cell(row=r_same, column=col["address"], value="引っ越し先")
    buf = BytesIO()
    wb.save(buf)
    res = await _upload(client, admin, buf.getvalue(), dry_run=False)
    assert res.status_code == 200, res.text
    assert res.json()["summary"]["patients_error"] == 0, res.text

    same_id = same.id
    d = await _reload(db, diff.id)
    assert (d.insurance_address, d.address, d.visit_same_as_insurance) == (
        "保険証の住所",
        "保険証の住所",
        True,
    )
    s = await _reload(db, same_id)
    assert (s.insurance_address, s.address, s.visit_same_as_insurance) == (
        "引っ越し先",
        "引っ越し先",
        True,
    )
    assert (float(d.lat), float(s.lat)) == (35.6, 35.6)  # 訪問先が変わった 2 名とも求め直す


@pytest.mark.asyncio
async def test_excel_new_format_new_patient_with_insurance_only(client, db) -> None:
    admin = await _make_user(db, "addr-xl-new@example.com")
    content = _xlsx(
        PATIENT_COLUMNS,
        [
            {
                "patient_code": "P-XL-NEW",
                "name": "新規 花子",
                "sex": "女性",
                "status": "稼働",
                "insurance_address": "千葉市稲毛区5-5",
                "visit_same_as_insurance": "○",
            }
        ],
    )
    res = await _upload(client, admin, content, dry_run=False)
    assert res.status_code == 200, res.text
    assert res.json()["summary"]["patients_new"] == 1, res.text
    p = await db.scalar(select(Patient).where(Patient.code == "P-XL-NEW"))
    assert p is not None
    assert (p.insurance_address, p.address, p.visit_same_as_insurance) == (
        "千葉市稲毛区5-5",
        "千葉市稲毛区5-5",
        True,
    )


@pytest.mark.asyncio
async def test_excel_legacy_file_keeps_insurance_address(client, db) -> None:
    """古い形式 (2 列なし・見出し「住所」): 保険証の住所を消さない・訪問先の編集は活かす."""
    admin = await _make_user(db, "addr-xl-legacy@example.com")
    keep = await _make_patient(
        db,
        code="P-XL-KEEP",
        insurance_address="保険証の住所",
        address="保険証の住所",
        visit_same_as_insurance=True,
    )
    moved = await _make_patient(
        db,
        code="P-XL-MOVE",
        insurance_address="保険証の住所2",
        address="保険証の住所2",
        visit_same_as_insurance=True,
    )
    moved_id = moved.id
    content = _xlsx(
        _LEGACY_COLUMNS,
        [
            {"patient_code": "P-XL-KEEP", "address": "保険証の住所", "note": "備考だけ"},
            {"patient_code": "P-XL-MOVE", "address": "新しい訪問先"},
        ],
        header_overrides=_LEGACY_HEADERS,
    )
    res = await _upload(client, admin, content, dry_run=False)
    assert res.status_code == 200, res.text
    assert res.json()["summary"]["patients_error"] == 0, res.text

    k = await _reload(db, keep.id)
    assert (k.insurance_address, k.address, k.visit_same_as_insurance, k.note) == (
        "保険証の住所",
        "保険証の住所",
        True,
        "備考だけ",
    )
    m = await _reload(db, moved_id)
    assert (m.insurance_address, m.address, m.visit_same_as_insurance) == (
        "保険証の住所2",
        "新しい訪問先",
        False,
    )


@pytest.mark.asyncio
async def test_excel_replace_all_legacy_file_keeps_insurance_address(client, db) -> None:
    admin = await _make_user(db, "addr-xl-ra@example.com")
    p = await _make_patient(
        db,
        code="P-XL-RA",
        insurance_address="保険証の住所",
        address="保険証の住所",
        visit_same_as_insurance=True,
    )
    content = _xlsx(
        _LEGACY_COLUMNS,
        [
            {
                "patient_id": str(p.id),
                "patient_code": "P-XL-RA",
                "name": p.name,
                "sex": "男性",
                "status": "稼働",
                "address": "保険証の住所",
            }
        ],
        header_overrides=_LEGACY_HEADERS,
    )
    res = await _upload(client, admin, content, mode="replace-all", dry_run=False)
    assert res.status_code == 200, res.text
    after = await _reload(db, p.id)
    assert (after.insurance_address, after.address, after.visit_same_as_insurance) == (
        "保険証の住所",
        "保険証の住所",
        True,
    )


def test_karte_shows_insurance_address_when_different_and_reads_back_visit_only() -> None:
    p = Patient(
        code="P-KARTE",
        name="カルテ 太郎",
        insurance_address="保険証の住所",
        address="訪問先",
        visit_same_as_insurance=False,
    )
    wb = build_karte_workbook(patient=p, fixed_visits=[], offices=[], course_templates=[])
    ws = wb.active
    assert ws["D6"].value == "訪問先\n保険証：保険証の住所"

    out = load_workbook(BytesIO(parse_karte_workbook(wb)))[SHEET_PATIENTS]
    header = tuple(c.value for c in out[1])
    assert out.cell(row=2, column=PATIENT_COL_INDEX["address"] + 1).value == "訪問先"
    # カルテには「同じ」の欄が無い → 古い形式として取り込む (同じ/別を変えない)。
    assert not has_address_columns(header)

    p.visit_same_as_insurance = True
    p.address = "保険証の住所"
    ws = build_karte_workbook(patient=p, fixed_visits=[], offices=[], course_templates=[]).active
    assert ws["D6"].value == "保険証の住所"


# ---------------------------------------------------------------------------
# 6. 訪問先の位置をサーバーで求める網 (2026-10-22 レビュー)
# ---------------------------------------------------------------------------

_NOT_FOUND = "訪問先の住所から位置が見つかりませんでした。緯度・経度を入れてください。"


def _geocode_finds_nothing(monkeypatch) -> None:
    async def _none(address: str, **kwargs):
        return None

    monkeypatch.setattr("app.services.geocoding.cached.geocode_address", _none)


@pytest.mark.asyncio
async def test_patch_insurance_only_regeocodes_visit_when_same(
    client, db, fake_visit_geocode
) -> None:
    admin = await _make_user(db, "geo-patch1@example.com")
    p = await _make_patient(
        db,
        code="P-GEO-1",
        insurance_address="旧住所",
        address="旧住所",
        visit_same_as_insurance=True,
        lat=35.0,
        lng=140.0,
    )
    res = await client.patch(
        f"/api/v1/patients/{p.id}", headers=_bearer(admin), json={"insurance_address": "新住所"}
    )
    assert res.status_code == 200, res.text
    body = res.json()
    assert (body["address"], body["lat"], body["lng"]) == ("新住所", 35.6, 140.1)
    assert fake_visit_geocode == ["新住所"]


@pytest.mark.asyncio
async def test_patch_coords_equal_to_saved_are_treated_as_missing(
    client, db, fake_visit_geocode
) -> None:
    """古い座標をそのまま送ってきた (画面が座標を作り直していない) → サーバーが求める."""
    admin = await _make_user(db, "geo-patch1b@example.com")
    p = await _make_patient(
        db,
        code="P-GEO-1B",
        insurance_address="旧住所",
        address="旧住所",
        visit_same_as_insurance=True,
        lat=35.0,
        lng=140.0,
    )
    res = await client.patch(
        f"/api/v1/patients/{p.id}",
        headers=_bearer(admin),
        json={"insurance_address": "新住所", "lat": 35.0, "lng": 140.0},
    )
    assert res.status_code == 200, res.text
    assert (res.json()["lat"], res.json()["lng"]) == (35.6, 140.1)


@pytest.mark.asyncio
async def test_patch_geocode_not_found_is_422(client, db, monkeypatch) -> None:
    _geocode_finds_nothing(monkeypatch)
    admin = await _make_user(db, "geo-patch2@example.com")
    p = await _make_patient(
        db,
        code="P-GEO-2",
        insurance_address="旧住所",
        address="旧住所",
        visit_same_as_insurance=True,
        lat=35.0,
        lng=140.0,
    )
    pid = p.id
    res = await client.patch(
        f"/api/v1/patients/{pid}",
        headers=_bearer(admin),
        json={"insurance_address": "どこにもない"},
    )
    assert res.status_code == 422, res.text
    assert res.json()["detail"] == _NOT_FOUND
    after = await _reload(db, pid)
    assert (after.insurance_address, after.address, float(after.lat)) == ("旧住所", "旧住所", 35.0)


@pytest.mark.asyncio
async def test_patch_same_with_different_address_is_422(client, db) -> None:
    admin = await _make_user(db, "geo-patch3@example.com")
    p = await _make_patient(
        db,
        code="P-GEO-3",
        insurance_address="保険証の住所",
        address="保険証の住所",
        visit_same_as_insurance=True,
    )
    pid = p.id
    res = await client.patch(
        f"/api/v1/patients/{pid}",
        headers=_bearer(admin),
        json={"address": "別の場所", "lat": 35.5, "lng": 140.5},
    )
    assert res.status_code == 422, res.text
    assert "『別の場所』を選んでください" in res.json()["detail"]
    # 同じ文字 (空白の違いだけ) なら通る。
    res = await client.patch(
        f"/api/v1/patients/{pid}", headers=_bearer(admin), json={"address": " 保険証の住所 "}
    )
    assert res.status_code == 200, res.text


@pytest.mark.asyncio
async def test_patch_different_to_same_without_coords_regeocodes(
    client, db, fake_visit_geocode
) -> None:
    admin = await _make_user(db, "geo-patch4@example.com")
    p = await _make_patient(
        db,
        code="P-GEO-4",
        insurance_address="保険証の住所",
        address="別の場所",
        visit_same_as_insurance=False,
        lat=35.0,
        lng=140.0,
    )
    res = await client.patch(
        f"/api/v1/patients/{p.id}", headers=_bearer(admin), json={"visit_same_as_insurance": True}
    )
    assert res.status_code == 200, res.text
    body = res.json()
    assert (body["address"], body["lat"], body["lng"]) == ("保険証の住所", 35.6, 140.1)
    assert fake_visit_geocode == ["保険証の住所"]


@pytest.mark.asyncio
async def test_patch_clearing_both_addresses_clears_coords(client, db) -> None:
    admin = await _make_user(db, "geo-patch5@example.com")
    p = await _make_patient(
        db,
        code="P-GEO-5",
        insurance_address="保険証の住所",
        address="保険証の住所",
        visit_same_as_insurance=True,
        lat=35.0,
        lng=140.0,
    )
    pid = p.id
    # 保険証の住所だけ消す → 訪問先は残す (きまり 4)・位置も残す。
    res = await client.patch(
        f"/api/v1/patients/{pid}", headers=_bearer(admin), json={"insurance_address": None}
    )
    assert res.status_code == 200, res.text
    body = res.json()
    assert (body["address"], body["lat"], body["visit_same_as_insurance"]) == (
        "保険証の住所",
        35.0,
        False,
    )
    # 訪問先も消す → 位置も空にする。
    res = await client.patch(
        f"/api/v1/patients/{pid}", headers=_bearer(admin), json={"address": None}
    )
    assert res.status_code == 200, res.text
    body = res.json()
    assert (body["address"], body["lat"], body["lng"]) == (None, None, None)


@pytest.mark.asyncio
async def test_create_without_coords_geocodes_and_not_found_is_422(
    client, db, monkeypatch, fake_visit_geocode
) -> None:
    admin = await _make_user(db, "geo-create@example.com")
    res = await client.post(
        "/api/v1/patients",
        headers=_bearer(admin),
        json={"name": "座標 なし", "insurance_address": "千葉市稲毛区9-9"},
    )
    assert res.status_code == 201, res.text
    assert (res.json()["lat"], res.json()["lng"]) == (35.6, 140.1)

    _geocode_finds_nothing(monkeypatch)
    res = await client.post(
        "/api/v1/patients",
        headers=_bearer(admin),
        json={"name": "見つからない", "insurance_address": "どこにもない"},
    )
    assert res.status_code == 422, res.text
    assert res.json()["detail"] == _NOT_FOUND


@pytest.mark.asyncio
async def test_pending_create_geocodes_visit_address_not_insurance(db, monkeypatch) -> None:
    from app.services.geocoding.client import GeocodeResult

    calls: list[str] = []

    async def _fake(address: str, **kwargs) -> GeocodeResult:
        calls.append(address)
        return GeocodeResult(
            lat=35.7, lng=140.3, formatted_address=address, place_id="f", status="OK"
        )

    monkeypatch.setattr("app.services.pending_request_applier.geocode_address", _fake)
    user = await _make_user(db, "geo-pending@example.com")
    pr = PendingRequest(
        requester_user_id=user.id,
        request_type="patient_create",
        payload={
            "code": "P-GEO-PR",
            "name": "申請 次郎",
            "status": "active",
            "insurance_address": "保険証の住所",
            "address": "実際の訪問先",
        },
        status="pending",
    )
    db.add(pr)
    await db.commit()
    await db.refresh(pr)
    await PendingRequestApplier().apply(db, pr)
    await db.commit()

    row = await db.scalar(select(Patient).where(Patient.code == "P-GEO-PR"))
    assert row is not None
    assert (row.insurance_address, row.address, row.visit_same_as_insurance) == (
        "保険証の住所",
        "実際の訪問先",
        False,
    )
    assert calls == ["実際の訪問先"]
    assert (float(row.lat), float(row.lng)) == (35.7, 140.3)


@pytest.mark.asyncio
async def test_excel_same_mark_with_different_visit_is_row_error(client, db) -> None:
    admin = await _make_user(db, "geo-xl1@example.com")
    p = await _make_patient(
        db,
        code="P-GEO-XL1",
        insurance_address="保険証の住所",
        address="保険証の住所",
        visit_same_as_insurance=True,
    )
    pid = p.id
    content = _xlsx(
        PATIENT_COLUMNS,
        [
            {
                "patient_code": "P-GEO-XL1",
                "address": "別の場所",
                "insurance_address": "保険証の住所",
                "visit_same_as_insurance": "○",
            }
        ],
    )
    res = await _upload(client, admin, content, dry_run=True)
    assert res.status_code == 200, res.text
    [row] = res.json()["patient_rows"]
    assert row["operation"] == "error"
    assert "○ を外すか、訪問先の住所を保険証の住所に合わせてください" in row["error_message"]
    assert (await _reload(db, pid)).address == "保険証の住所"


@pytest.mark.asyncio
async def test_excel_blank_mark_with_same_text_is_same(client, db) -> None:
    admin = await _make_user(db, "geo-xl2@example.com")
    content = _xlsx(
        PATIENT_COLUMNS,
        [
            {
                "patient_code": "P-GEO-XL2",
                "name": "同じ文字",
                "sex": "女性",
                "status": "稼働",
                "address": "千葉市稲毛区１－１",
                "insurance_address": "千葉市稲毛区1-1",
            }
        ],
    )
    res = await _upload(client, admin, content, dry_run=False)
    assert res.status_code == 200, res.text
    assert res.json()["summary"]["patients_new"] == 1, res.text
    p = await db.scalar(select(Patient).where(Patient.code == "P-GEO-XL2"))
    assert p is not None
    assert (p.address, p.visit_same_as_insurance) == ("千葉市稲毛区1-1", True)
    assert (float(p.lat), float(p.lng)) == (35.6, 140.1)


@pytest.mark.asyncio
async def test_excel_changed_visit_geocodes_and_failure_is_row_error_in_preview(
    client, db, monkeypatch, fake_visit_geocode
) -> None:
    admin = await _make_user(db, "geo-xl3@example.com")
    p = await _make_patient(
        db,
        code="P-GEO-XL3",
        insurance_address="旧住所",
        address="旧住所",
        visit_same_as_insurance=True,
        lat=35.0,
        lng=140.0,
    )
    pid = p.id
    row = {
        "patient_code": "P-GEO-XL3",
        "insurance_address": "新住所",
        "visit_same_as_insurance": "○",
        "lat": 35.0,  # 保存済みと同じ = 新しい座標ではない
        "lng": 140.0,
    }
    # プレビューで、求めた位置が変更として見える。
    res = await _upload(client, admin, _xlsx(PATIENT_COLUMNS, [row]), dry_run=True)
    assert res.status_code == 200, res.text
    [r] = res.json()["patient_rows"]
    assert r["operation"] == "update"
    changes = {c["field"]: c["new_value"] for c in r["changes"]}
    assert changes["address"] == "新住所" and changes["lat"] == 35.6 and changes["lng"] == 140.1

    res = await _upload(client, admin, _xlsx(PATIENT_COLUMNS, [row]), dry_run=False)
    assert res.status_code == 200, res.text
    after = await _reload(db, pid)
    assert (after.address, float(after.lat), float(after.lng)) == ("新住所", 35.6, 140.1)

    # 求められない → 行エラー (プレビューで分かる)・何も変えない。
    _geocode_finds_nothing(monkeypatch)
    row2 = {**row, "insurance_address": "どこにもない", "lat": None, "lng": None}
    res = await _upload(client, admin, _xlsx(PATIENT_COLUMNS, [row2]), dry_run=False)
    assert res.status_code == 200, res.text
    [r] = res.json()["patient_rows"]
    assert r["operation"] == "error" and r["error_message"] == _NOT_FOUND
    assert (await _reload(db, pid)).insurance_address == "新住所"
