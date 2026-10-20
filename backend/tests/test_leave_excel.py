"""休みの段 #7: 職員 Excel (半休・有給の書き出し・往復・全置換の差分) と書き口の grep 試験.

正典: docs/plans/leave-and-kinmu-sheet-design-2026-10-20.md §4-3・§2-1・§11-1。
"""

from __future__ import annotations

import io
import re
from datetime import time, timedelta
from pathlib import Path
from uuid import uuid4

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app.core.security import create_access_token, hash_password
from app.models import Patient, PendingRequest, Staff, StaffWeeklyOverride, User, Visit
from app.services.leave_calendar import iso_key
from app.services.op_log_service import set_staff_override
from app.services.staff_excel.schema import (
    OVERRIDE_EDIT_COL_INDEX,
    OVERRIDE_EDIT_WEEKDAY_KEYS,
    SHEET_OVERRIDE_EDIT,
    format_override_cell,
    parse_override_cell,
)
from app.services.staff_leave import today_jst

APP_DIR = Path(__file__).resolve().parents[1] / "app"


# ---------------------------------------------------------------------------
# セルの表記
# ---------------------------------------------------------------------------


@pytest.mark.parametrize(
    ("override_type", "token"),
    [("off", "休"), ("am_off", "午前休"), ("pm_off", "午後休"), ("paid_off", "有給")],
)
def test_cell_round_trip_leave_types(override_type, token) -> None:
    assert format_override_cell(override_type, None, None) == token
    assert parse_override_cell(token) == (override_type, None, None)


def test_cell_custom_time_and_legacy_values() -> None:
    assert format_override_cell("custom_time", "09:00", "15:00") == "09:00-15:00"
    assert parse_override_cell("9:00-15:00") == ("custom_time", "09:00", "15:00")
    assert parse_override_cell("off") == ("off", None, None)
    assert parse_override_cell("休み") == ("off", None, None)
    with pytest.raises(ValueError):
        parse_override_cell("半休")


# ---------------------------------------------------------------------------
# 往復・全置換
# ---------------------------------------------------------------------------


def _bearer(user: User) -> dict[str, str]:
    token = create_access_token(subject=user.id, role=user.role, staff_id=user.staff_id)
    return {"Authorization": f"Bearer {token}"}


async def _admin(db) -> User:
    u = User(
        email=f"xl-{uuid4().hex[:8]}@example.com", password_hash=hash_password("pw"), role="admin"
    )
    db.add(u)
    await db.commit()
    await db.refresh(u)
    return u


async def _export(client, admin) -> bytes:
    res = await client.get("/api/v1/staff/import-export/export", headers=_bearer(admin))
    assert res.status_code == 200, res.text
    return res.content


async def _replace(client, admin, content: bytes, *, dry_run: bool):
    files = {
        "file": (
            "staff.xlsx",
            content,
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
    }
    return await client.post(
        f"/api/v1/staff/import-export/replace-all?dry_run={'true' if dry_run else 'false'}",
        headers=_bearer(admin),
        files=files,
    )


async def _rows(db, staff_id) -> dict[tuple[int, int, int], StaffWeeklyOverride]:
    rows = (
        await db.scalars(
            select(StaffWeeklyOverride)
            .where(StaffWeeklyOverride.staff_id == staff_id)
            .execution_options(populate_existing=True)
        )
    ).all()
    return {(r.iso_year, r.iso_week, r.weekday): r for r in rows}


@pytest.mark.asyncio
async def test_export_replace_round_trip_keeps_types_and_source(client, db) -> None:
    admin = await _admin(db)
    staff = Staff(code="XL1", name="エクセル 花子")
    db.add(staff)
    await db.flush()
    staff_id = staff.id
    req = PendingRequest(
        requester_user_id=admin.id,
        request_type="staff_off",
        payload={"override_type": "am_off"},
        target_staff_id=staff_id,
        status="approved",
    )
    db.add(req)
    await db.flush()
    req_id = req.id
    base = today_jst() + timedelta(days=14)
    days = [base, base + timedelta(days=1), base + timedelta(days=2)]
    for d, t in zip(days, ("am_off", "paid_off", "custom_time"), strict=True):
        y, w, wd = iso_key(d)
        await set_staff_override(
            db,
            staff_id=staff_id,
            iso_year=y,
            iso_week=w,
            weekday=wd,
            override_type=t,
            start_time="09:00" if t == "custom_time" else None,
            end_time="15:00" if t == "custom_time" else None,
            source="request" if t == "am_off" else None,
            pending_request_id=req_id if t == "am_off" else None,
        )
    await db.commit()

    content = await _export(client, admin)
    wb = load_workbook(io.BytesIO(content))
    ws = wb[SHEET_OVERRIDE_EDIT]
    tokens = {
        ws.cell(row=r, column=OVERRIDE_EDIT_COL_INDEX[k] + 1).value
        for r in range(3, ws.max_row + 1)
        for k in OVERRIDE_EDIT_WEEKDAY_KEYS
    }
    assert {"午前休", "有給", "09:00-15:00"} <= tokens

    res = await _replace(client, admin, content, dry_run=False)
    assert res.status_code == 200, res.text
    assert res.json()["transaction_applied"] is True

    rows = await _rows(db, staff_id)
    assert {r.override_type for r in rows.values()} == {"am_off", "paid_off", "custom_time"}
    am = rows[iso_key(days[0])]
    assert am.source == "request"
    assert am.pending_request_id == req_id


@pytest.mark.asyncio
async def test_replace_with_new_half_leave_unassigns_visit(client, db) -> None:
    admin = await _admin(db)
    staff = Staff(code="XL2", name="エクセル 次郎")
    patient = Patient(code="XL-P1", name="エクセル患者", status="active", special_week_active=[])
    db.add_all([staff, patient])
    await db.flush()
    day = today_jst() + timedelta(days=10)
    visit = Visit(
        patient_id=patient.id,
        primary_staff_id=staff.id,
        visit_date=day,
        start_time=time(14, 0),
        end_time=time(15, 0),
        type="regular",
        status="planned",
        source="auto",
    )
    db.add(visit)
    await db.flush()
    staff_id, staff_code, staff_name, visit_id = staff.id, staff.code, staff.name, visit.id
    await db.commit()

    wb = load_workbook(io.BytesIO(await _export(client, admin)))
    ws = wb[SHEET_OVERRIDE_EDIT]
    row = ws.max_row + 1
    y, w, wd = iso_key(day)
    values = {
        "staff_id": str(staff_id),
        "staff_code": staff_code,
        "staff_name": staff_name,
        "iso_year": y,
        "iso_week": w,
        OVERRIDE_EDIT_WEEKDAY_KEYS[wd]: "午後休",
    }
    for key, value in values.items():
        ws.cell(row=row, column=OVERRIDE_EDIT_COL_INDEX[key] + 1, value=value)
    buf = io.BytesIO()
    wb.save(buf)
    content = buf.getvalue()

    preview = await _replace(client, admin, content, dry_run=True)
    assert preview.status_code == 200, preview.text
    assert preview.json()["summary"]["leave_unassign_visits"] == 1

    res = await _replace(client, admin, content, dry_run=False)
    assert res.status_code == 200, res.text
    got = await db.get(Visit, visit_id, populate_existing=True)
    assert got is not None
    assert got.primary_staff_id is None
    assert got.manual_staff_override is True
    rows = await _rows(db, staff_id)
    assert rows[(y, w, wd)].override_type == "pm_off"
    assert rows[(y, w, wd)].source == "admin"


# ---------------------------------------------------------------------------
# grep 試験 (休み設計 §2-1・段階 0 §0-1)
# ---------------------------------------------------------------------------


def _py_files() -> list[Path]:
    return [p for p in APP_DIR.rglob("*.py") if "__pycache__" not in p.parts]


def test_no_off_equality_outside_core_leave() -> None:
    """``== "off"`` (と ``!=``) を ``app/core/leave.py`` 以外に書かない."""
    pattern = re.compile(r"""[!=]=\s*['"]off['"]""")
    offenders = [
        f"{p.relative_to(APP_DIR)}:{i}"
        for p in _py_files()
        if p.relative_to(APP_DIR).as_posix() != "core/leave.py"
        for i, line in enumerate(p.read_text(encoding="utf-8").splitlines(), 1)
        if pattern.search(line)
    ]
    assert offenders == []


def test_override_rows_written_only_by_set_staff_override() -> None:
    """``StaffWeeklyOverride(`` の生成と ``delete(StaffWeeklyOverride`` は ``set_staff_override``
    の中だけ。``db.delete(`` で override の行を消す箇所も無い (範囲は backend/app)."""
    create = re.compile(r"(?<![A-Za-z0-9_.])StaffWeeklyOverride\(")
    bulk_delete = re.compile(r"delete\(\s*StaffWeeklyOverride\b")
    offenders: list[str] = []
    for p in _py_files():
        rel = p.relative_to(APP_DIR).as_posix()
        text = p.read_text(encoding="utf-8")
        lines = text.splitlines()
        for i, line in enumerate(lines, 1):
            if line.lstrip().startswith(("class ", "#")):
                continue
            if create.search(line) or bulk_delete.search(line):
                offenders.append(f"{rel}:{i}")
        # override の行を db.delete(...) で消している箇所 (変数名に ov / override を含む)。
        for i, line in enumerate(lines, 1):
            if re.search(r"db\.delete\(\s*(ov|override|existing_ov|row_ov)\w*\s*\)", line):
                offenders.append(f"{rel}:{i}")
    allowed_create = {"services/op_log_service.py"}
    remaining = [o for o in offenders if o.split(":")[0] not in allowed_create]
    assert remaining == []
    # op_log_service.py の中でも set_staff_override の本体だけ。
    src = (APP_DIR / "services/op_log_service.py").read_text(encoding="utf-8")
    body = src.split("async def set_staff_override(", 1)[1].split("\nasync def ", 1)[0]
    assert len(create.findall(src)) == len(create.findall(body)) == 1
    assert "await db.delete(existing)" in body
