"""QR 打刻履歴の月次レポート (A4 縦 HTML + Excel) を作る — 読み取り専用の抽出結果から。

画面機能 (/records の「打刻履歴」タブ) が本番に入るまでのつなぎ。入力は
``extract.sql`` の出力 (JSON 配列)。氏名を含むので出力先は gitignore 済みの
``docs/reports/`` にする。

    ssh root@72.60.211.213 'docker exec -i carelink-postgres psql -U carelink -d carelink -At' \
        < docs/tools/visit-history/extract.sql > rows.json
    python docs/tools/visit-history/build_report.py rows.json 2026-09 docs/reports/2026-09-30-qr-history-2026-09
"""

from __future__ import annotations

import datetime as dt
import json
import sys
from collections import OrderedDict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

WD = "月火水木金土日"
ORG = "訪問看護ステーション よりより"
# 到着と退出の間がこれ未満なら「まとめて読んだ」可能性として印を付ける。
SHORT_STAY_MIN = 5


def load(path: str) -> list[dict]:
    raw = Path(path).read_text(encoding="utf-8")
    return json.loads(raw[raw.index("[") :])


def minute_of(ts: str | None) -> int | None:
    """'YYYY-MM-DD HH:MM:SS' (JST) → その日の分。秒は切り捨て (画面の HH:MM と揃える)。"""
    if not ts:
        return None
    return int(ts[11:13]) * 60 + int(ts[14:16])


def hm(m: int | None) -> str:
    return "" if m is None else f"{m // 60:02d}:{m % 60:02d}"


def normalize(rows: list[dict]) -> list[dict]:
    out = []
    for r in rows:
        d = dt.date.fromisoformat(r["visit_date"])
        arr, dep = minute_of(r["arr_at"]), minute_of(r["dep_at"])
        stay = dep - arr if arr is not None and dep is not None else None
        cancelled = r["status"] == "cancelled" or r["deleted"]
        remarks = []
        if arr is not None and dep is None:
            remarks.append("退出なし")
        if r["is_unplanned"]:
            remarks.append("予定外の訪問")
        if r["substitute"]:
            remarks.append(f"代行（予定: {r['planned_staff'] or '未割当'}）")
        if r["arr_match"] in ("review", "mismatch", "no_gps") and arr is not None:
            remarks.append("場所 要確認")
        if cancelled and arr is not None:
            remarks.append("取消済みの予定に記録")
        if stay is not None and stay < SHORT_STAY_MIN:
            remarks.append("到着と退出が近い")
        out.append(
            {
                "date": d,
                "wd": WD[d.weekday()],
                "office": r["office"] or "",
                "patient": r["patient"],
                # 予定外は打刻時刻がそのまま予定欄に入っているので、予定としては出さない。
                "ps": None if r["is_unplanned"] else r["start_time"],
                "pe": None if r["is_unplanned"] else r["end_time"],
                "planned": None if r["is_unplanned"] else r["planned_staff"],
                "nurse": r["arr_staff"] or r["dep_staff"],
                "arr": arr,
                "dep": dep,
                "stay": stay,
                "remarks": remarks,
                "cancelled": cancelled,
                "unplanned": bool(r["is_unplanned"]),
            }
        )
    out.sort(key=lambda x: (x["date"], x["ps"] or hm(x["arr"]), x["patient"]))
    return out


def summarize(rows: list[dict]) -> OrderedDict:
    """看護師ごとの件数。予定は「予定の担当」、読み取りは「実際に読んだ人」で数える。"""
    names = sorted({r["planned"] for r in rows if r["planned"]} | {r["nurse"] for r in rows if r["nurse"]})
    s: OrderedDict = OrderedDict()
    for n in names:
        planned = [r for r in rows if r["planned"] == n and not r["cancelled"]]
        mine = [r for r in rows if r["nurse"] == n and r["arr"] is not None]
        s[n] = {
            "planned": len(planned),
            "arr": len(mine),
            "dep": sum(1 for r in mine if r["dep"] is not None),
            "nodep": sum(1 for r in mine if r["dep"] is None),
        }
    return s


# ────────────────────────── Excel ──────────────────────────
HEAD_FILL = PatternFill("solid", fgColor="F3F6F8")
THIN = Side(style="thin", color="D7DDE2")
COLS = ["日付", "曜日", "拠点", "利用者", "予定開始", "予定終了", "予定の担当", "訪問した看護師", "到着", "退出", "滞在(分)", "備考"]
WIDTHS = [11, 5, 7, 18, 9, 9, 14, 14, 8, 8, 9, 38]


def _time(v: str | int | None):
    if v is None or v == "":
        return None
    if isinstance(v, int):
        return dt.time(v // 60, v % 60)
    return dt.time(int(v[:2]), int(v[3:5]))


def _sheet(ws, rows: list[dict], with_state: bool) -> None:
    cols = COLS + (["QR 読み取り"] if with_state else [])
    ws.append(cols)
    for r in rows:
        line = [r["date"], r["wd"], r["office"], r["patient"], _time(r["ps"]), _time(r["pe"]), r["planned"], r["nurse"],
                _time(r["arr"]), _time(r["dep"]), r["stay"], "、".join(r["remarks"])]
        if with_state:
            line.append("あり" if r["arr"] is not None else "なし")
        ws.append(line)
    for i, w in enumerate(WIDTHS + ([11] if with_state else []), 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for c in ws[1]:
        c.font, c.fill, c.alignment = Font(bold=True), HEAD_FILL, Alignment(horizontal="center")
        c.border = Border(bottom=Side(style="thin", color="9AA7B0"))
    for row in ws.iter_rows(min_row=2):
        row[0].number_format = "yyyy/mm/dd"
        for i in (4, 5, 8, 9):
            row[i].number_format = "hh:mm"
            row[i].alignment = Alignment(horizontal="center")
        row[1].alignment = Alignment(horizontal="center")
        for i in (8, 9):
            row[i].font = Font(bold=True)
        for c in row:
            c.border = Border(bottom=THIN)
    ws.freeze_panes = "E2"
    ws.auto_filter.ref = ws.dimensions
    ws.page_setup.orientation, ws.page_setup.paperSize = "landscape", ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "1:1"


def build_xlsx(rows: list[dict], summary: OrderedDict, label: str, notes: list[str], path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "QR読み取りあり"
    _sheet(ws, [r for r in rows if r["arr"] is not None], with_state=False)
    ws2 = wb.create_sheet("全予定")
    _sheet(ws2, rows, with_state=True)
    ws3 = wb.create_sheet("看護師別")
    ws3.append(["看護師", "予定の件数", "到着の読み取り", "退出の読み取り", "退出なし"])
    for n, s in summary.items():
        ws3.append([n, s["planned"], s["arr"], s["dep"], s["nodep"]])
    for c in ws3[1]:
        c.font, c.fill = Font(bold=True), HEAD_FILL
    for i, w in enumerate([16, 12, 14, 14, 10], 1):
        ws3.column_dimensions[get_column_letter(i)].width = w
    ws4 = wb.create_sheet("読み方")
    ws4.column_dimensions["A"].width = 110
    ws4.append([f"訪問時刻の記録（QR 読み取り） {label}　{ORG}"])
    ws4["A1"].font = Font(bold=True, size=13)
    for n in notes:
        ws4.append([n])
    for row in ws4.iter_rows(min_row=2):
        row[0].alignment = Alignment(wrap_text=True, vertical="top")
    wb.save(path)


# ────────────────────────── HTML (A4 縦) ──────────────────────────
HTML = r"""<!doctype html>
<html lang="ja"><head><meta charset="utf-8">
<title>訪問時刻の記録（QR 読み取り） __LABEL__</title>
<style>
html{color-scheme:light}
body{margin:0;background:#e9ecef;color:#1e2a33;font-family:"BIZ UDPGothic","Noto Sans JP","Hiragino Sans","Yu Gothic UI",Meiryo,sans-serif;font-size:9pt;line-height:1.5}
.bar{position:sticky;top:0;background:#fff;border-bottom:1px solid #d7dde2;padding:8px 16px;display:flex;gap:12px;align-items:center;font-size:13px;z-index:2}
.bar button{height:32px;padding:0 14px;border-radius:6px;border:1px solid #1e2a33;background:#1e2a33;color:#fff;font-size:13px;cursor:pointer}
.sheet{background:#fff;width:210mm;height:297mm;margin:8mm auto;padding:12mm 13mm 10mm;box-sizing:border-box;box-shadow:0 2px 14px rgba(20,30,40,.12);display:flex;flex-direction:column}
.sheet header{border-bottom:2px solid #1e2a33;padding-bottom:2mm;margin-bottom:2mm;display:flex;justify-content:space-between;align-items:flex-end;gap:6mm}
.sheet header h1{font-size:14pt;margin:0;letter-spacing:.02em}
.sheet header .org{font-size:8.6pt;color:#55656f;text-align:right;line-height:1.4}
.body{flex:1;min-height:0;overflow:hidden}
h2{font-size:10.6pt;margin:3.5mm 0 1.4mm;padding-left:2.5mm;border-left:3px solid #1e2a33}
h2 small{font-size:8.4pt;font-weight:400;color:#55656f;margin-left:3mm}
p{margin:1mm 0 2mm}
.lead{background:#fff7e6;border-left:4px solid #9a6b00;padding:2.5mm 4mm;margin:1mm 0 3mm;font-size:9.4pt}
.kpi{display:flex;gap:3mm;margin:0 0 3mm;flex-wrap:wrap}
.kpi span{background:#f3f6f8;border:1px solid #d7dde2;padding:1.5mm 3.5mm;border-radius:2mm;font-size:8.8pt}
.kpi b{font-size:12pt;margin-left:1mm}
ul{margin:1mm 0 2mm;padding-left:5mm} li{margin-bottom:.8mm}
table{width:100%;border-collapse:collapse;font-variant-numeric:tabular-nums;font-size:8.8pt;table-layout:fixed}
th{text-align:left;font-size:8pt;color:#55656f;background:#f3f6f8;border-bottom:1px solid #9aa7b0;padding:1.2mm 1.6mm;white-space:nowrap}
td{border-bottom:1px solid #d7dde2;padding:1.2mm 1.6mm;white-space:nowrap;overflow:hidden;text-overflow:ellipsis}
td.b{font-weight:700;font-size:9.6pt} td.n{text-align:right} td.rm{white-space:normal;color:#55656f;font-size:8pt}
.sheet footer{border-top:1px solid #d7dde2;padding-top:1.5mm;margin-top:2mm;font-size:7.8pt;color:#55656f;display:flex;justify-content:space-between;gap:6mm}
@page{size:A4 portrait;margin:0}
@media print{body{background:#fff}.bar{display:none}.sheet{margin:0;box-shadow:none;height:296mm;break-after:page}.sheet:last-child{break-after:auto}}
</style></head><body>
<div class="bar"><b>訪問時刻の記録（QR 読み取り） __LABEL__</b><span id="info"></span><span style="flex:1"></span><button onclick="window.print()">印刷 / PDF に保存</button></div>
<div id="sheets"></div>
<script>
const DATA = __DATA__;
const root = document.getElementById('sheets');
let body, tbody;
// 行を 1 つずつ流し込み、用紙からあふれたら次の用紙へ送る (固定行数で切らない)。
function newSheet() {
  const s = document.createElement('div'); s.className = 'sheet';
  s.innerHTML = `<header><h1>訪問時刻の記録（QR 読み取り）　${DATA.label}</h1><div class="org">${DATA.org}<br>出力 ${DATA.generated}</div></header><div class="body"></div><footer><span>${DATA.footer}</span><span class="pg"></span></footer>`;
  root.appendChild(s); body = s.querySelector('.body'); tbody = null;
}
const over = () => body.scrollHeight > body.clientHeight + 1;
function addBlock(html) {
  const d = document.createElement('div'); d.innerHTML = html; body.appendChild(d);
  if (over() && body.children.length > 1) { d.remove(); newSheet(); body.appendChild(d); }
}
function startTable(t, cont) {
  const els = [];
  if (t.title) { const h = document.createElement('h2'); h.innerHTML = `${t.title}${cont ? '（続き）' : ''}<small>${t.sub || ''}</small>`; body.appendChild(h); els.push(h); }
  const tb = document.createElement('table');
  tb.innerHTML = `<thead><tr>${t.cols.map(c => `<th${c[1] ? ` style="width:${c[1]}"` : ''}>${c[0]}</th>`).join('')}</tr></thead><tbody></tbody>`;
  body.appendChild(tb); els.push(tb); tbody = tb.querySelector('tbody'); return els;
}
function addTable(t) {
  let els = startTable(t, false);
  for (const r of t.rows) {
    const tr = document.createElement('tr');
    tr.innerHTML = r.map((c, i) => `<td class="${t.cls[i] || ''}">${c}</td>`).join('');
    tbody.appendChild(tr);
    if (over()) {
      tr.remove();
      if (!tbody.children.length) els.forEach(e => e.remove());
      newSheet(); els = startTable(t, true); tbody.appendChild(tr);
    }
  }
}
newSheet();
for (const b of DATA.blocks) b.type === 'html' ? addBlock(b.html) : addTable(b);
const sheets = root.querySelectorAll('.sheet');
sheets.forEach((s, i) => s.querySelector('.pg').textContent = `${i + 1} / ${sheets.length}`);
document.getElementById('info').textContent = `A4 縦 ${sheets.length} 枚`;
</script></body></html>
"""


def esc(s: str | None) -> str:
    return (s or "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def build_html(rows: list[dict], summary: OrderedDict, label: str, start_note: str, notes: list[str], generated: str, path: Path) -> None:
    hit = [r for r in rows if r["arr"] is not None]
    active = [r for r in rows if not r["cancelled"]]
    nodep = sum(1 for r in hit if r["dep"] is None)
    blocks: list[dict] = [
        {"type": "html", "html": f'<div class="lead">{esc(start_note)}</div>'
            f'<div class="kpi"><span>この月の訪問<b>{len(active)}</b> 件</span><span>到着の読み取り<b>{len(hit)}</b> 件</span>'
            f'<span>退出まで読み取り<b>{len(hit) - nodep}</b> 件</span><span>退出なし<b>{nodep}</b> 件</span></div>'},
        {"type": "html", "html": "<h2>この表の読み方</h2><ul>" + "".join(f"<li>{esc(n)}</li>" for n in notes) + "</ul>"},
        {"type": "table", "title": "看護師別の件数", "sub": "予定は担当として組まれた件数、読み取りは実際に QR を読んだ件数",
         "cols": [["看護師", "42mm"], ["予定の件数", "28mm"], ["到着の読み取り", "30mm"], ["退出の読み取り", "30mm"], ["退出なし", ""]],
         "cls": ["", "n", "n", "n", "n"],
         "rows": [[esc(n), s["planned"], s["arr"], s["dep"], s["nodep"]] for n, s in summary.items()]},
    ]
    for n in summary:
        mine = [r for r in hit if r["nurse"] == n]
        if not mine:
            continue
        blocks.append({
            "type": "table", "title": esc(n), "sub": f"{len(mine)} 件",
            "cols": [["日付", "19mm"], ["利用者", "34mm"], ["予定", "25mm"], ["到着", "15mm"], ["退出", "15mm"], ["滞在", "14mm"], ["備考", ""]],
            "cls": ["", "", "", "b", "b", "", "rm"],
            "rows": [[f"{r['date'].month}/{r['date'].day} ({r['wd']})", esc(r["patient"]),
                      f"{r['ps']}–{r['pe']}" if r["ps"] else "—", hm(r["arr"]) or "—", hm(r["dep"]) or "—",
                      f"{r['stay']}分" if r["stay"] is not None else "—", esc("、".join(r["remarks"]))] for r in mine],
        })
    data = {"label": label, "org": ORG, "generated": generated, "blocks": blocks,
            "footer": "到着・退出は QR を読み取った時刻です（家に入ってから読むため、実際の到着より数分遅いことがあります）。"}
    path.write_text(HTML.replace("__LABEL__", label).replace("__DATA__", json.dumps(data, ensure_ascii=False)), encoding="utf-8")


def main() -> None:
    src, month, out = sys.argv[1], sys.argv[2], Path(sys.argv[3])
    y, m = map(int, month.split("-"))
    label = f"{y} 年 {m} 月"
    rows = normalize(load(src))
    summary = summarize(rows)
    first = min((r["date"] for r in rows if r["arr"] is not None), default=None)
    hit = sum(1 for r in rows if r["arr"] is not None)
    total = sum(1 for r in rows if not r["cancelled"])
    start_note = (f"QR の読み取りは運用を始めたばかりで、この月は {total} 件の訪問のうち {hit} 件に到着の記録があります。"
                  "記録のない訪問は「訪問していない」という意味ではありません。紙の時間確認表とあわせてご確認ください。")
    notes = [
        "到着・退出は、スタッフが利用者宅の QR を読み取った時刻です。家に入ってから読むため、インターホン待ちなどの時間は含まれず、実際の到着より数分遅いことがあります。",
        "「訪問した看護師」は実際に QR を読み取った人です。予定の担当と違う場合は備考に「代行」と出ます。",
        "「退出なし」は到着だけ読み取り、退出の読み取りが無い訪問です。滞在時間は計算していません。",
        "「予定外の訪問」は予定に無い訪問を QR で記録したものです。予定の欄は空です。",
        f"「到着と退出が近い」は間が {SHORT_STAY_MIN} 分未満のものです。訪問の後にまとめて読み取った可能性があります。",
        "「場所 要確認」は読み取り時の位置が利用者宅から離れていた、または位置が取れなかったものです。",
        "「取消済みの予定に記録」は、読み取りの後に予定の側が取り消されたものです。訪問の事実として載せています。",
        "時刻は分単位（秒は切り捨て）です。",
    ]
    if first:
        notes.insert(0, f"この月で最初の読み取りは {first.month}/{first.day} です。QR カードを配布した 9/18 より前は、ほとんど記録がありません。")
    generated = dt.datetime.now().strftime("%Y/%m/%d %H:%M")
    out.parent.mkdir(parents=True, exist_ok=True)
    build_html(rows, summary, label, start_note, notes, generated, out.with_suffix(".html"))
    build_xlsx(rows, summary, label, [start_note] + notes, out.with_suffix(".xlsx"))
    print(f"rows={len(rows)} arrival={hit} nurses={len(summary)} -> {out.with_suffix('.html').name}, {out.with_suffix('.xlsx').name}")


if __name__ == "__main__":
    main()
