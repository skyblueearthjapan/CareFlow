"""ならし案の検査 (計算とは別に決まりを数え直す) と出力 (Excel・A4)。利用者名入り = git に入れない・公開しない。"""

from __future__ import annotations

import html
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from solver import BUFFER_MIN, LUNCH_MIN_HARD, LUNCH_WIN, bucket, fmt, travel_min

WD = "月火水木金土日"


def day_label(day) -> str:
    return f"{day:%m/%d}({WD[day.weekday()]})"


def new_assignments(out) -> dict:
    """visit_id -> (staff_id, start)"""
    new = {}
    for _day, (_vehicles, routes, _dropped) in out.items():
        for sid, seq in routes.items():
            for j, t in seq:
                for v, off in zip(j.visits, j.offsets, strict=True):
                    new[v.id] = (sid, t + off)
    return new


def _hm(t: str) -> int:
    return int(t[:2]) * 60 + int(t[3:5])


def staff_hours(d, staff_id: str, day) -> tuple[int, int] | None:
    """その日の勤務 (開始, 終了)。計算とは別に、取り出した元のデータから組み立て直す。"""
    wd = day.weekday()
    hours = None
    for r in d.get("shifts") or []:
        if (
            r["staff_id"] == staff_id
            and r["weekday"] == wd
            and r["is_on"]
            and r["start_time"]
            and r["end_time"]
        ):
            hours = (_hm(r["start_time"]), _hm(r["end_time"]))
    for r in d.get("weekly_overrides") or []:
        if r["staff_id"] != staff_id or r["weekday"] != wd:
            continue
        kind = r["override_type"]
        if kind == "off":
            return None
        if kind == "custom_time" and r.get("start_time") and r.get("end_time"):
            hours = (_hm(r["start_time"]), _hm(r["end_time"]))
        elif kind == "am_off" and hours:
            hours = (max(hours[0], 12 * 60), hours[1])
        elif kind == "pm_off" and hours:
            hours = (hours[0], min(hours[1], 12 * 60))
    if hours and hours[0] >= hours[1]:
        return None
    return hours


def staff_events(d, staff_id: str, day) -> list[tuple[int, int]]:
    """その日の予定 (前後 15 分込み = 決まりの正典 §4・計算側の EVENT_BUFFER_MIN と同じ値)。

    日をまたぐ予定は、初日は開始〜24:00、間の日は終日、最終日は 0:00〜終了として数える。
    """
    out = []
    for e in d.get("events") or []:
        if e["staff_id"] != staff_id:
            continue
        first, last = e["date"], e.get("end_date") or e["date"]
        if not first <= str(day) <= last:
            continue
        a = _hm(e["start"]) if str(day) == first else 0
        b = _hm(e["end"]) if str(day) == last else 24 * 60
        if b > a:
            out.append((a - 15, b + 15))
    return out


def lunch_free(seq) -> int:
    """11:30〜13:30 の中で取れる一番長い昼休み (分)。

    訪問と訪問の間では、移動 (＋ゆとり) を昼休みの前にも後にも置ける。窓の外の空きで
    移動を済ませられれば、その分だけ窓の中の空きが昼休みに使える。
    """
    lo, hi = LUNCH_WIN
    gaps = []  # (空きの始まり, 終わり, その間に要る移動)
    prev_end = None
    for k, (j, t) in enumerate(seq):
        if k == 0:
            gaps.append((lo, t, 0))  # 最初の訪問の前 (始点からの移動は数えない)
        else:
            pj = seq[k - 1][0]
            tr = travel_min(pj.visits[0].loc, j.visits[0].loc)
            gaps.append((prev_end, t, tr + (BUFFER_MIN if tr else 0)))
        prev_end = t + j.dur
    gaps.append((prev_end if seq else lo, hi, 0))  # 最後の訪問の後
    best = 0
    for s0, e0, need in gaps:
        a0, b0 = max(s0, lo), min(e0, hi)
        if b0 <= a0:
            continue
        outside = (a0 - s0) + (e0 - b0)  # 窓の外で移動に使える時間
        best = max(best, (b0 - a0) - max(0, need - outside))
    return best


def check(d, out, cap_regular: int, allow_over: int) -> tuple[list[str], list[str]]:
    """決まりを数え直す (計算の結果を元データと突き合わせる)。(問題, 気を付ける点) を返す。"""
    errors, warnings = [], []
    for day, (vehicles, routes, dropped) in out.items():
        vmap = {s.id: s for s in vehicles}
        for sid, seq in routes.items():
            if not seq:
                continue
            s = vmap[sid]
            hours = staff_hours(d, sid, day)
            if hours is None:
                errors.append(f"{day} {s.name}: 休みの日に訪問がある")
                continue
            st, en = hours
            hard_end = max(en, 18 * 60) if s.manager else en
            cnt = sum(j.count for j, _ in seq)
            if not s.manager and cnt > cap_regular + allow_over:
                errors.append(f"{day} {s.name}: {cnt} 名（上限超え）")
            elif not s.manager and cnt > cap_regular:
                warnings.append(f"{day} {s.name}: {cnt} 名（{cap_regular} 名を超える）")
            if seq[0][1] < st:
                errors.append(f"{day} {s.name}: 勤務開始前 {fmt(seq[0][1])}")
            last_end = seq[-1][1] + seq[-1][0].dur
            if last_end > hard_end:
                errors.append(f"{day} {s.name}: 終わりが遅すぎる {fmt(last_end)}")
            elif s.manager and last_end > en:
                warnings.append(f"{day} {s.name}: 勤務時刻を超える {fmt(last_end)}")
            for (a, ta), (b, tb) in zip(seq, seq[1:]):  # noqa: B905
                tr = travel_min(a.visits[0].loc, b.visits[0].loc)
                if tb < ta + a.dur + tr + (BUFFER_MIN if tr else 0):
                    errors.append(f"{day} {s.name}: {fmt(ta)}→{fmt(tb)} 間に合わない")
            evs = staff_events(d, sid, day)
            for j, t in seq:
                for ea, eb in evs:
                    if t < eb and ea < t + j.dur:
                        errors.append(f"{day} {s.name}: {fmt(t)} の訪問が予定と重なる")
                if j.count == 2:
                    a, b = j.visits
                    if j.offsets[1] + b.dur > j.dur or a.dur > j.dur:
                        errors.append(f"{day} {s.name}: 同じ建物の 2 名が枠に収まらない")
                for v, off in zip(j.visits, j.offsets, strict=True):
                    if sid in v.ng:
                        errors.append(f"{day} {s.name}: NG の職員")
                    if v.female_only and s.sex == "male":
                        errors.append(f"{day} {s.name}: 女性限定に男性")
                    if not (v.win[0] <= t + off <= v.win[1]):
                        errors.append(f"{day} {s.name}: 希望の範囲外 {fmt(t + off)}")
                    if (t + off) % 5:
                        errors.append(f"{day} {s.name}: 5 分刻みでない {fmt(t + off)}")
            if st <= LUNCH_WIN[0] and hard_end >= LUNCH_WIN[1]:
                gap = lunch_free(seq)
                if gap < LUNCH_MIN_HARD:
                    errors.append(f"{day} {s.name}: 昼休みが 30 分取れない")
                elif gap < 45:
                    warnings.append(f"{day} {s.name}: 昼休みが {gap} 分（45 分未満）")
        for j in dropped:
            errors.append(f"{day}: 入らない訪問 {j.count} 件")
    return errors, warnings


def board_stats(visits, assign):
    """職員×日の件数・重なり・間に合わない・移動。同じ建物は同時 2 名まで重なってよい。"""
    by_id = {v.id: v for v in visits}
    per = defaultdict(list)
    for vid, (sid, t) in assign.items():
        if sid is not None:
            per[(by_id[vid].day, sid)].append((t, by_id[vid]))
    overlaps, late, travel, served = [], 0, 0, 0
    for key, lst in per.items():
        lst.sort(key=lambda x: x[0])
        served += len(lst)
        active = []  # (終了, 訪問)
        for t, v in lst:
            active = [(e, w) for e, w in active if e > t]
            same = [w for _e, w in active if bucket(w.loc) == bucket(v.loc)]
            other = [w for _e, w in active if bucket(w.loc) != bucket(v.loc)]
            if other or len(same) >= 2:
                overlaps.append((key, (other or same)[0], v, t))
            active.append((t + v.dur, v))
        for (ta, a), (tb, b) in zip(lst, lst[1:]):  # noqa: B905
            tr = travel_min(a.loc, b.loc)
            travel += tr
            if tr and ta + a.dur <= tb < ta + a.dur + tr + BUFFER_MIN:
                late += 1
    return per, overlaps, late, travel, served


def flags_for(out, hist_by_day) -> dict:
    """visit_id -> 気を付ける点の一覧 (Jev の確認と札の材料)"""
    flags = defaultdict(list)
    for day, (vehicles, routes, _dropped) in out.items():
        vmap = {s.id: s for s in vehicles}
        hist = hist_by_day.get(day, {})
        for sid, seq in routes.items():
            s = vmap[sid]
            en = s.shift[day.weekday()][1]
            for j, t in seq:
                for v, off in zip(j.visits, j.offsets, strict=True):
                    st = t + off
                    if st != v.fixed_start:
                        flags[v.id].append(("move", st - v.fixed_start))
                    if v.office not in s.offices:
                        flags[v.id].append(("cross", 1))
                    if s.manager and st + v.dur > en:
                        flags[v.id].append(("overtime", st + v.dur - en))
                    if (hist.get(v.patient_id) or [None])[0] == sid:
                        flags[v.id].append(("same_as_last", 1))
    return flags


def write_excel(
    path: Path, title: str, d, visits, out, errors, warnings, cap_regular, jev=None, diag=None
):
    S = {s["id"]: s["name"] for s in d["staff"]}
    new = new_assignments(out)
    cur = {v.id: (v.cur_staff, v.fixed_start) for v in visits}
    cur_per, cur_ov, cur_late, cur_tr, cur_n = board_stats(visits, cur)
    new_per, new_ov, new_late, new_tr, new_n = board_stats(visits, new)
    over_days = sum(
        1
        for (_day, sid), lst in new_per.items()
        if len(lst) > cap_regular
        and not any(s.id == sid and s.manager for _d, (vh, _r, _x) in out.items() for s in vh)
    )
    jev = jev or {}

    wb = Workbook()
    bold, hdr = Font(bold=True), PatternFill("solid", fgColor="E7F4F0")
    warn = PatternFill("solid", fgColor="FBF1DD")

    def sheet(name, header, rows, widths):
        ws = wb.create_sheet(name)
        ws.append(header)
        for c in ws[1]:
            c.font, c.fill = bold, hdr
        for r in rows:
            ws.append(r)
        for i, w in enumerate(widths):
            ws.column_dimensions[chr(65 + i)].width = w
        ws.freeze_panes = "A2"
        return ws

    ws = wb.active
    ws.title = "概要"
    ws.append([title + " — 案（まだ反映していません）"])
    ws["A1"].font = Font(bold=True, size=12)
    ws.append([f"本番の写し {d.get('extracted_at', '')[:16]} から計算。本番は変えていません。"])
    ws.append([])
    ws.append(["", "今の盤面", "ならし案"])
    ws.append(["同じ職員の時間の重なり（件・同じ建物の 2 名は除く）", len(cur_ov), len(new_ov)])
    ws.append(["移動が間に合わない並び（件）", cur_late, new_late])
    ws.append(
        [
            "担当なしの訪問（件）",
            sum(1 for v in visits if v.cur_staff is None),
            sum(1 for v in visits if v.id not in new),
        ]
    )
    ws.append(["担当のある訪問（件）", cur_n, new_n])
    ws.append(["移動の合計（分・担当のある訪問だけ）", cur_tr, new_tr])
    ws.append(
        [
            "1 件あたりの移動（分）",
            round(cur_tr / cur_n, 1) if cur_n else "",
            round(new_tr / new_n, 1) if new_n else "",
        ]
    )
    ws.append([f"正規が {cap_regular} 名を超える日（人・日）", "—", over_days])
    ws.append(["決まりの検査で見つかった問題", "—", len(errors)])
    ws.append(["気を付ける点（勤務超え・昼休み 45 分未満など）", "—", len(warnings)])
    if diag:
        ws.append(["アプリの物差し: 移動が平均の 1.5 倍を超える（人・日）", len(diag["cur_high_days"]), len(diag["new_high_days"])])
        ws.append(["アプリの物差し: あと 1 手で良くなる所（見落とし）", "—", diag["missed"]])
    ws.append([])
    staff_order = []
    for _day, (vehicles, _r, _dr) in sorted(out.items()):
        for s in vehicles:
            if s not in staff_order:
                staff_order.append(s)
    ws.append(["日", "訪問"] + [s.name for s in staff_order])
    for day in sorted(out):
        row = [day_label(day), sum(1 for v in visits if v.day == day)]
        for s in staff_order:
            n, c = len(new_per.get((day, s.id), [])), len(cur_per.get((day, s.id), []))
            row.append(f"{n}（今 {c}）" if (n or c) else "")
        ws.append(row)
    ws.column_dimensions["A"].width = 32
    for i in range(1, len(staff_order) + 2):
        ws.column_dimensions[chr(65 + i)].width = 13

    rows = []
    for v in sorted(visits, key=lambda v: (v.day, v.fixed_start)):
        sid, st = new.get(v.id, (None, None))
        if sid is not None and sid == v.cur_staff and st == v.fixed_start:
            continue
        rows.append(
            [
                day_label(v.day),
                v.patient,
                fmt(v.fixed_start),
                S.get(v.cur_staff, "担当なし"),
                fmt(st) if st is not None else "入らない",
                S.get(sid, ""),
                (st - v.fixed_start) if st is not None and st != v.fixed_start else "",
                jev.get(v.id, ""),
            ]
        )
    ws2 = sheet(
        "変更一覧",
        [
            "日",
            "利用者",
            "今の時刻",
            "今の担当",
            "案の時刻",
            "案の担当",
            "ずらす(分)",
            "Jev: 要確認の度合い",
        ],
        rows,
        [11, 22, 9, 14, 9, 14, 10, 18],
    )
    for r in ws2.iter_rows(min_row=2):
        if isinstance(r[7].value, str) and r[7].value.startswith("要確認"):
            for c in r:
                c.fill = warn

    rows = [
        [day_label(k[0]), S.get(k[1]), f"{fmt(a.fixed_start)} {a.patient}", f"{fmt(t)} {b.patient}"]
        for k, a, b, t in sorted(cur_ov, key=lambda x: (x[0][0], x[3]))
    ]
    sheet("今の重なり", ["日", "職員", "先の訪問", "重なる訪問"], rows, [11, 14, 28, 28])
    sheet(
        "決まりの検査",
        ["結果"],
        [[e] for e in errors]
        or [
            [
                "問題なし（人数・勤務時刻・休み・予定・移動・NG・女性限定・希望の範囲・5 分刻み・昼休み・同じ建物の枠）"
            ]
        ],
        [90],
    )
    sheet("気を付ける点", ["内容"], [[w] for w in warnings] or [["なし"]], [90])
    if diag:
        _diag_sheet(wb, diag, bold, hdr, warn)
    wb.save(path)
    return {
        "cur_overlaps": len(cur_ov),
        "new_overlaps": len(new_ov),
        "cur_late": cur_late,
        "new_late": new_late,
        "cur_served": cur_n,
        "new_served": new_n,
        "cur_travel_served_only": cur_tr,
        "new_travel": new_tr,
        "over_cap_staff_days": over_days,
    }


def write_a4(path: Path, title: str, d, visits, out, hist_by_day, jev=None, diag=None):
    S = {s["id"]: s["name"] for s in d["staff"]}
    e = html.escape
    jev = jev or {}
    pages = []
    for day in sorted(out):
        vehicles, routes, dropped = out[day]
        wd = day.weekday()
        cols = []
        for s in vehicles:
            seq = routes.get(s.id) or []
            st, en = s.shift[wd]
            cards, prev = [], None
            for j, t in seq:
                if prev is not None:
                    pj, pt = prev
                    tr = travel_min(pj.visits[0].loc, j.visits[0].loc)
                    free = t - (pt + pj.dur) - tr - (BUFFER_MIN if tr else 0)
                    cards.append(
                        f'<div class="mv">↓ 移動 {tr} 分{"＋ゆとり " + str(BUFFER_MIN) if tr else "（同じ建物）"}・空き {free} 分</div>'
                    )
                for v, off in zip(j.visits, j.offsets, strict=True):
                    vs = t + off
                    tags = []
                    if vs != v.fixed_start:
                        tags.append(f'<span class="tg t">時刻 {fmt(v.fixed_start)}→</span>')
                    if v.cur_staff != s.id:
                        tags.append(
                            f'<span class="tg s">今 {e(S.get(v.cur_staff, "担当なし").split(chr(12288))[0])}</span>'
                        )
                    if v.office not in s.offices:
                        tags.append('<span class="tg x">拠点またぎ</span>')
                    if s.manager and vs + v.dur > en:
                        tags.append('<span class="tg o">勤務超え</span>')
                    if (hist_by_day.get(day, {}).get(v.patient_id) or [None])[0] == s.id:
                        tags.append('<span class="tg r">前回と同じ</span>')
                    if str(jev.get(v.id, "")).startswith("要確認"):
                        tags.append('<span class="tg j">Jev 要確認</span>')
                    pair = '<span class="pr">同じ建物</span>' if j.count == 2 else ""
                    cards.append(
                        f'<div class="cd"><div class="tm">{fmt(vs)}〜{fmt(vs + v.dur)} {pair}</div>'
                        f'<div class="pt">{e(v.patient)} <small>{e(v.office)}・{e(v.time_type)}</small></div>'
                        f"<div>{''.join(tags)}</div></div>"
                    )
                prev = (j, t)
            cnt = sum(j.count for j, _ in seq)
            role = "マネージャー" if s.manager else ""
            cols.append(
                f'<div class="col{" mg" if s.manager else ""}"><div class="hd">{e(s.name)} <small>{role}</small>'
                f'<div class="sub">{fmt(st)}〜{fmt(en)}・{cnt} 名・昼の空き {lunch_free(seq) if seq else 120} 分</div></div>'
                f"{''.join(cards) or '<div class=em>訪問なし</div>'}</div>"
            )
        drop = "".join(
            f"<li>{e(v.patient)} {fmt(v.fixed_start)}</li>" for j in dropped for v in j.visits
        )
        pages.append(
            f'<section class="pg"><h2>{day:%Y年%m月%d日}（{WD[wd]}）<small>訪問 {sum(1 for v in visits if v.day == day)} 件</small></h2>'
            f'<div class="grid" style="grid-template-columns:repeat({max(1, len(cols))},1fr)">{"".join(cols)}</div>'
            + (f'<div class="dr">入らない: <ul>{drop}</ul></div>' if drop else "")
            + "</section>"
        )
    path.write_text(
        f"""<!doctype html><html lang="ja"><head><meta charset="utf-8"><title>{e(title)}</title>
<style>
@page {{ size: A4 landscape; margin: 8mm; }}
body {{ font-family: "Noto Sans JP", "Yu Gothic", sans-serif; color:#2b2620; background:#fff; margin:0; font-size:9.5px; }}
.pg {{ page-break-after: always; padding: 4mm 2mm; }}
h1 {{ font-size:15px; margin:4mm 2mm 0; }} .note {{ margin:1mm 2mm 3mm; color:#6b5e54; }}
h2 {{ font-size:13px; margin:0 0 2mm; }} h2 small {{ font-weight:normal; color:#6b5e54; margin-left:6px; }}
.grid {{ display:grid; gap:2mm; align-items:start; }}
.col {{ border:1px solid #e2dacd; border-radius:4px; padding:1.5mm; }} .col.mg {{ background:#fdf5f6; }}
.hd {{ font-weight:bold; font-size:11px; border-bottom:1px solid #e2dacd; padding-bottom:1mm; margin-bottom:1mm; }}
.hd small {{ color:#b5476b; font-weight:normal; }} .sub {{ font-weight:normal; font-size:9px; color:#6b5e54; }}
.cd {{ border:1px solid #cabfad; border-radius:3px; padding:1mm 1.5mm; margin:0.8mm 0; break-inside:avoid; }}
.tm {{ font-weight:bold; }} .pt {{ font-size:10px; }} .pt small {{ color:#6b5e54; }}
.mv {{ color:#9c9087; font-size:8.5px; text-align:center; }}
.tg {{ display:inline-block; border-radius:2px; padding:0 2px; margin:0.3mm 0.6mm 0 0; font-size:8.5px; }}
.t {{ background:#fbf1dd; }} .s {{ background:#e7f4f0; }} .x {{ background:#e8e4f7; }} .o {{ background:#fde2e4; }} .r {{ background:#eee; }} .j {{ background:#ffe08a; }}
.dg {{ border-collapse:collapse; margin:2mm; font-size:11px; }} .dg th, .dg td {{ border:1px solid #e2dacd; padding:1mm 3mm; text-align:right; }} .dg th:first-child {{ text-align:left; font-weight:normal; }}
h3 {{ font-size:11.5px; margin:3mm 2mm 1mm; }} ul {{ margin:0 2mm; padding-left:5mm; font-size:10px; }}
.pr {{ font-size:8.5px; color:#0b6e5e; font-weight:normal; }} .em {{ color:#9c9087; }} .dr {{ color:#b00; margin-top:2mm; }}
</style></head><body>
<h1>{e(title)} — 案（まだ反映していません）</h1>
<div class="note">決まり: 曜日は変えない／固定の時刻 → 希望の範囲内／正規の人数上限／勤務時刻を守る（マネージャーはあふれの受け皿）／NG・女性限定は絶対／同じ建物の 2 名は 1 人で 90 分枠／昼休み 11:30〜13:30／移動は直線 20km/h＋ゆとり 8 分。
札: <span class="tg t">時刻 →</span> 元の時刻／<span class="tg s">今 ○○</span> 今の担当／<span class="tg x">拠点またぎ</span>／<span class="tg o">勤務超え</span>／<span class="tg r">前回と同じ</span>／<span class="tg j">Jev 要確認</span></div>
{diag_page(diag) if diag else ""}{"".join(pages)}
</body></html>""",
        encoding="utf-8",
    )


# ---------------------------------------------------------------- アプリの物差し (diagnose.py)

DIAG_NOTE = (
    "アプリの「スケジュール診断」と同じ計算（直線距離・{speed:g}km/h・別の住所へ移るとき ゆとり {buf} 分・"
    "1 件目への移動は数えない）で、今の盤面とならし案を「職員 × 日 = 1 本の順路」として数えた。"
    "アプリの画面はコースごとに数えるので、画面の数字とは分け方が違う。"
)


def _diag_summary_rows(diag):
    c, n = diag["cur"], diag["new"]
    return [
        ["移動（分）", c["travel_minutes"], n["travel_minutes"]],
        ["移動の距離（km）", c["travel_km"], n["travel_km"]],
        ["ゆとり（分）", c["buffer_minutes"], n["buffer_minutes"]],
        ["待ち時間（分・移動しても余る空き）", c["gap_minutes"], n["gap_minutes"]],
        ["担当のある訪問（件）", c["visits"], n["visits"]],
        ["移動が その日の平均の 1.5 倍を超える（人・日）", len(diag["cur_high_days"]), len(diag["new_high_days"])],
        ["週の移動が 職員の平均の 1.5 倍を超える（人）", len(diag["cur_high_week"]), len(diag["new_high_week"])],
    ]


def _diag_sheet(wb, diag, bold, hdr, warn):
    ws = wb.create_sheet("アプリの物差し")
    cfg = diag["config"]
    ws.append([DIAG_NOTE.format(speed=cfg["travel_speed_kmh"], buf=cfg["visit_buffer_min"])])
    ws.append(
        [
            "物差しの一致（道具の移動の計算 = アプリの計算）: "
            + ("全順路で一致" if diag["yardstick_match"] else "食い違いあり " + "、".join(diag["yardstick_mismatch"]))
        ]
    )
    ws.append([])
    ws.append(["", "今の盤面", "ならし案"])
    for c in ws[ws.max_row]:
        c.font, c.fill = bold, hdr
    for r in _diag_summary_rows(diag):
        ws.append(r)
    ws.append([])
    ws.append(["移動が多い順路（案）", "職員", "移動(分)", "km", "その日の平均", "いちばん長い移動"])
    for c in ws[ws.max_row]:
        c.font, c.fill = bold, hdr
    for h in diag["new_high_days"] or []:
        ws.append([h["day"], h["staff"], h["travel_minutes"], h["travel_km"], h["day_avg"], h["longest_leg"]])
    if not diag["new_high_days"]:
        ws.append(["なし"])
    ws.append([])
    ws.append(["あと 1 手", "日", "手", f"減る移動＋ゆとり(分・{diag['threshold_min']} 分以上)", "選ばなかった理由"])
    for c in ws[ws.max_row]:
        c.font, c.fill = bold, hdr
    for mv in diag["moves"]:
        ws.append([mv["kind"], mv["label"], mv["move"], mv["saving_min"], mv["why_not"]])
        if mv["kind"] == "見落とし":
            for c in ws[ws.max_row]:
                c.fill = warn
    if not diag.get("moves_checked", True):
        ws.append(["確認していません（案が決まりの検査に落ちている・入らない訪問がある）"])
    elif not diag["moves"]:
        ws.append(["なし（決まりを守ったまま 1 手で良くなる所は残っていない）"])
    ws.append([])
    ws.append(["日", "職員", "今 件", "今 移動(分)", "今 km", "案 件", "案 移動(分)", "案 km", "案 待ち時間(分)"])
    for c in ws[ws.max_row]:
        c.font, c.fill = bold, hdr
    for r in diag["rows"]:
        ws.append(
            [r["day"], r["staff"], r["cur_visits"], r["cur_travel"], r["cur_km"],
             r["new_visits"], r["new_travel"], r["new_km"], r["new_gap"]]
        )
    for col, w in zip("ABCDEFGHI", [34, 14, 60, 22, 40, 9, 11, 9, 14]):  # noqa: B905
        ws.column_dimensions[col].width = w


def diag_page(diag) -> str:
    e = html.escape
    cfg = diag["config"]
    rows = "".join(
        f"<tr><th>{e(str(a))}</th><td>{e(str(b))}</td><td><b>{e(str(c))}</b></td></tr>"
        for a, b, c in _diag_summary_rows(diag)
    )
    hi = "".join(
        f"<li>{e(h['day'])} {e(h['staff'])}: 移動 {e(str(h['travel_minutes']))} 分（{e(str(h['travel_km']))} km・その日の平均 {e(str(h['day_avg']))} 分）"
        f" — {e(h['longest_leg'])}</li>"
        for h in diag["new_high_days"]
    ) or "<li>なし</li>"
    missed = [m for m in diag["moves"] if m["kind"] == "見落とし"]
    declined = [m for m in diag["moves"] if m["kind"] != "見落とし"]
    reasons = Counter(r for m in declined for r in m["why_not"].split("・") if r)
    mv = "".join(
        f"<li><b>見落とし</b> {e(m['label'])} {e(m['move'])}（−{m['saving_min']} 分）</li>" for m in missed[:10]
    ) + "".join(
        f"<li>{e(m['label'])} {e(m['move'])}（−{m['saving_min']} 分）— 選ばなかった理由: {e(m['why_not'])}</li>"
        for m in declined[:5]
    ) + (
        f"<li>わざと選ばなかった手は全部で {len(declined)} 件（理由: "
        + "、".join(f"{e(r)} {n} 件" for r, n in reasons.most_common())
        + "）。一覧は Excel の「アプリの物差し」</li>"
        if declined
        else ""
    ) or "<li>なし（決まりを守ったまま 1 手で良くなる所は残っていない）</li>"
    if not diag.get("moves_checked", True):
        verdict = "案が決まりの検査に落ちているので、あと 1 手の確認はしていません（先に検査の問題を直す）。"
    elif not missed:
        verdict = "アプリの物差しで見ても、決まりを守ったまま 1 手で良くなる所は残っていません。"
    else:
        verdict = (
            f"1 手で良くなる所が {len(missed)} 件残っています（計算の打ち切りの可能性）。"
            "--seconds を増やして計算し直すと取り込める可能性があります。計算し直した後にもう一度この診断を確かめてください。"
        )
    return f"""<section class="pg"><h2>アプリの物差しで見た案（スケジュール診断と同じ計算）</h2>
<div class="note">{e(DIAG_NOTE.format(speed=cfg["travel_speed_kmh"], buf=cfg["visit_buffer_min"]))}
物差しの一致: {"全順路で一致" if diag["yardstick_match"] else "食い違いあり"}</div>
<p style="font-size:12px;margin:2mm"><b>結論:</b> {e(verdict)}</p>
<table class="dg"><tr><th></th><th>今の盤面</th><th>ならし案</th></tr>{rows}</table>
<h3>移動が その日の平均の 1.5 倍を超える順路（案）</h3><ul>{hi}</ul>
<h3>あと 1 手（{diag["threshold_min"]} 分以上 減るもの・1 手ずつの評価で、同時に入れると足し算にならない）</h3><ul>{mv}</ul>
</section>"""
